"""Build the human-readable Word design and SLT PowerPoint from repository sources.

Optional authoring dependencies: python-docx and python-pptx. No business data.
"""
from pathlib import Path
import re
import shutil
from docx import Document
from docx.shared import Inches as DI, Pt as DP, RGBColor as DC
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'docs'/'delivery';OUT.mkdir(exist_ok=True)

def plain(s):
    return re.sub(r'\[([^]]+)\]\([^)]+\)',r'\1',s).replace('**','').replace('`','')

doc=Document();section=doc.sections[0]
section.top_margin=DI(.75);section.bottom_margin=DI(.7)
style=doc.styles['Normal'];style.font.name='Arial';style.font.size=DP(10)
style.paragraph_format.space_after=DP(7)
for name in ('Title','Heading 1','Heading 2','Heading 3'):
    doc.styles[name].font.name='Arial';doc.styles[name].font.color.rgb=DC.from_string('132C3B')
doc.add_paragraph('INV STUDIO',style='Subtitle')
doc.add_heading('From supplier invoice\nto checked Excel',0)
doc.add_paragraph('Solution design • operating model • verification record • build stories',style='Subtitle')
doc.add_paragraph('5 October 2026\nFinal revision inv-studio-api-00009-vxw has 100% traffic and uses the tested core image plus the cancel/footer fixes. Reader, workbook, lookup, deletion/model and cancel-dialog checks passed. The allowlist is owner-only, the owner account remained email-verified, unauthenticated API calls are denied and the temporary identity was deleted. The private 412,139-row evidence catalog completed atomic cutover; authenticated API and References UI checks passed without mutation. 214 backend tests pass with PostgreSQL included and no skips; Ruff, both Chromium workflows, all four GitHub checks and all three packaged reader warm-ups pass. Native and Paddle measurements remain bounded to their evaluated layouts; hosted adaptive Paddle reached 96/99 facts on one challenging two-page scan, while Docling recovered only 10/23 lines on that scan.')
doc.add_paragraph('For senior leaders, product owners, AP operators and the engineering team. Read the decision and process first; use the data contract and stories when building or accepting a release.')
doc.add_page_break()
for file in ['solution-design.md','data-contract.md','operator-guide.md','validation.md','stories.md']:
    lines=(ROOT/'docs'/file).read_text().splitlines();i=0;in_code=False
    while i<len(lines):
        line=lines[i]
        if line.startswith('```'):
            in_code=not in_code
            if in_code and file == 'solution-design.md':doc.add_paragraph('Flow: Upload → Confirm rules → Template or conservative native layout / OCR / permitted managed Vertex or selected AI → Match approved references → Review evidence → approved export and receipt reservation. Separate outputs: combined EXTRACTION_REVIEW_ONLY and all-field DRAFT_UNVALIDATED, both without approval or reservation.')
            i+=1;continue
        if in_code or not line.strip():i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                values=[plain(x.strip()) for x in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',x) for x in values):rows.append(values)
                i+=1
            table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1'
            for n,v in enumerate(rows[0]):table.rows[0].cells[n].text=v
            for row in rows[1:]:
                cells=table.add_row().cells
                for n,v in enumerate(row[:len(cells)]):cells[n].text=v
            continue
        if line.startswith('#'):
            depth=len(line)-len(line.lstrip('#'));doc.add_heading(plain(line.lstrip('# ').strip()),min(depth,3))
        elif line.startswith('- '):doc.add_paragraph(plain(line[2:]),style='List Bullet')
        else:doc.add_paragraph(plain(line))
        i+=1
    if file!='stories.md':doc.add_page_break()
section.footer.paragraphs[0].text='Inv Studio • Implementation and acceptance guide • Synthetic examples only'
doc.save(OUT/'Invoice_Studio_Design_and_Stories.docx')

prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
navy='112C3D';ink='163548';muted='53717B';teal='037E83';gold='E5BC66';white='FFFFFF';bg='F3F7F7'
def box(slide,x,y,w,h,text,size=22,color=ink,bold=False,fill=None):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)) if fill else slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    if fill:shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(fill);shape.line.fill.background()
    tf=shape.text_frame;tf.word_wrap=True;tf.margin_left=Inches(.12);tf.margin_right=Inches(.12)
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(12)
    return shape

def slide(title,kicker='INVOICE STUDIO / DECISION PACK'):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(bg)
    box(s,.6,.3,12,.3,kicker,11,teal,True);box(s,.6,.95,12,1.3,title,32,navy,True)
    box(s,.6,7.02,11,.3,'5 OCT 2026  •  MEASURED EVIDENCE AND PROPOSED NEXT STAGES',9,muted)
    box(s,12,6.95,.6,.35,str(len(prs.slides)).zfill(2),12,teal,True)
    return s

def cards(title,items,footer=None):
    s=slide(title);n=len(items);w=11.9/n
    for i,(head,body) in enumerate(items):
        x=.65+i*(w+.03);box(s,x,2.3,w-.15,3.8,'',fill=white)
        box(s,x+.15,2.52,w-.45,1.05,head,21,teal,True);box(s,x+.15,3.7,w-.45,2.2,body,19,ink)
    if footer:box(s,.7,6.35,11.8,.45,footer,14,muted)
    return s
s=slide('From supplier invoice\nto checked Excel','INV STUDIO / SLT PROPOSAL AND WORKING MVP')
box(s,.7,2.8,7.6,1.5,'Build the business controls in house.\nReuse local readers. Keep AI replaceable.',28,ink)
box(s,9,2.65,3.5,2.6,'01\nProve the invoice workflow first',28,white,True,teal)
box(s,.7,5.45,11.5,1.05,'Revision 00009-vxw is live. 214 tests, PostgreSQL, browsers, CI and reader warm-ups pass.',17,muted)
cards('The problem is bigger than reading text',[
 ('Identity','Which supplier site, buying company, market route and item does this invoice belong to?'),
 ('Evidence','Was the quantity accepted? Was it already invoiced? Does the unit price and tax match?'),
 ('Handoff','Can we generate the receiving workbook without assembling rows manually?')])
s=slide('The to-be process in six clear steps')
steps=[('01','Upload','One invoice per file; two workers, remaining files queued.'),('02','Confirm','Reader, AI fallback, references and rules.'),('03','Read','Template, native layout, OCR or permitted AI.'),('04','Validate','Approved references; lookup remains evidence only.'),('05','Review','Edit and save each invoice before batch selection.'),('06','Export','Review-only, approved or manual draft workbook.')]
for i,(num,title,body) in enumerate(steps):
    x=.65+(i%3)*4.15;y=2.3+(i//3)*2.05
    box(s,x,y,3.95,1.8,'',fill=white);box(s,x+.1,y+.1,3.7,.45,num+'  '+title,22,teal,True);box(s,x+.1,y+.65,3.7,1,body,18)
cards('Buy the readers. Build the decision layer.',[
 ('Local readers','invoice2data for repeatable layouts. PaddleOCR for scans. Docling for document structure.'),
 ('AI adapter','Managed Vertex or a selected API/subscription model proposes structured fields when local reading is incomplete.'),
 ('Our product','Exact matching, evidence review, route rules, receipt balances, duplicate control and Excel contract.')], 'No extraction score is treated as proof of business correctness.')
cards('The exact output leadership can inspect',[
 ('Header · 13 columns','One row per invoice. Supplier site, order, location, date and invoice totals.'),
 ('Tax · 3 columns','Transaction number, tax code and tax basis. Current MVP uses one tax treatment per invoice.'),
 ('Details · 6 columns','Transaction number, item, UPC, cost, quantity and unit tax code.')], 'Approved export revalidates and reserves receipts. Manual draft uses the same shape but is visibly DRAFT_UNVALIDATED and reserves nothing.')
cards('Three controlled routes to the workbook',[
 ('Approved export','Ready and reviewed invoices are revalidated. Receipt allocations and export receipt commit atomically.'),
 ('Combined review','Saved processed facts join into EXTRACTION_REVIEW_ONLY without approval, allocation or ledger change.'),
 ('Manual draft','Every target field is editable during OCR. Explicit acknowledgement; DRAFT_UNVALIDATED reserves nothing.')], 'All three use exact 13/3/6 sheets. Only approved export can reserve receipt quantity.')
cards('Cross-market supply needs an explicit route',[
 ('Supplier legal entity','A supplier name is not enough. Operational sites have their own identities and approved names.'),
 ('Buying company','Printed buyer name is evidence, not an internal code. PO scope must identify the approved buyer.'),
 ('Route and currency','UAE → Kuwait with AED billing remains a distinct route. Currency equality does not prove approval.')])
cards('Reference files are evidence, not assumptions',[
 ('Known source profile','412,139 source rows are live: 114,940 items and 297,199 PO/GRN rows. Search and filter checks passed.'),
 ('Show conflicts','Keep provenance and duplicates. Name, price, quantity and unit clues rank candidates; a person confirms.'),
 ('Approve separately','Lookup never approves matching. Supplier/site, receipt, prior invoicing, route and tax meaning still need owners.')], 'Completed target examples and lookup hits are evidence; neither establishes canonical master-data approval.')
cards('Cloud delivery is a controlled pilot',[
 ('Firebase frontend','Responsive portal, approved account sign-in and authenticated document downloads.'),
 ('GCP application','Cloud Run readers/API and managed Vertex through the service identity. Cloud SQL is 1 vCPU / 3.75 GB; estimated compute about USD 49/month, plus storage.'),
 ('Release status','Revision 00009-vxw is live. Reader, lookup and control checks passed; the allowlist is owner-only.')], 'The 412,139-row evidence catalog is live; catalog evidence is not canonical approval.')
cards('Owner subscription paths keep a hard boundary',[
 ('Saved vs verified','An encrypted saved key is not verified access. Model retrieval or an explicit probe verifies the connection.'),
 ('ChatGPT move','Local OAuth uses the application registration. Export deletes local credentials and clears the active account.'),
 ('Model selection','Offer a connected model only when none is selected; never replace a still-connected user choice.')], 'No automated result proves subscription eligibility, model access or live inference quality.')
cards('Review and deletion keep meanings explicit',[
 ('Buyer identity','Printed buyer name stays separate from the approved internal buyer/company code.'),
 ('Reader trace','Scanned invoice2data names Paddle input. Failed and text-only attempts are separate; coverage is not accuracy.'),
 ('Permanent deletion','Confirmed deletion removes live evidence; minimal ledger tombstones remain, and provider retention still applies.')])
cards('Release checks are green; deployment remains bounded',[
 ('Repository gate','214 backend tests pass with PostgreSQL included and no skips. Ruff, both Chromium workflows and all four GitHub checks pass.'),
 ('Reader image','All three packaged reader warm-ups recovered two synthetic lines.'),
 ('Hosted boundary','Revision 00009-vxw has 100% traffic. The allowlist is owner-only; public and direct-service APIs returned 401.')])
cards('Measured AI-off reader evidence',[
 ('Native tables','15/15 PDFs, 570/570 rows and 3,990/3,990 tested line facts exact in one family. Hosted: 193 rows and 1,351/1,351 line facts in 10.21 s.'),
 ('PaddleOCR','Adaptive two-page path: 90/99 to 96/99 (97.0%), 23/23 rows locally; revision 00008 returned the same hosted result in 242.98 s.'),
 ('Docling','Native examples and three small scans yielded 6/6 rows. A challenging scan gave 10/23 poor matches; ten pages exceeded 528 s.')], 'These bounded cohorts do not prove the expected 95% across suppliers, scans, languages and lengths.')
cards('Managed Vertex: useful evidence, narrow cohort',[
 ('Measured result','All source-tested fields were exact across 15 documents and 570 lines from one layout family.'),
 ('Measured usage','36,450 input and 91,892 output tokens including reasoning across 15 successful calls.'),
 ('Calculated cost','USD 0.3719325 at the published time-bounded rate; USD 24.7955 per 1,000 at the same mix.')], 'Not an actual bill. Excludes retries, hosting, database, storage and network. Broad accuracy remains unproven.')
cards('Define 90% before measuring it',[
 ('Header exact','Target ≥90% over every frozen gold header cell, including expected nulls. Failures and timeouts score zero.'),
 ('Line exact','Target ≥90% exact cells after adjudicated line alignment; also report missing, extra and misaligned lines.'),
 ('Zero correction','Target ≥90% of all attempted invoices needing no extraction edit, plus zero false-ready approvals.')], 'Stratify by supplier, native/scan, language, quality and length. Completeness, schema and arithmetic are diagnostics, not accuracy.')
cards('A staged build plan with clear exits',[
 ('Data contract\n1–2 weeks','Approve mappings and missing values. Accept the target workbook with the receiving team.'),
 ('Supplier evaluation\n2–3 weeks','Evaluate at least 100 representative invoices and the difficult exception cases.'),
 ('Production pilot\n3–5 weeks','Add role separation, durable workers, backup restoration, monitoring and controlled UAT.')], 'Planning ranges assume available AP/data owners and a deliberately bounded initial invoice contract.')
cards('Extend the platform after invoice proof',[
 ('New store or market','Search approved coverage; create supplier discovery and qualification cases only for gaps.'),
 ('Supplier item intake','Normalize item identity, codes, packaging and units. Route ambiguity to a data steward.'),
 ('Digital channels','Maintain versioned content, images, translations and channel rules separately from purchasing identity.')])
cards('The next business decisions',[
 ('Name the owners','AP process owner, reference steward, receiving owner and finance/tax approver.'),
 ('Accept the evidence','Confirm receipt status meaning, invoice baseline, route scope and the initial supplier test set.'),
 ('Authorize the rollout','Set approved users, provider choices and budget; sign off downstream workbook compatibility.')], 'Sixteen acceptance-driven implementation stories accompany the design and operator guide.')
prs.save(OUT/'Invoice_Studio_SLT.pptx')
downloads=ROOT/'docs'/'review-artifact'/'downloads';downloads.mkdir(exist_ok=True)
for name in ('Invoice_Studio_Design_and_Stories.docx','Invoice_Studio_SLT.pptx'):
    shutil.copy2(OUT/name,downloads/name)
print('Created Word design and',len(prs.slides),'slide SLT deck in docs/delivery and review-artifact/downloads')
