import hashlib
import io
import json
from copy import deepcopy

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.excel import HEADERS
from app.models import Invoice


MUTATION = {"X-Studio-Request": "1"}
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def draft_payload():
    return {
        "acknowledge_unvalidated": True,
        "Header": {
            "Transaction Number": 1,
            "Document": "0000380",
            "Supplier Site": "0091001",
            "Order No": "0070001",
            "Location": "000900001",
            "Location Type": "Warehouse (W)",
            "Document Date": "2026-10-04",
            "Total Cost Ex Tax": "20.0000",
            "Tax Amount": "1.0000",
            "Ref No. 1": "",
            "Ref No. 2": "",
            "Ref No. 3": "",
            "Comment": "Entered while OCR was running",
        },
        "Tax_Breakdown": {
            "Transaction Number": 1,
            "Tax Code": "VAT5",
            "Tax Basis": "20.0000",
        },
        "Details": [
            {
                "Transaction Number": 1,
                "Item": "000042",
                "UPC": "00012345678905",
                "Unit Cost": "5.0000",
                "Quantity": "4.0000",
                "Unit Tax Code": "VAT5",
            }
        ],
    }


@pytest.fixture
def local_client(tmp_path, monkeypatch):
    monkeypatch.setenv("INV_STUDIO_DATA", str(tmp_path / "module-default"))
    from app import main

    app = main.create_app(tmp_path / "data")
    with TestClient(app) as client:
        yield client, app
    app.state.pool.shutdown(wait=True)


@pytest.fixture
def cloud_client(tmp_path, monkeypatch):
    monkeypatch.setenv("INV_STUDIO_DATA", str(tmp_path / "module-default"))
    from app import cloud_store, main
    from app.store import Store

    class FakeIdentity:
        cloud = True

        def config(self):
            return {"cloud": True, "firebase": {"providers": ["password"]}}

        def verify(self, value):
            if value == "Bearer valid":
                return {"email": "owner@example.test", "uid": "owner"}
            raise ValueError("Sign in")

    class FakeCloudStore(Store):
        def sync_templates(self):
            pass

    monkeypatch.setattr(main, "CloudIdentity", FakeIdentity)
    monkeypatch.setattr(cloud_store, "PostgresStore", FakeCloudStore)
    app = main.create_app(tmp_path / "cloud-data")
    with TestClient(app) as client:
        yield client, app
    app.state.pool.shutdown(wait=True)


def seed_processing_job(app):
    invoice = Invoice(
        number="OCR-ORIGINAL",
        site="existing-site",
        po="existing-order",
        location="existing-location",
        date="2025-01-02",
        taxCode="OLD-TAX",
        net="999.00",
        tax="49.95",
    ).model_dump(mode="json")
    job = {
        "id": "processing-job",
        "filename": "slow-scan.pdf",
        "status": "processing",
        "revision": 7,
        "reviewed": False,
        "invoice": invoice,
        "validation": {"ready": False, "issues": [], "matches": []},
        "trace": [],
        "provenance": [],
    }
    app.state.store.job(job["id"], job)
    app.state.store.set("references", {"version": "approved-v1", "sites": [{"id": "keep"}]})
    return job


def test_manual_draft_downloads_during_processing_without_mutating_job_or_ledger(local_client):
    client, app = local_client
    original = seed_processing_job(app)
    before_job = deepcopy(app.state.store.job(original["id"]))
    before_references = deepcopy(app.state.store.get("references"))
    with app.state.store.connection() as connection:
        before_ledger = app.state.store.ledger(connection)
        before_exports = connection.execute("SELECT COUNT(*) FROM exports").fetchone()[0]

    response = client.post(
        f"/api/jobs/{original['id']}/draft", headers=MUTATION, json=draft_payload()
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MIME
    assert response.headers["content-disposition"] == (
        'attachment; filename="Invoice_0000380_DRAFT_UNVALIDATED.xlsx"'
    )
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"

    workbook = load_workbook(io.BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ["Header", "Tax_Breakdown", "Details"]
    assert [cell.value for cell in workbook["Header"][1]] == HEADERS["Header"]
    assert [cell.value for cell in workbook["Tax_Breakdown"][1]] == HEADERS["Tax_Breakdown"]
    assert [cell.value for cell in workbook["Details"][1]] == HEADERS["Details"]
    assert workbook["Header"].max_column == 13
    assert workbook["Tax_Breakdown"].max_column == 3
    assert workbook["Details"].max_column == 6
    assert workbook["Header"]["B2"].value == "0000380"
    assert workbook["Header"]["B2"].data_type == "s"
    assert workbook["Details"]["C2"].value == "00012345678905"
    assert workbook["Details"]["C2"].data_type == "s"
    assert workbook["Header"]["A2"].value == 1
    assert workbook["Tax_Breakdown"]["A2"].value == 1
    assert workbook["Details"]["A2"].value == 1
    assert not any(
        cell.data_type == "f"
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )

    after_job = app.state.store.job(original["id"])
    assert after_job == before_job
    assert after_job["status"] == "processing"
    assert after_job["revision"] == 7
    assert after_job["invoice"]["number"] == "OCR-ORIGINAL"
    assert app.state.store.get("references") == before_references
    with app.state.store.connection() as connection:
        assert app.state.store.ledger(connection) == before_ledger
        assert connection.execute("SELECT COUNT(*) FROM exports").fetchone()[0] == before_exports

    audit = client.get(f"/api/jobs/{original['id']}/audit")
    assert audit.status_code == 200
    event = audit.json()[-1]
    assert event["event"] == "manual_draft_downloaded"
    assert event["payload"] == {
        "kind": "manual_draft",
        "draft": True,
        "reference_validated": False,
        "receipt_reserved": False,
        "transaction_number": 1,
        "detail_rows": 1,
        "filename": "Invoice_0000380_DRAFT_UNVALIDATED.xlsx",
        "job_id": "processing-job",
        "workbook_sha256": hashlib.sha256(response.content).hexdigest(),
        "actor": "local-operator",
    }


@pytest.mark.parametrize("acknowledgement", [None, False])
def test_manual_draft_requires_explicit_unvalidated_acknowledgement(
    local_client, acknowledgement
):
    client, app = local_client
    job = seed_processing_job(app)
    body = draft_payload()
    if acknowledgement is None:
        body.pop("acknowledge_unvalidated")
    else:
        body["acknowledge_unvalidated"] = acknowledgement

    response = client.post(f"/api/jobs/{job['id']}/draft", headers=MUTATION, json=body)

    assert response.status_code == 422
    assert "acknowledge_unvalidated" in response.json()["detail"]
    with app.state.store.connection() as connection:
        assert connection.execute("SELECT COUNT(*) FROM audit").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM exports").fetchone()[0] == 0
    assert app.state.store.job(job["id"]) == job


def test_claude_subscription_import_and_delete_never_expose_plaintext(local_client):
    client, app = local_client
    secret = "claude-setup-token-" + "x" * 40

    imported = client.post(
        "/api/subscriptions/claude/import",
        headers=MUTATION,
        json={"setup_token": secret},
    )
    assert imported.status_code == 200, imported.text
    assert imported.json() == {"connected": True, "provider": "claude_local"}
    assert secret not in imported.text
    assert app.state.store.secret("claude_subscription")["token"] == secret

    state = client.get("/api/state")
    assert state.status_code == 200
    assert secret not in state.text
    assert secret.encode() not in app.state.store.path.read_bytes()
    assert isinstance(state.json()["connections"]["claude_local"], bool)

    deleted = client.delete("/api/subscriptions/claude", headers=MUTATION)
    assert deleted.status_code == 200
    assert deleted.json() == {"connected": False}
    assert app.state.store.secret("claude_subscription") is None
    assert secret not in deleted.text
    with app.state.store.connection() as connection:
        events = [row[0] for row in connection.execute("SELECT event FROM audit ORDER BY id")]
    assert events == ["subscription_connected", "subscription_disconnected"]


def test_chatgpt_credential_transfer_enforces_local_and_cloud_sides(
    local_client, cloud_client
):
    local, _ = local_client
    cloud, _ = cloud_client
    marker = "credential-marker-that-must-not-be-stored"

    local_import = local.post(
        "/api/chatgpt/import", headers=MUTATION, json={"bundle": {"marker": marker}}
    )
    assert local_import.status_code == 403
    assert marker not in local.get("/api/state").text

    cloud_export = cloud.post(
        "/api/chatgpt/export",
        headers={**MUTATION, "Authorization": "Bearer valid"},
        json={"id": "account-that-does-not-need-to-exist"},
    )
    assert cloud_export.status_code == 403
    assert "local Invoice Studio" in cloud_export.json()["detail"]


@pytest.mark.parametrize(
    ("method", "path", "body"),
    [
        ("POST", "/api/jobs/unknown/draft", draft_payload()),
        ("POST", "/api/subscriptions/claude/import", {"setup_token": "x" * 40}),
        ("DELETE", "/api/subscriptions/claude", None),
        ("POST", "/api/chatgpt/export", {"id": "account"}),
        ("POST", "/api/chatgpt/import", {"bundle": {}}),
        ("GET", "/api/reference-lookup/summary", None),
        ("GET", "/api/reference-lookup/search?kind=item&q=widget", None),
    ],
)
def test_cloud_authentication_protects_every_extension_route(
    cloud_client, method, path, body
):
    client, _ = cloud_client
    request = {"headers": MUTATION}
    if body is not None:
        request["json"] = body
    response = client.request(method, path, **request)
    assert response.status_code == 401
    assert response.json() == {"detail": "Sign in"}


def test_reference_lookup_routes_are_read_only_source_evidence(local_client, tmp_path):
    import gzip
    from app.reference_lookup import ReferenceLookup
    client, app = local_client
    source_hash = "a" * 64
    rows = [{"kind": "item", "source_hash": source_hash, "source_sheet": "Items",
             "source_row": index, "keys": [code],
             "data": {"sku": code, "description": description}, "flags": []}
            for index, code, description in [(1, "000042", "Widget small"), (2, "000043", "Widget large")]]
    archive = tmp_path / "catalog.ndjson.gz"
    with gzip.open(archive, "wt") as output:
        for row in rows: output.write(json.dumps(row) + "\n")
    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps({"version": 1, "expected_counts": {"item": 2, "po": 0, "total": 2},
                                    "source_hashes": [source_hash]}))
    ReferenceLookup(app.state.store).import_archive(archive, manifest)
    summary = client.get("/api/reference-lookup/summary")
    assert summary.status_code == 200
    assert summary.json()["approved_for_matching"] is False
    assert summary.json()["counts"] == {"item": 2, "po": 0, "total": 2}
    page_one = client.get("/api/reference-lookup/search", params={"kind": "item", "q": "Widget", "limit": 1})
    assert page_one.status_code == 200, page_one.text
    first = page_one.json()
    assert first["approved_for_matching"] is False
    assert len(first["records"]) == 1
    assert first["next_cursor"]
    page_two = client.get("/api/reference-lookup/search", params={
        "kind": "item", "q": "widget", "limit": 1, "cursor": first["next_cursor"]})
    assert page_two.status_code == 200, page_two.text
    second = page_two.json()
    assert second["next_cursor"] is None
    result_rows = first["records"] + second["records"]
    assert {row["data"]["sku"] for row in result_rows} == {"000042", "000043"}
    assert all(row["source_hash"] == source_hash and row["requires_confirmation"] for row in result_rows)
    assert app.state.store.get("references") is None
    with app.state.store.connection() as connection:
        assert connection.execute("SELECT COUNT(*) FROM audit").fetchone()[0] == 0
