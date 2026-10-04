import io
from copy import deepcopy

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.drafts import (
    ManualDraftRequest,
    build_draft_workbook,
    draft_filename,
    draft_public_metadata,
    validate_draft,
)
from app.excel import HEADERS


def valid_payload():
    return {
        "acknowledge_unvalidated": True,
        "Header": {
            "Transaction Number": 1,
            "Document": "0000380",
            "Supplier Site": "0091001",
            "Order No": "0070001",
            "Location": "000900001",
            "Location Type": "Warehouse (W)",
            "Document Date": "2026-10-04",
            "Total Cost Ex Tax": "20.0000",
            "Tax Amount": "1.0000",
            "Ref No. 1": "00001",
            "Ref No. 2": "reference 2",
            "Ref No. 3": "reference 3",
            "Comment": "Manual draft requested by operator",
        },
        "Tax_Breakdown": {
            "Transaction Number": 1,
            "Tax Code": "VAT5",
            "Tax Basis": "20.0000",
        },
        "Details": [
            {
                "Transaction Number": 1,
                "Item": "000042",
                "UPC": "00012345678905",
                "Unit Cost": "5.0000",
                "Quantity": "2.0000",
                "Unit Tax Code": "VAT5",
            },
            {
                "Transaction Number": 1,
                "Item": "000043",
                "UPC": "00012345678912",
                "Unit Cost": "2.5000",
                "Quantity": "4.0000",
                "Unit Tax Code": "VAT5",
            },
        ],
    }


def test_manual_draft_builds_exact_target_workbook_with_visible_disclosure():
    payload = valid_payload()
    model = validate_draft(payload)
    assert isinstance(model, ManualDraftRequest)

    workbook = load_workbook(io.BytesIO(build_draft_workbook(model)), data_only=False)
    assert workbook.sheetnames == ["Header", "Tax_Breakdown", "Details"]
    assert workbook.properties.title == "Invoice Studio DRAFT_UNVALIDATED"
    assert "reference data was not validated" in workbook.properties.description
    assert "no receipt number was reserved" in workbook.properties.description

    expected_rows = {"Header": 2, "Tax_Breakdown": 2, "Details": 3}
    for name, columns in HEADERS.items():
        sheet = workbook[name]
        assert [cell.value for cell in sheet[1]] == columns
        assert sheet.max_column == len(columns)
        assert sheet.max_row == expected_rows[name]
        assert "DRAFT_UNVALIDATED" in sheet["A1"].comment.text
        assert "no receipt number was reserved" in sheet["A1"].comment.text

    header = workbook["Header"]
    assert [header.cell(2, column).value for column in range(1, 14)] == [
        1,
        "0000380",
        "0091001",
        "0070001",
        "000900001",
        "Warehouse (W)",
        header["G2"].value,
        20,
        1,
        "00001",
        "reference 2",
        "reference 3",
        "Manual draft requested by operator",
    ]
    assert header["G2"].value.date() == model.header.document_date
    for coordinate in ("B2", "C2", "D2", "E2", "J2"):
        assert header[coordinate].data_type == "s"
    assert workbook["Tax_Breakdown"]["A2"].value == 1
    assert [workbook["Details"].cell(row, 1).value for row in (2, 3)] == [1, 1]
    assert workbook["Details"]["B2"].value == "000042"
    assert workbook["Details"]["C2"].value == "00012345678905"
    assert workbook["Details"]["C2"].number_format == "@"

    formulas = [
        cell.coordinate
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    assert formulas == []


def test_manual_draft_filename_and_public_audit_metadata_are_explicit_and_safe():
    payload = valid_payload()
    payload["Header"]["Document"] = "INV/0000380 ? final"
    name = draft_filename(payload)
    assert name == "Invoice_INV_0000380_final_DRAFT_UNVALIDATED.xlsx"
    assert "/" not in name

    metadata = draft_public_metadata(payload, job_id="job-1")
    assert metadata == {
        "kind": "manual_draft",
        "draft": True,
        "reference_validated": False,
        "receipt_reserved": False,
        "transaction_number": 1,
        "detail_rows": 2,
        "filename": name,
        "job_id": "job-1",
    }
    assert "0000380" not in repr({key: value for key, value in metadata.items() if key != "filename"})


@pytest.mark.parametrize(
    ("mutate", "message"),
    [
        (lambda body: body.pop("Tax_Breakdown"), "Tax_Breakdown"),
        (lambda body: body.update({"unexpected": True}), "Extra inputs"),
        (lambda body: body.update({"acknowledge_unvalidated": False}), "Input should be True"),
        (lambda body: body["Header"].update({"Transaction Number": 2}), "less than or equal to 1"),
        (lambda body: body["Details"][0].update({"Transaction Number": "1"}), "valid integer"),
        (lambda body: body["Header"].update({"Document": "=HYPERLINK('x')"}), "formula"),
        (lambda body: body["Details"][0].update({"Item": "+1+1"}), "formula"),
        (lambda body: body["Header"].update({"Comment": "@SUM(A:A)"}), "formula"),
        (lambda body: body["Header"].update({"Document Date": "2026-02-30"}), "calendar date"),
        (lambda body: body["Header"].update({"Total Cost Ex Tax": "2e1"}), "plain"),
        (lambda body: body["Header"].update({"Total Cost Ex Tax": 20.0}), "exact decimal"),
        (lambda body: body["Details"][0].update({"Quantity": "0"}), "greater than 0"),
        (lambda body: body["Details"][0].update({"Unit Cost": "-5"}), "plain"),
        (lambda body: body["Header"].update({"Supplier Site": ""}), "must not be blank"),
        (lambda body: body["Tax_Breakdown"].update({"Tax Basis": "19.9999"}), "Tax Basis"),
        (lambda body: body["Details"][0].update({"Unit Cost": "4.9999"}), "detail Unit Cost"),
    ],
)
def test_manual_draft_rejects_incomplete_malformed_or_incoherent_input(mutate, message):
    payload = deepcopy(valid_payload())
    mutate(payload)
    with pytest.raises(ValidationError, match=message):
        validate_draft(payload)


def test_manual_draft_requires_every_target_column_key():
    payload = valid_payload()
    payload["Details"][0].pop("UPC")
    with pytest.raises(ValidationError, match="UPC"):
        build_draft_workbook(payload)
