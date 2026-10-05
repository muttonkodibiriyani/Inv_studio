"""Review fields, banner and download from ONE fine-rules result with evidence. Synthetic data only."""

import io
from datetime import datetime
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import matching
from app.excel import HEADERS, rules_workbook
from app.matching import printed_evidence, rules_validation, rules_view
from tests.test_fine_rules_api import CONFIG, INVOICE, imported

H = {"x-studio-request": "1"}
PRINTED = "Total before tax 70.00\nVAT 0.00"


def lineage(target, value, source, reference="", line=None):
    return {"target": target, "line": line, "original": value, "value": value, "rule": "R-T", "source": source,
            "reference": reference, "confidence": "Exact"}


def approved_result(**extra):
    header = {"Document": "INV-1", "Supplier Site": "22001", "Order No": "13000001", "Location": "38091",
              "Location Type": "Store (S)", "Document Date": "2026-01-15", "Currency": "KWD", "Market": "Kuwait",
              "Net Amount": "30", "Tax Amount": "0", "Validation Status": "Approved"}
    lines = [{"Item": "345000001", "UPC": "0012345678905", "Unit Cost": "10", "Quantity": "3",
              "Unit Tax Code": "VAT0", "Validation Status": "Matched", "Source Row": "line 1"}]
    trace = [lineage("Document", "INV-1", "Invoice Document"),
             lineage("Supplier Site", "22001", "Item Master SUPPLIER", "Item Master!2"),
             lineage("Order No", "13000001", "POGRN RMS_ORDER_NO", "POGRN!2"),
             lineage("Location", "38091", "Accepted POGRN LOCATION", "POGRN!2"),
             lineage("Location Type", "Store (S)", "location master", "38091|v1"),
             lineage("Document Date", "2026-01-15", "Invoice Document Date"),
             lineage("Currency", "KWD", "Supplier sites", "22001|v1"),
             lineage("Market", "Kuwait", "Location-market mapping", "38091|v1"),
             lineage("Net Amount", "30", "Invoice totals", "page 1"),
             lineage("Tax Amount", "0", "Invoice totals", "page 1"),
             lineage("Item", "345000001", "Item Master ITEM_PARENT", "Item Master!2", 1),
             lineage("UPC", "0012345678905", "Invoice barcode", "", 1),
             lineage("Unit Cost", "10", "Invoice line", "", 1), lineage("Quantity", "3", "Invoice line", "", 1),
             lineage("Unit Tax Code", "VAT0", "Owner VAT code table", "KWT|v1", 1)]
    return {"status": "Approved", "header": header, "lines": lines, "lineage": trace, "exceptions": [], **extra}


def test_every_shown_value_carries_its_evidence():
    view = rules_view(approved_result())
    site = view["fields"]["site"]
    assert site["value"] == "22001" and site["evidence"][0]["kind"] == "sheet"
    assert site["evidence"][0]["reference"] == "Item Master!2"
    assert view["fields"]["market"]["evidence"][0]["kind"] == "table"
    assert view["fields"]["number"]["evidence"][0]["kind"] == "printed"
    assert view["fields"]["taxCode"]["value"] == "VAT0" and view["fields"]["taxCode"]["evidence"]
    for f in view["fields"].values():
        assert (f["value"] is None) == f["flagged"] and (f["value"] is None or f["evidence"])


def test_value_without_evidence_stays_empty_and_is_flagged_never_guessed():
    result = approved_result()
    result["lineage"] = [x for x in result["lineage"] if x["target"] != "Location"]
    view = rules_view(result)
    assert view["fields"]["location"] == {"label": "Delivery location", "target": "Location", "value": None,
                                          "evidence": [], "flagged": True, "reason": "No evidence"}
    # Not in the owner's target and in no owner sheet: not offered at all, so never filled.
    assert not {"seller", "buyer", "origin"} & set(view["fields"])
    blocking = {i["message"] for i in view["issues"] if i["blocking"]}
    assert "Delivery location has no evidence and is left empty" in blocking
    assert rules_validation(view, reviewed=True)["ready"] is False


def test_printed_amount_evidence_only_when_found_on_the_document():
    result = approved_result()
    result["lineage"] = [x for x in result["lineage"] if x["target"] not in ("Net Amount", "Tax Amount")]
    boxes = [{"page": 2, "text": "Net 30.00", "box": [1, 2, 3, 4]}, {"page": 2, "text": "VAT 0.00", "box": [1, 9, 3, 4]}]
    view = rules_view(result, "", boxes)
    assert view["fields"]["net"]["evidence"][0]["reference"] == "page 2 box [1, 2, 3, 4]"
    assert view["fields"]["tax"]["value"] == "0"
    missing = rules_view(result, "nothing printed here", [])
    assert missing["fields"]["net"]["value"] is None and missing["fields"]["net"]["flagged"]
    # A longer number is not a match for the amount.
    assert printed_evidence("30", "Total 130.00") is None


def test_item_line_rate_below_95_goes_to_owner_review():
    result = approved_result()
    result["lines"] = result["lines"] * 20
    result["lineage"] += [lineage(c, v, "s", "Item Master!2" if c == "Item" else "", n) for n in range(2, 21)
                          for c, v in (("Item", "345000001"), ("Unit Cost", "10"), ("Quantity", "3"),
                                       ("Unit Tax Code", "VAT0"))]
    assert rules_view(result)["item_lines"]["owner_review"] is False
    result["lines"] = [*result["lines"][:19], {**result["lines"][0], "Item": ""}]
    view = rules_view(result)
    assert view["item_lines"]["resolved"] == 19 and view["item_lines"]["rate"] == "0.9500"
    result["lines"] = [*result["lines"][:18], *[{**result["lines"][0], "Item": ""}] * 2]
    view = rules_view(result)
    assert view["item_lines"]["owner_review"] is True
    assert any(i["code"] == "Owner Review" and i["blocking"] for i in view["issues"])
    # RULES' own numbers win when present.
    view = rules_view(approved_result(item_resolution={"resolved": 1, "total": 2, "rate": "0.5", "below": True,
                                                       "definition": "one ITEM_PARENT agreeing with POGRN"}))
    assert view["item_lines"]["resolved"] == 1 and view["item_lines"]["owner_review"] is True
    assert view["item_lines"]["definition"] == "one ITEM_PARENT agreeing with POGRN"


def test_banner_uses_rules_exceptions_verbatim():
    result = approved_result(status="Review", exceptions=[{
        "Exception Type": "Missing/Ambiguous PO", "Description": "d", "Owner": "Buyer", "Line No.": "",
        "Rule ID": "ALG-008", "Candidates / Evidence": "", "blocking": True}])
    issues = rules_validation(rules_view(result), reviewed=True)["issues"]
    assert issues[0] == {"code": "Missing/Ambiguous PO", "message": "d", "owner": "Buyer", "line": None,
                         "rule": "ALG-008", "evidence": "", "blocking": True, "check": "", "type": ""}


def test_download_matches_the_owner_target_format():
    content, evidence = rules_workbook([rules_view(approved_result())])
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Header", "Tax_Breakdown", "Details"]
    header = wb["Header"]
    assert [c.value for c in header[1]] == HEADERS["Header"] and len(HEADERS["Header"]) == 13
    assert [header.cell(2, c).value for c in range(1, 7)] == [1, "INV-1", 22001, 13000001, 38091, "Store (S)"]
    assert header["G2"].value == datetime(2026, 1, 15) and header["G2"].number_format == "m/d/yyyy"
    assert header["H2"].value == 30 and header["I2"].value == 0
    assert all(header.cell(2, c).value is None for c in range(10, 14))
    tax = wb["Tax_Breakdown"]
    assert [c.value for c in tax[2]] == [1, "VAT0", 30] and tax.max_row == 2
    details = wb["Details"]
    assert [c.value for c in details[2]] == [1, 345000001, "0012345678905", 10, 3, "VAT0"]
    assert details["D2"].number_format == "General"
    assert evidence[1]["Header!C2"][0]["reference"] == "Item Master!2"
    assert evidence[1]["Details!B2"][0]["reference"] == "Item Master!2"
    no_upc, _ = rules_workbook([rules_view(approved_result())], upc="empty")
    assert load_workbook(io.BytesIO(no_upc))["Details"]["C2"].value is None


def test_download_refuses_unapproved_or_unevidenced_values():
    with pytest.raises(ValueError, match="approved"):
        rules_workbook([rules_view(approved_result(status="Review"))])
    result = approved_result()
    result["lineage"] = [x for x in result["lineage"] if x["target"] != "Order No"]
    with pytest.raises(ValueError, match="Purchase order has no evidenced value"):
        rules_workbook([rules_view(result)])


def job(text=PRINTED):
    return {"id": "job-1", "status": "review", "reviewed": False, "revision": 1, "filename": "synthetic.pdf",
            "invoice": INVOICE.model_dump(mode="json"), "text": text, "boxes": [], "provenance": []}


@pytest.fixture
def production(tmp_path, monkeypatch):
    monkeypatch.setenv("INV_STUDIO_DATA", str(tmp_path / "module-default"))
    monkeypatch.delenv("INV_STUDIO_DEMO_REFERENCES", raising=False)
    from app.main import create_app

    def never(*_, **__):
        raise AssertionError("demo validation references used in production")
    monkeypatch.setattr("app.main.enrich", never)
    monkeypatch.setattr("app.main.validate", never)
    imported(tmp_path, monkeypatch)
    app = create_app(tmp_path / "data")
    with TestClient(app) as client:
        assert client.post("/api/fine-rules/config", json=CONFIG, headers=H).status_code == 200
        yield app, client


def test_review_banner_and_download_come_from_one_rules_result(production):
    app, client = production
    app.state.store.job("job-1", job())
    shown = client.get("/api/jobs/job-1").json()
    rules = shown["rules"]
    assert rules["status"] == "Approved" and rules["revision"] == shown["revision"]
    assert rules["fields"]["site"]["value"] == "22001"
    assert rules["fields"]["site"]["evidence"][0]["reference"].startswith("Item Master!2")
    assert rules["fields"]["net"]["evidence"][0]["kind"] == "printed"
    assert rules["item_lines"] == {**rules["item_lines"], "resolved": 2, "total": 2, "owner_review": False}
    assert shown["validation"]["source"] == "fine_rules"
    assert [i["code"] for i in shown["validation"]["issues"] if i["blocking"]] == ["REVIEW"]
    # A rules 'Approved' alone never makes the job approved: export stays held until a person confirms review.
    assert shown["status"] == "review" and shown["validation"]["ready"] is False
    held = client.post("/api/jobs/job-1/export", headers=H, json={"revision": shown["revision"]})
    assert held.status_code == 409
    # The invoice the reviewer edits is not filled from derived values.
    assert shown["invoice"]["site"] is None
    reviewed = client.post("/api/jobs/job-1/review", headers=H,
                           json={"invoice": shown["invoice"], "revision": shown["revision"], "confirm": True}).json()
    assert reviewed["validation"]["ready"] is True and reviewed["rules"]["revision"] == reviewed["revision"]
    applied = [e for e in client.get("/api/jobs/job-1/audit").json() if e["event"] == "exported"]
    assert applied == []
    exported = client.post("/api/jobs/job-1/export", headers=H, json={"revision": reviewed["revision"]}).json()
    wb = load_workbook(io.BytesIO(client.get(exported["url"]).content))
    assert wb["Header"]["C2"].value == int(reviewed["rules"]["fields"]["site"]["value"])
    assert [wb["Details"].cell(r, 2).value for r in (2, 3)] == [345000001, 345000002]
    evidence = client.get(f"/api/exports/{exported['id']}/evidence").json()["cells"]
    assert evidence["Header!C2"][0]["reference"].startswith("Item Master!2")
    audit = [e for e in client.get("/api/jobs/job-1/audit").json() if e["event"] == "exported"][0]["payload"]
    assert "evidence" not in audit and audit["evidence_cells"] == len(evidence)


def test_empty_not_guessed_blocks_export(production):
    app, client = production
    app.state.store.job("job-1", job(text="no totals printed"))
    shown = client.get("/api/jobs/job-1").json()
    assert shown["rules"]["fields"]["net"] == {**shown["rules"]["fields"]["net"], "value": None, "flagged": True}
    reviewed = client.post("/api/jobs/job-1/review", headers=H,
                           json={"invoice": shown["invoice"], "revision": shown["revision"], "confirm": True}).json()
    assert reviewed["validation"]["ready"] is False
    assert "Net total not found" not in str(reviewed["validation"]["issues"])
    assert any(i["rule"] == "EVIDENCE" and i["blocking"] for i in reviewed["validation"]["issues"])
    assert client.post("/api/jobs/job-1/export", headers=H, json={"revision": reviewed["revision"]}).status_code == 409


def test_rules_rerun_when_mapping_tables_change(production):
    app, client = production
    app.state.store.job("job-1", job())
    first = client.get("/api/jobs/job-1").json()["rules"]
    assert client.post("/api/fine-rules/config", json={}, headers=H).status_code == 200
    second = client.get("/api/jobs/job-1").json()["rules"]
    assert second["signature"] != first["signature"] and second["status"] != "Approved"


def test_demo_references_are_test_only(production, monkeypatch, tmp_path):
    _, client = production
    assert client.post("/api/references/demo", headers=H).status_code == 404
    assert client.get("/api/state").json()["references"] is None
    monkeypatch.setenv("INV_STUDIO_DEMO_REFERENCES", "1")
    monkeypatch.setenv("INV_STUDIO_CLOUD", "1")
    from app.main import create_app
    with pytest.raises(RuntimeError, match="test-only"):
        create_app(tmp_path / "cloud")


def test_legacy_demo_validation_is_still_available_to_tests():
    assert matching.validate and matching.enrich and Decimal("0.95") == matching.ITEM_THRESHOLD


def test_extraction_evidence_backs_printed_fields_and_disagreement_is_a_warning():
    result={"status":"Approved","header":{"Document":"SYN-1"},"lines":[],"lineage":[],"exceptions":[]}
    printed={"header":{"number":{"quote":"Invoice SYN-1","page":1,"box":[0.1,0.1,0.3,0.12],"source":"ai",
                                 "review":{"reason":"OCR read a different number","other_value":"SYN-7"}}},
             "lines":[{"qty":{"quote":"2","page":1,"source":"ocr"}}]}
    view=rules_view(result,printed=printed)
    number=view["fields"]["number"]
    assert number["value"]=="SYN-1" and not number["flagged"]
    assert number["evidence"][0]["kind"]=="printed" and number["evidence"][0]["reference"]=="page 1 box [0.1, 0.1, 0.3, 0.12]"
    warn=[i for i in view["issues"] if i["rule"]=="EVID-OCR"]
    assert len(warn)==1 and warn[0]["blocking"] is False and "SYN-1" not in warn[0]["message"]
    # Without extraction evidence the same value stays empty and flagged.
    assert rules_view(result)["fields"]["number"]["value"] is None


def test_owner_rule_and_printed_evidence_kinds_and_buyer_name():
    result={"status":"Approved","header":{"Document":"SYN-2","Buyer Name":"Synthetic Buyer"},"lines":[],"exceptions":[],
            "lineage":[{"target":"Buyer Name","line":None,"original":"","value":"Synthetic Buyer","rule":"BUYER-NAME",
                        "source":"Owner rule","reference":"owner rule BUYER-NAME (2026-10-05)","confidence":"Owner rule",
                        "evidence_kind":"owner_rule"},
                       {"target":"Document","line":None,"original":"Invoice SYN-2","value":"SYN-2","rule":"",
                        "source":"Invoice","reference":"page 1","confidence":"Exact","evidence_kind":"printed"}]}
    view=rules_view(result)
    buyer=view["fields"]["buyer_name"]
    assert buyer["value"]=="Synthetic Buyer" and not buyer["flagged"] and buyer["evidence"][0]["kind"]=="owner_rule"
    assert view["fields"]["number"]["evidence"][0]["kind"]=="printed"
    assert not any("Buyer" in i["message"] for i in view["issues"])
    result["header"]["Buyer Name"]=""
    assert "buyer_name" not in rules_view(result)["fields"]


def test_owner_review_is_not_doubled_when_rules_raises_it():
    result={"status":"Review","header":{},"lines":[],"lineage":[],
            "item_resolution":{"resolved":1,"total":2,"rate":"0.5","threshold":"0.95","below":True,"definition":"d"},
            "exceptions":[{"Exception Type":"Owner Validation","Description":"Item-line check 1/2 is below 95 %","Rule ID":"ITEM-LINE-95","blocking":True}]}
    view=rules_view(result)
    assert view["item_lines"]["owner_review"] is True
    assert [i["rule"] for i in view["issues"] if i["rule"] in ("ITEM-95","ITEM-LINE-95")]==["ITEM-LINE-95"]


def test_failure_status_check_and_po_candidates_reach_the_view():
    result={"status":"Review","header":{},"lines":[],"lineage":[],"po_candidates":3,
            "exceptions":[{"Exception Type":"Missing/Ambiguous PO","Engine Type":"Ambiguous PO","Check ID":"C-12",
                           "Description":"3 orders","Rule ID":"POG-001","blocking":True}]}
    view=rules_view(result)
    issue=next(i for i in view["issues"] if i["rule"]=="POG-001")
    assert (issue["code"],issue["type"],issue["check"])==("Missing/Ambiguous PO","Ambiguous PO","C-12")
    assert view["po_candidates"]==3
