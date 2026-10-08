"""Bulk downloads run in the background so no single request outlasts the 60 s Firebase Hosting proxy."""
import io
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.models import Invoice
from tests.test_extraction_draft import _rules_lines

HEADERS = {'X-Studio-Request': '1'}


def _app(tmp_path, monkeypatch, delay=0.0):
    monkeypatch.setenv('INV_STUDIO_DATA', str(tmp_path / 'module-default'))
    import app.main
    from tests.test_rules_wiring import approved_result

    def slow(entries, *a, **k):
        time.sleep(delay)
        return [approved_result() for _ in entries]
    monkeypatch.setattr(app.main, 'run_batch', slow)
    app_ = app.main.create_app(tmp_path)
    invoice = Invoice(number='INV-1', net='30', lines=[{'item_id': '00042', 'sku': 'VENDOR-1', 'qty': '3', 'price': '10'}]).model_dump(mode='json')
    for n in range(1, 4):
        app_.state.store.job(f'synthetic-{n}', {'id': f'synthetic-{n}', 'filename': f'synthetic-{n}.pdf', 'status': 'review', 'revision': 2,
                                                'reviewed': False, 'invoice': invoice, 'rules': {**_rules_lines(), 'revision': 1}})
    return app_


def _jobs(*ids):
    return {'jobs': [{'id': i, 'revision': 2} for i in ids], 'acknowledge_unvalidated': True}


def test_each_request_waits_less_than_the_hosting_proxy():
    import app.main
    assert app.main.BACKGROUND_WAIT < 60


def test_a_quick_bulk_download_comes_back_on_the_first_request(tmp_path, monkeypatch):
    with TestClient(_app(tmp_path, monkeypatch)) as client:
        response = client.post('/api/background/extraction-batch', headers=HEADERS, json=_jobs('synthetic-1', 'synthetic-2'))
        assert response.status_code == 200 and 'attachment' in response.headers['content-disposition']
        assert load_workbook(io.BytesIO(response.content))['Details']['B2'].value == 345000001


def test_a_slow_bulk_download_is_polled_until_the_workbook_is_ready(tmp_path, monkeypatch):
    import app.main
    monkeypatch.setattr(app.main, 'BACKGROUND_WAIT', 0.05)
    with TestClient(_app(tmp_path, monkeypatch, delay=0.4)) as client:
        response = client.post('/api/background/extraction-batch', headers=HEADERS, json=_jobs('synthetic-1', 'synthetic-2', 'synthetic-3'))
        assert response.status_code == 202
        polls = 0
        while response.status_code == 202:
            polls += 1
            task = response.json()['id']
            response = client.get(f"/api/background/tasks/{task}")
        assert polls >= 1 and response.status_code == 200
        book = load_workbook(io.BytesIO(response.content))
        assert book['Details']['B2'].value == 345000001 and book['Header'].max_row == 4
        # Collected once: a second poll of the same id is gone.
        assert client.get(f"/api/background/tasks/{task}").status_code == 404


def test_a_refusal_reaches_the_browser_with_its_status_and_reason(tmp_path, monkeypatch):
    import app.main
    monkeypatch.setattr(app.main, 'BACKGROUND_WAIT', 0.05)
    with TestClient(_app(tmp_path, monkeypatch, delay=0.2)) as client:
        response = client.post('/api/background/extraction-batch', headers=HEADERS,
                               json={'jobs': [{'id': 'synthetic-1', 'revision': 1}], 'acknowledge_unvalidated': True})
        while response.status_code == 202:
            response = client.get(f"/api/background/tasks/{response.json()['id']}")
        assert response.status_code == 409 and 'changed' in response.json()['detail']
        assert client.post('/api/background/not-a-download', headers=HEADERS, json={}).status_code == 404
        assert client.post('/api/background/extraction-batch', headers=HEADERS, json={'jobs': 'x'}).status_code in (400, 422)
