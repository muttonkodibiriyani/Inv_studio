"""Owner mapping tables (VAT codes, location master, supplier sites). Every table here is synthetic."""

import json
import os
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook

from app import fine_rules as fr
from app import fine_rules_tables as tables
from tests.test_fine_rules import invoice, run, types

D = Decimal


def workbook(path):
    book = Workbook()
    vat = book.active
    vat.title = "VAT CODES"
    vat.append(["ITEM", "VAT_REGION", "VAT_TYPE", "VAT_CODE", "VAT_REGION_NAME", "VAT_RATE", "ACTIVE_DATE"])
    since = datetime(2025, 10, 1)
    vat.append([1, 1, "C", "TKPV00", "TESTLAND Vat Region", 0, since])
    vat.append([1, 1, "R", "TKSV00", "TESTLAND Vat Region", 0, since])
    vat.append([1, 2, "C", "ZZPV05", "ZEDLAND Vat Region", 5, since])
    vat.append([1, 3, "C", "CXPV00", "CUSTOM Vat Region", 0, since])
    vat.append([1, 4, "C", "QQPV09", "QLAND Vat Region", "nine", since])  # row 6: bad rate
    locations = book.create_sheet("LOCATIONS")
    locations.append(["LOCATION", "WH/STORE", "ENTITY AND CURENCY", "COUNTRY"])
    locations.append([38091, "STORE", "RT1TKD", "TESTLAND"])
    locations.append([800901, "W/H", "RT1TKD", "TESTLAND"])
    locations.append([38092, "SHOP", "RZ1ZZD", "ZEDLAND"])  # row 4: unknown type
    locations.append([38091, "STORE", "RT1TKD", "TESTLAND"])  # row 5: repeat
    sites = book.create_sheet("SUPPLIER SITES")
    sites.append(["Supplier CODE", "Supplier Name", "Supplier Site", "Supplier Site Name", "Status Description",
                  "Currency"])
    sites.append([1, "Sample Supplier", 22001, "ABC001RT1TKD", "Active", "KWD"])
    sites.append([1, "Sample Supplier", 22002, "ABC001RZ1ZZD", "Inactive", "AED"])
    sites.append([2, "Other Sample", None, None, "Active", "USD"])
    sites.append([3, "Third Sample", 99013, "XYZ001RT1TKD", "Active", ""])  # row 5: no currency
    sites.append([1, "Sample Supplier", 22001, "ABC001RT1TKD", "Active", "KWD"])  # row 6: repeat
    book.save(path)
    return path


def owner_config(**changes):
    config = {
        # The synthetic items carry the confirmed RA1 (Kuwait) suffix; the owner sheet writes countries in capitals.
        "location_master": {"38091": {"type": fr.STORE, "market": "KUWAIT", "country": "KUWAIT"},
                            "800901": {"type": fr.WAREHOUSE, "market": "KUWAIT", "country": "KUWAIT"}},
        "supplier_sites": [{"supplier_site": "22001", "currency": "KWD", "status": "Active"}],
        "vat_codes": [{"code": "TKPV00", "region": "KUWAIT", "rate": "0", "active_from": "2025-10-01"},
                      {"code": "ZZPV05", "region": "ZEDLAND", "rate": "5", "active_from": "2025-10-01"}],
        "version": "owner-test",
    }
    config.update(changes)
    return config


def test_loader_reads_three_sheets_counts_and_flags_bad_rows(tmp_path):
    loaded, report = tables.load(workbook(tmp_path / "owner.xlsx"))
    assert [v["code"] for v in loaded["vat_codes"]] == ["TKPV00", "ZZPV05", "CXPV00"]
    assert loaded["location_master"]["800901"] == {"type": fr.WAREHOUSE, "market": "TESTLAND",
                                                   "country": "TESTLAND", "entity_currency": "RT1TKD"}
    assert loaded["supplier_sites"][1] == {"supplier_site": "22002", "currency": "AED", "status": "Inactive",
                                           "supplier_code": "1", "supplier_name": "Sample Supplier",
                                           "site_name": "ABC001RZ1ZZD"}
    assert {k: v for k, v in report.items() if k != "flagged"} == {
        "vat_codes_c_pv": 3, "vat_rows_excluded": 1, "locations": 2,
        "locations_by_type": {fr.STORE: 1, fr.WAREHOUSE: 1}, "supplier_sites": 2, "supplier_sites_active": 1,
        "supplier_sites_inactive": 1, "supplier_rows_without_site_id": 1, "site_currencies": 2}
    assert [(f["sheet"], f["row"]) for f in report["flagged"]] == [
        ("VAT CODES", 6), ("LOCATIONS", 4), ("LOCATIONS", 5), ("SUPPLIER SITES", 5), ("SUPPLIER SITES", 6)]
    fr.RulesConfig.from_dict(loaded)


def test_cli_writes_a_private_merged_config_and_prints_counts_only(tmp_path, capsys):
    current = tmp_path / "current.json"
    current.write_text(json.dumps({"value_tolerance": "0.5", "supplier_sites": [], "version": "old"}))
    out = tmp_path / "private.json"
    assert tables.main([str(workbook(tmp_path / "owner.xlsx")), "--out", str(out), "--merge", str(current),
                        "--version", "owner-1"]) == 0
    assert os.stat(out).st_mode & 0o777 == 0o600
    written = json.loads(out.read_text())
    assert written["value_tolerance"] == "0.5" and written["version"] == "owner-1"
    assert len(written["supplier_sites"]) == 2
    printed = capsys.readouterr().out
    assert json.loads(printed)["supplier_sites"] == 2
    assert "22001" not in printed and "Sample" not in printed and "TKPV00" not in printed


def test_owner_tables_give_market_currency_and_unit_tax_code():
    result = run(invoice(taxCode=None), config=owner_config())
    assert result["status"] == "Approved", result["exceptions"]
    header = result["header"]
    assert (header["Market"], header["Currency"], header["Location Type"]) == ("KUWAIT", "KWD", fr.STORE)
    assert {line["Unit Tax Code"] for line in result["lines"]} == {"TKPV00"}
    sources = {x["target"]: x["source"] for x in result["lineage"]}
    assert sources["Currency"] == "Supplier-site table"
    assert sources["Unit Tax Code"] == "Owner VAT code table (C/PV), receiving market"


def test_inactive_supplier_site_currency_is_left_for_review():
    sites = [{"supplier_site": "22001", "currency": "KWD", "status": "Inactive"}]
    result = run(invoice(taxCode=None), config=owner_config(supplier_sites=sites))
    assert result["status"] == "Review" and result["header"]["Currency"] == ""
    assert "Currency Mapping" in types(result, "R-012")


def test_site_missing_from_the_table_is_a_currency_exception():
    result = run(invoice(taxCode=None), config=owner_config(supplier_sites=[]))
    assert result["header"]["Currency"] == "" and "Currency Mapping" in types(result, "R-012")


def test_usd_site_still_needs_an_approved_usd_exception():
    sites = [{"supplier_site": "22001", "currency": "USD", "status": "Active"}]
    state = fr.Run(invoice(currency="USD"), fr.RulesConfig.from_dict(owner_config(supplier_sites=sites)), "")
    assert fr.resolve_currency(state, "22001", "KUWAIT", None) is None
    assert [e["Engine Type"] for e in state.exceptions] == ["USD Review"]


def test_tax_code_needs_a_region_row_an_active_date_and_the_printed_rate():
    no_region = run(invoice(taxCode=None), config=owner_config(vat_codes=[]))
    assert {line["Unit Tax Code"] for line in no_region["lines"]} == {""}
    assert "Tax Code" in types(no_region, "R-016")
    later = [{"code": "TKPV00", "region": "KUWAIT", "rate": "0", "active_from": "2026-02-01"}]
    assert "Tax Code" in types(run(invoice(taxCode=None), config=owner_config(vat_codes=later)), "R-016")
    taxed = run(invoice(taxCode=None, tax=D("3.50")), config=owner_config())
    assert taxed["status"] == "Review" and any("does not agree" in e["Description"] for e in taxed["exceptions"])
    five = [{"code": "TKPV05", "region": "KUWAIT", "rate": "5", "active_from": "2025-10-01"}]
    assert run(invoice(taxCode=None, tax=D("3.50")), config=owner_config(vat_codes=five))["status"] == "Approved"


def test_tax_code_region_is_the_receiving_location_country_only():
    zed = {"38091": {"type": fr.STORE, "market": "ZEDLAND", "country": "ZEDLAND"}}
    result = run(invoice(taxCode=None, tax=D("3.50")), config=owner_config(location_master=zed))
    assert {line["Unit Tax Code"] for line in result["lines"]} == {"ZZPV05"}


def test_printed_tax_code_is_kept_as_reviewed():
    result = run(invoice(taxCode="VAT0"), config=owner_config())
    assert {line["Unit Tax Code"] for line in result["lines"]} == {"VAT0"}


def test_location_master_type_overrides_prefix_and_market_ignores_case():
    master = {"38091": {"type": fr.WAREHOUSE, "market": "KUWAIT"}}
    result = run(invoice(taxCode=None), config=owner_config(location_master=master))
    assert "Location master type" in json.dumps(result["pogrn_validation"], default=str)
    result = run(invoice(taxCode=None), config=owner_config())
    assert "Market Mapping" not in types(result) and result["status"] == "Approved", result["exceptions"]
    zed = {"38091": {"type": fr.STORE, "market": "ZEDLAND"}}
    assert "Market Mapping" in types(run(invoice(taxCode=None), config=owner_config(location_master=zed)))


def test_config_rejects_conflicting_owner_rows():
    import pytest

    with pytest.raises(ValueError):
        fr.RulesConfig.from_dict(owner_config(supplier_sites=[{"supplier_site": "1", "currency": "KWD",
                                                               "status": "Active"}] * 2))
    with pytest.raises(ValueError):
        fr.RulesConfig.from_dict(owner_config(vat_codes=[{"code": "X", "region": "R", "rate": "-1",
                                                          "active_from": "2025-10-01"}]))


def test_config_audit_keeps_table_sizes_not_rows(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from app.main import create_app

    monkeypatch.setenv("INV_STUDIO_DATA", str(tmp_path / "module-default"))
    app = create_app(tmp_path / "data")
    with TestClient(app) as client:
        saved = client.post("/api/fine-rules/config", json=owner_config(), headers={"x-studio-request": "1"})
        assert saved.status_code == 200
        assert client.get("/api/fine-rules/config").json()["supplier_sites"][0]["supplier_site"] == "22001"
    with app.state.store.connection() as c:
        audits = [json.loads(r[0]) for r in c.execute("SELECT payload FROM audit WHERE event='fine_rules_config_changed'")]
    assert audits[-1]["supplier_sites"] == 1 and audits[-1]["location_master"] == 2
    assert "22001" not in json.dumps(audits)


def test_banner_names_neither_list_once_owner_tables_are_loaded():
    import re
    import shutil
    import subprocess
    from pathlib import Path

    import pytest

    if not shutil.which("node"):
        pytest.skip("node is not installed")
    script = (Path(__file__).parents[1] / "app/static/app.js").read_text()
    body = re.search(r"function fineRulesConfigGaps\(config\) \{.*?\n\}\n", script, re.S).group(0)
    probe = body + "console.log(JSON.stringify([fineRulesConfigGaps({}), fineRulesConfigGaps(%s)]));" % json.dumps(
        owner_config())
    empty, loaded = json.loads(subprocess.run(["node", "-e", probe], capture_output=True, text=True,
                                              check=True).stdout)
    assert any("V-007" in g for g in empty) and any("V-010" in g for g in empty)
    assert loaded == []
