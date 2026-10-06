import io
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from app.models import Invoice
from app.extraction_draft import extraction_workbook


def test_review_copy_preserves_printed_values_and_leaves_unknown_codes_blank():
    invoice=Invoice(number='=NOT_A_FORMULA',date='2026-01-02',net='10.01',tax='0.50',
        lines=[{'sku':'00077','gtin':'00012345678905','qty':'3','price':'3.33','description':'Test only'}])
    book=load_workbook(io.BytesIO(extraction_workbook(invoice,'synthetic.pdf',2)))
    assert book.sheetnames==['Header','Tax_Breakdown','Details']
    assert [sheet.max_column for sheet in book]==[13,3,6]
    assert book['Header']['B2'].value=='=NOT_A_FORMULA'
    assert book['Header']['B2'].data_type=='s'
    assert book['Header']['H2'].value==10.01
    assert book['Header']['C2'].value is None
    assert book['Details']['B2'].value is None
    assert '00077' in book['Details']['B2'].comment.text
    assert book['Details']['C2'].value=='00012345678905'
    assert book['Details']['D2'].value==3.33
    assert all(cell.data_type!='f' for sheet in book for row in sheet for cell in row)
    assert 'EXTRACTION_REVIEW_ONLY' in book['Header']['M2'].value


def test_confirmed_internal_item_is_written_without_substituting_supplier_sku():
    invoice=Invoice(number='SYN-REF',lines=[{'item_id':'00042','sku':'VENDOR-9','qty':'2','price':'5'}])
    book=load_workbook(io.BytesIO(extraction_workbook(invoice,'synthetic.pdf',3)))
    assert book['Details']['B2'].value=='00042'
    assert book['Details']['B2'].data_type=='s'
    assert 'VENDOR-9' in book['Details']['B2'].comment.text
    assert 'business validation is still required' in book['Details']['B2'].comment.text


def test_review_copy_download_requires_ack_and_never_approves_or_changes_job(tmp_path,monkeypatch):
    monkeypatch.setenv('INV_STUDIO_DATA',str(tmp_path/'module-default'))
    from app.main import create_app
    app=create_app(tmp_path)
    record={'id':'synthetic','filename':'synthetic.pdf','status':'review','revision':2,'reviewed':False,
            'invoice':Invoice(number='REVIEW-1',net='10',lines=[{'sku':'0001','qty':'2','price':'5'}]).model_dump(mode='json')}
    app.state.store.job('synthetic',record)
    with TestClient(app) as client:
        headers={'X-Studio-Request':'1'};url='/api/jobs/synthetic/extraction-draft'
        assert client.post(url,headers=headers,json={'revision':2}).status_code==400
        assert client.post(url,headers=headers,json={'revision':1,'acknowledge_unvalidated':True}).status_code==409
        response=client.post(url,headers=headers,json={'revision':2,'acknowledge_unvalidated':True})
        assert response.status_code==200
        assert 'EXTRACTION_REVIEW_ONLY' in response.headers['Content-Disposition']
        assert app.state.store.job('synthetic')==record
        with app.state.store.connection() as connection:assert app.state.store.ledger(connection)==[]


def test_batch_review_workbook_links_each_invoice_and_rejects_stale_or_unfinished(tmp_path,monkeypatch):
    monkeypatch.setenv('INV_STUDIO_DATA',str(tmp_path/'module-default'))
    from app.main import create_app
    app=create_app(tmp_path)
    originals=[]
    for index,count in enumerate((1,2),1):
        record={'id':str(index),'filename':f'synthetic-{index}.pdf','status':'review','revision':2,'reviewed':False,
                'invoice':Invoice(number=f'INVOICE-{index}',net=str(count*10),tax='0',
                    lines=[{'sku':f'SKU-{index}-{n}','gtin':f'000{index}{n}','qty':'2','price':'5'} for n in range(count)]).model_dump(mode='json')}
        originals.append(record);app.state.store.job(str(index),record)
    with TestClient(app) as client:
        url='/api/exports/extraction-batch';headers={'X-Studio-Request':'1'}
        selected=[{'id':'1','revision':2},{'id':'2','revision':2}]
        assert client.post(url,headers=headers,json={'jobs':selected}).status_code==400
        body={'jobs':selected,'acknowledge_unvalidated':True}
        response=client.post(url,headers=headers,json=body)
        assert response.status_code==200
        book=load_workbook(io.BytesIO(response.content))
        assert book.sheetnames==['Header','Tax_Breakdown','Details']
        assert [s.max_column for s in book]==[13,3,6]
        assert [r[0] for r in book['Header'].iter_rows(min_row=2,values_only=True)]==[1,2]
        assert [r[0] for r in book['Tax_Breakdown'].iter_rows(min_row=2,values_only=True)]==[1,2]
        assert [r[0] for r in book['Details'].iter_rows(min_row=2,values_only=True)]==[1,2,2]
        assert book['Header']['B3'].value=='INVOICE-2'
        assert book['Details']['C4'].value=='00021'
        assert [app.state.store.job(str(i)) for i in (1,2)]==originals
        with app.state.store.connection() as c:assert app.state.store.ledger(c)==[]
        assert client.post(url,headers=headers,json={**body,'jobs':[selected[0],selected[0]]}).status_code==400
        assert client.post(url,headers=headers,json={**body,'jobs':[{'id':'1','revision':1}]}).status_code==409
        for status in ('processing','queued'):
            app.state.store.job('2',{**originals[1],'status':status})
            assert client.post(url,headers=headers,json=body).status_code==409
        app.state.store.job('2',{**originals[1],'invoice':Invoice().model_dump(mode='json')})
        assert client.post(url,headers=headers,json=body).status_code==409


def _rules_lines(item='00099',flagged=False,evidence=True):
    cell={'value':None if flagged else item,'flagged':flagged,'reason':'Not found' if flagged else '',
          'evidence':[{'kind':'sheet','source':'Item Master ITEM_PARENT','reference':'row 2','rule':'ALG-021'}] if evidence and not flagged else []}
    return {'revision':4,'lines':[{'line':1,'cells':{'Item':cell}}],
            'issues':[{'rule':'ALG-016-VAR','code':'Item Review','message':'Synthetic variant check','line':1},
                      {'rule':'ALG-018-OCR','code':'Item Review','message':'Synthetic OCR check','line':2}]}


def test_rules_matched_item_parent_fills_the_details_item_when_the_review_box_is_empty():
    invoice=Invoice(number='SYN-RULES',lines=[{'sku':'VENDOR-1','qty':'1','price':'2'},{'sku':'VENDOR-2','qty':'1','price':'2'}])
    book=load_workbook(io.BytesIO(extraction_workbook(invoice,'synthetic.pdf',4,_rules_lines())))
    assert book['Details']['B2'].value=='00099'
    assert book['Details']['B2'].data_type=='s'
    assert 'Rules match, not validated: Item Master ITEM_PARENT (ALG-021)' in book['Details']['B2'].comment.text
    assert 'Review flag ALG-016-VAR: Synthetic variant check' in book['Details']['B2'].comment.text
    assert 'ALG-018-OCR' not in book['Details']['B2'].comment.text
    assert book['Details']['B3'].value is None
    assert 'requires confirmation' in book['Details']['B3'].comment.text


def test_with_rules_the_draft_item_is_the_target_item_and_a_typed_item_never_overrides_it():
    invoice=Invoice(number='SYN-RULES',lines=[{'sku':'VENDOR-1','qty':'1','price':'2'}])
    for lines in (_rules_lines(flagged=True),_rules_lines(evidence=False),None):
        book=load_workbook(io.BytesIO(extraction_workbook(invoice,'synthetic.pdf',4,lines)))
        assert book['Details']['B2'].value is None
    typed=Invoice(number='SYN-RULES',lines=[{'item_id':'00042','sku':'VENDOR-1','qty':'1','price':'2'}])
    book=load_workbook(io.BytesIO(extraction_workbook(typed,'synthetic.pdf',4,_rules_lines())))
    assert book['Details']['B2'].value=='00099'
    assert 'review candidate only for the rules: 00042' in book['Details']['B2'].comment.text
    book=load_workbook(io.BytesIO(extraction_workbook(typed,'synthetic.pdf',4,_rules_lines(flagged=True))))
    assert book['Details']['B2'].value is None
    book=load_workbook(io.BytesIO(extraction_workbook(typed,'synthetic.pdf',4,None)))
    assert book['Details']['B2'].value=='00042'


def test_draft_download_uses_only_rules_of_the_current_revision(tmp_path,monkeypatch):
    monkeypatch.setenv('INV_STUDIO_DATA',str(tmp_path/'module-default'))
    from app.main import create_app
    app=create_app(tmp_path)
    invoice=Invoice(number='REVIEW-2',net='2',lines=[{'sku':'VENDOR-1','qty':'1','price':'2'}]).model_dump(mode='json')
    headers={'X-Studio-Request':'1'};url='/api/jobs/synthetic/extraction-draft'
    with TestClient(app) as client:
        for rules_revision,expected in ((2,'00099'),(1,None)):
            record={'id':'synthetic','filename':'synthetic.pdf','status':'review','revision':2,'reviewed':False,
                    'invoice':invoice,'rules':{**_rules_lines(),'revision':rules_revision}}
            app.state.store.job('synthetic',record)
            response=client.post(url,headers=headers,json={'revision':2,'acknowledge_unvalidated':True})
            assert response.status_code==200
            assert load_workbook(io.BytesIO(response.content))['Details']['B2'].value==expected
            assert app.state.store.job('synthetic')==record


def test_draft_item_equals_the_target_export_item_line_by_line():
    from app.excel import rules_workbook
    from app.matching import rules_view
    from tests.test_rules_wiring import approved_result
    invoice=Invoice(number='INV-1',lines=[{'item_id':'00042','sku':'VENDOR-1','qty':'3','price':'10'}])
    view=rules_view(approved_result())
    export,_=rules_workbook([view])
    draft=extraction_workbook(invoice,'synthetic.pdf',1,view)
    item=lambda content:load_workbook(io.BytesIO(content))['Details']['B2'].value
    assert item(draft)==item(export)==345000001
    # Matched first, then moved to Exception by a later check: the export cell is empty, so the draft's is too.
    result=approved_result(status='Exception')
    result['lines'][0]['Item']='';result['lines'][0]['Validation Status']='Exception'
    result['lineage']=[x for x in result['lineage'] if x.get('target')!='Item']
    assert item(extraction_workbook(invoice,'synthetic.pdf',1,rules_view(result))) is None
