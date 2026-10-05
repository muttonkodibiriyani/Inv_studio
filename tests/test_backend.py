import io
import json
import threading
import time
from copy import deepcopy
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.excel import batch_workbook
from app.matching import enrich, key, validate
from app.models import Invoice, Policy
from app.references import import_references, reference_workbook, validate_references


ROOT = Path(__file__).resolve().parents[1]
MUTATION = {"X-Studio-Request": "1"}
OPTIONS = {
    "engine": "auto",
    "ai_fallback": False,
    "provider": "openai",
    "model": "",
    "language": "en",
}


def sample(name):
    return json.loads((ROOT / "samples" / name).read_text())


def test_inbox_load_uses_one_database_connection_for_many_invoices(tmp_path,monkeypatch):
    from contextlib import contextmanager
    monkeypatch.setenv('INV_STUDIO_DATA',str(tmp_path/'module-default'))
    from app.main import create_app
    app=create_app(tmp_path/'inbox');store=app.state.store
    for index in range(40):
        store.job(str(index),{'id':str(index),'status':'review','reviewed':False,'revision':1,
                             'invoice':Invoice(number=f'SYNTHETIC-{index}').model_dump(mode='json')})
    original=store.connection;calls=[]
    @contextmanager
    def counted(transaction=False):
        calls.append(transaction)
        with original(transaction) as c:yield c
    monkeypatch.setattr(store,'connection',counted)
    with TestClient(app) as client:
        response=client.get('/api/state')
        assert response.status_code==200
        assert len(response.json()['jobs'])==40
    assert calls==[False]


@pytest.fixture
def refs():
    return import_references((ROOT / "samples" / "reference.json").read_bytes(), "reference.json")


def prepared_invoice(refs, name="invoice.json"):
    invoice, _ = enrich(Invoice.model_validate(sample(name)), refs)
    return invoice


def issue_codes(result, line=None):
    return {
        issue["code"]
        for issue in result["issues"]
        if line is None or issue.get("line") == line
    }


def test_confirmed_internal_item_still_requires_exact_approved_identity_and_scope(refs):
    invoice = prepared_invoice(refs)
    invoice.lines[0].item_id = "345000101"
    assert validate(invoice, refs, Policy(), reviewed=True)["ready"]
    invoice.lines[0].item_id = "345000102"
    conflict = validate(invoice, refs, Policy(), reviewed=True)
    assert not conflict["ready"] and "ITEM_CONFLICT" in issue_codes(conflict, 1)
    invoice.lines[0].item_id = "UNAPPROVED-CANDIDATE"
    missing = validate(invoice, refs, Policy(), reviewed=True)
    assert not missing["ready"] and "ITEM" in issue_codes(missing, 1)
    invoice.lines[0].item_id = "345000101"
    invoice.lines[0].sku = None
    invoice.lines[0].gtin = None
    assert validate(invoice, refs, Policy(), reviewed=True)["ready"]


def test_exact_reference_enrichment_and_validation(refs):
    invoice, provenance = enrich(Invoice.model_validate(sample("invoice.json")), refs)

    assert (invoice.seller, invoice.site, invoice.buyer, invoice.location) == (
        "LUMENA_AE",
        "91001",
        "RETAIL_AE",
        "38001",
    )
    assert invoice.taxCode == "UEPV05"
    assert {entry["field"] for entry in provenance} >= {
        "seller",
        "site",
        "buyer",
        "location",
        "taxCode",
    }

    result = validate(invoice, refs, Policy(), reviewed=True)
    assert result["ready"] is True
    assert result["issues"] == []
    assert [match["item"] for match in result["matches"]] == ["345000101", "345000102"]
    assert result["matches"][0]["gtin"] == "00012345678905"


@pytest.mark.parametrize(
    ("change", "expected"),
    [
        (lambda invoice: setattr(invoice.lines[0], "qty", None), "QTY"),
        (lambda invoice: setattr(invoice.lines[0], "gtin", "00000000000000"), "ITEM"),
        (lambda invoice: setattr(invoice.lines[0], "sku", "LUM-CRM50"), "ITEM_CONFLICT"),
    ],
)
def test_line_identity_and_required_value_failures_are_held(refs, change, expected):
    invoice = prepared_invoice(refs)
    change(invoice)

    result = validate(invoice, refs, Policy(), reviewed=True)

    assert result["ready"] is False
    assert expected in issue_codes(result, line=1)


def test_only_accepted_receipts_count_and_repeated_lines_accumulate(refs):
    refs = deepcopy(refs)
    refs["receipts"].extend(
        [
            {"id": "PENDING", "po": "70003", "line": "1", "qty": "100", "status": "pending"},
            {"id": "REJECTED", "po": "70003", "line": "1", "qty": "100", "status": "rejected"},
            {"id": "RETURN", "po": "70003", "line": "1", "qty": "-1", "status": "accepted"},
        ]
    )
    validate_references(refs)
    raw = sample("invoice.json")
    raw.update(number="REPEATED-1", po="70003", net="480", tax="24")
    raw["lines"] = [
        {**raw["lines"][0], "qty": "4"},
        {**raw["lines"][0], "qty": "4"},
    ]
    invoice, _ = enrich(Invoice.model_validate(raw), refs)

    result = validate(invoice, refs, Policy(), reviewed=True)

    receipt_issues = [issue for issue in result["issues"] if issue["code"] == "RECEIPT"]
    assert [(issue["line"], issue["message"]) for issue in receipt_issues] == [
        (2, "Needs 8 units; 5 accepted received units remain")
    ]

    ledger = [{"invoice_key": "other|buyer|invoice", "allocations": [{"key": "70003|1", "qty": "3"}]}]
    result = validate(invoice, refs, Policy(), ledger=ledger, reviewed=True)
    assert any(issue["code"] == "RECEIPT" and issue["line"] == 1 for issue in result["issues"])


def test_duplicate_invoice_key_and_missing_references_are_explicit(refs):
    invoice = prepared_invoice(refs)
    duplicate = [{"invoice_key": key(invoice), "allocations": []}]

    assert "DUPLICATE" in issue_codes(validate(invoice, refs, Policy(), duplicate, reviewed=True))

    missing = validate(invoice, {}, Policy(), reviewed=True)
    assert {"SITE", "ROUTE", "PO", "ITEM", "TAX"} <= issue_codes(missing)


def test_supplier_name_must_match_approved_name_or_alias(refs):
    invoice = prepared_invoice(refs)
    invoice.supplier_name = "LUMENA beauty-trading, LLC"
    assert "SUPPLIER_NAME" not in issue_codes(
        validate(invoice, refs, Policy(), reviewed=True)
    )

    refs = deepcopy(refs)
    refs["sites"][0]["aliases"] = ["Lumena Trading"]
    invoice.supplier_name = "Lumena Trading"
    assert validate(invoice, refs, Policy(), reviewed=True)["ready"] is True

    invoice.supplier_name = "Different Supplier LLC"
    assert "SUPPLIER_NAME" in issue_codes(
        validate(invoice, refs, Policy(), reviewed=True)
    )


def test_three_decimal_currency_rounding_drives_net_and_tax(refs):
    refs = deepcopy(refs)
    order = next(order for order in refs["orders"] if order["id"] == "70001")
    order["currency"] = "KWD"
    order["lines"][0]["price"] = "1.2344"
    route = next(route for route in refs["routes"] if route["site"] == "91001" and route["buyer"] == "RETAIL_AE")
    route["currency"] = "KWD"
    refs["taxRules"] = [
        {"origin": "AE", "market": "AE", "currency": "KWD", "code": "KW05", "rate": "0.05"}
    ]
    raw = sample("invoice.json")
    raw.update(number="KWD-001", currency="KWD", net="2.469", tax="0.123")
    raw["lines"] = [{**raw["lines"][0], "qty": "2", "price": "1.2344"}]
    invoice, _ = enrich(Invoice.model_validate(raw), refs)
    policy = Policy(total_tolerance=Decimal("0"), currency_decimals={"KWD": 3})

    result = validate(invoice, refs, policy, reviewed=True)
    assert result["ready"] is True
    assert result["net"] == "2.469"

    invoice.tax = Decimal("0.124")
    assert "TAX_TOTAL" in issue_codes(validate(invoice, refs, policy, reviewed=True))
    invoice.tax = Decimal("0.123")
    invoice.net = Decimal("2.468")
    assert "TOTAL" in issue_codes(validate(invoice, refs, policy, reviewed=True))


def test_batch_workbook_joins_transactions_and_uses_exact_identifier_types(refs):
    refs = deepcopy(refs)
    first_order = next(order for order in refs["orders"] if order["id"] == "70001")
    first_order["location"] = "900001"
    first_order["location_type"] = "Warehouse (W)"
    first = prepared_invoice(refs, "invoice.json")
    first.number = "0380"
    second = prepared_invoice(refs, "invoice-2.json")
    entries = [
        (first, validate(first, refs, Policy(), reviewed=True)),
        (second, validate(second, refs, Policy(), reviewed=True)),
    ]

    workbook = load_workbook(io.BytesIO(batch_workbook(entries)), data_only=False)

    assert [workbook["Header"].cell(row, 1).value for row in (2, 3)] == [1, 2]
    assert [workbook["Tax_Breakdown"].cell(row, 1).value for row in (2, 3)] == [1, 2]
    assert [workbook["Details"].cell(row, 1).value for row in (2, 3, 4)] == [1, 1, 2]
    assert workbook["Details"].max_column == 6
    assert workbook["Header"]["B2"].value == "0380"
    assert workbook["Header"]["B2"].data_type == "s"
    assert workbook["Header"]["C2"].value == 91001
    assert workbook["Header"]["D2"].value == 70001
    assert workbook["Header"]["E2"].value == 900001
    assert workbook["Header"]["F2"].value == "Warehouse (W)"
    assert workbook["Details"]["C2"].value == "00012345678905"
    assert workbook["Details"]["C2"].data_type == "s"
    assert workbook["Details"]["C2"].number_format == "@"


def test_location_type_must_come_from_the_purchase_order_reference(refs):
    refs = deepcopy(refs)
    order = next(order for order in refs["orders"] if order["id"] == "70001")
    order["location_type"] = ""
    invoice = prepared_invoice(refs)

    assert "LOCATION_TYPE" in issue_codes(
        validate(invoice, refs, Policy(), reviewed=True)
    )
    with pytest.raises(ValueError, match="location_type"):
        validate_references(refs)


def test_reference_xlsx_round_trip_and_rejects_formulas_duplicates_and_bad_sheets(refs):
    content = reference_workbook(refs)
    restored = import_references(content, "references.xlsx")
    assert restored["items"][0]["gtin"] == "00012345678905"
    assert restored["orders"][0]["lines"][0]["price"] == "60"

    workbook = load_workbook(io.BytesIO(content))
    workbook["Items"]["C2"] = "=1+1"
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(ValueError, match="formulas are not accepted"):
        import_references(output.getvalue(), "formula.xlsx")

    duplicate = deepcopy(refs)
    duplicate["items"].append(deepcopy(duplicate["items"][0]))
    with pytest.raises(ValueError, match="Duplicate keys in items"):
        import_references(json.dumps(duplicate).encode(), "duplicate.json")

    workbook = load_workbook(io.BytesIO(content))
    del workbook["Receipts"]
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(ValueError, match="Missing worksheet Receipts"):
        import_references(output.getvalue(), "missing-sheet.xlsx")

    with pytest.raises(ValueError, match="damaged or is not a valid"):
        import_references(b"not a zip archive", "broken.xlsx")


@pytest.mark.parametrize(
    ("mutate", "message"),
    [
        (
            lambda data: data["orders"][0].update(lines={"id": "not-a-list"}),
            "PO lines must be a non-empty list",
        ),
        (
            lambda data: data["sites"][0].update(name={"nested": "not-a-name"}),
            "required values must be strings or numbers",
        ),
        (
            lambda data: data["sites"][0].update(aliases="not-a-list"),
            "Supplier aliases must be a list",
        ),
        (
            lambda data: data["sites"][0].update(aliases=[{"nested": "not-a-name"}]),
            "Supplier aliases must be a list",
        ),
    ],
)
def test_reference_json_rejects_nested_values_before_semantic_matching(refs, mutate, message):
    malformed = deepcopy(refs)
    mutate(malformed)

    with pytest.raises(ValueError, match=message):
        import_references(json.dumps(malformed).encode(), "references.json")


@pytest.fixture
def api_client(tmp_path, monkeypatch):
    monkeypatch.setenv("INV_STUDIO_DATA", str(tmp_path / "module-default"))
    import app.main as main

    def fake_process(path, options, store, ai_reader, progress):
        progress("fixture", "Reading synthetic invoice")
        invoice = json.loads(path.read_text())
        return {
            "invoice": invoice,
            "text": "synthetic fixture",
            "boxes": [],
            "trace": [{"engine": "fixture", "status": "extracted"}],
            "selected_engine": "fixture",
            "completeness": 1,
        }

    monkeypatch.setattr(main, "process", fake_process)
    app = main.create_app(tmp_path / "data")
    with TestClient(app) as client:
        yield client, app
    app.state.pool.shutdown(wait=True)


def preflight(client, files, *, confirm=True, options=None):
    options = options or OPTIONS
    response = client.post(
        "/api/preflight",
        headers=MUTATION,
        json={"options": options, "files": files},
    )
    assert response.status_code == 200, response.text
    token = response.json()["token"]
    if confirm:
        response = client.post(
            "/api/preflight/confirm", headers=MUTATION, json={"token": token}
        )
        assert response.status_code == 200, response.text
    return token


def upload_invoice(client, invoice, name="invoice.json"):
    content = json.dumps(invoice).encode()
    token = preflight(client, [{"name": name, "size": len(content)}])
    response = client.post(
        "/api/invoices",
        headers=MUTATION,
        data={"options": json.dumps(OPTIONS), "preflight_token": token},
        files={"file": (name, content, "application/json")},
    )
    assert response.status_code == 200, response.text
    job_id = response.json()["id"]
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] not in ("queued", "processing"):
            return job
        time.sleep(0.01)
    pytest.fail("synthetic extraction did not finish")


def review_job(client, job, *, confirm=True):
    response = client.post(
        f"/api/jobs/{job['id']}/review",
        headers=MUTATION,
        json={"invoice": job["invoice"], "revision": job["revision"], "confirm": confirm},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_preflight_requires_confirmation_is_file_bound_and_invalidates_on_rule_change(api_client):
    client, _ = api_client
    content = json.dumps(sample("invoice.json")).encode()
    plan = client.post(
        "/api/preflight",
        headers=MUTATION,
        json={"options": OPTIONS, "files": [{"name": "invoice.json", "size": len(content)}]},
    )
    assert plan.status_code == 200
    assert any("No reference files loaded" in warning for warning in plan.json()["warnings"])
    token = plan.json()["token"]

    unconfirmed = client.post(
        "/api/invoices",
        headers=MUTATION,
        data={"options": json.dumps(OPTIONS), "preflight_token": token},
        files={"file": ("invoice.json", content, "application/json")},
    )
    assert unconfirmed.status_code == 409
    client.post("/api/preflight/confirm", headers=MUTATION, json={"token": token}).raise_for_status()

    wrong_file = client.post(
        "/api/invoices",
        headers=MUTATION,
        data={"options": json.dumps(OPTIONS), "preflight_token": token},
        files={"file": ("different.json", content, "application/json")},
    )
    assert wrong_file.status_code == 409
    accepted = client.post(
        "/api/invoices",
        headers=MUTATION,
        data={"options": json.dumps(OPTIONS), "preflight_token": token},
        files={"file": ("invoice.json", content, "application/json")},
    )
    assert accepted.status_code == 200
    repeated = client.post(
        "/api/invoices",
        headers=MUTATION,
        data={"options": json.dumps(OPTIONS), "preflight_token": token},
        files={"file": ("invoice.json", content, "application/json")},
    )
    assert repeated.status_code == 409

    changed_token = preflight(
        client, [{"name": "invoice.json", "size": len(content)}]
    )
    policy = Policy(total_tolerance=Decimal("0.02")).model_dump(mode="json")
    client.post("/api/policy", headers=MUTATION, json=policy).raise_for_status()
    changed = client.post(
        "/api/invoices",
        headers=MUTATION,
        data={"options": json.dumps(OPTIONS), "preflight_token": changed_token},
        files={"file": ("invoice.json", content, "application/json")},
    )
    assert changed.status_code == 409
    assert "Rules or references changed" in changed.json()["detail"]


def test_invalid_upload_does_not_consume_confirmed_file_entry(api_client):
    client, _ = api_client
    content = b"not a pdf"
    token = preflight(client, [{"name": "broken.pdf", "size": len(content)}])
    request = {
        "headers": MUTATION,
        "data": {"options": json.dumps(OPTIONS), "preflight_token": token},
        "files": {"file": ("broken.pdf", content, "application/pdf")},
    }

    first = client.post("/api/invoices", **request)
    second = client.post("/api/invoices", **request)

    assert first.status_code == second.status_code == 400
    assert first.json()["detail"] == second.json()["detail"] == "File is not a PDF"


def test_invalid_options_json_returns_sanitized_client_error(api_client):
    client, _ = api_client
    marker = "credential-must-not-be-echoed"
    response = client.post(
        "/api/invoices",
        headers=MUTATION,
        data={
            "options": '{"engine":"auto","api_key":"' + marker + '"',
            "preflight_token": "unused",
        },
        files={"file": ("invoice.json", b"{}", "application/json")},
    )

    assert response.status_code == 422
    assert "Invalid input" in response.json()["detail"]
    assert marker not in response.text


def test_rule_change_during_extraction_never_enriches_or_readies_job(tmp_path, monkeypatch):
    monkeypatch.setenv("INV_STUDIO_DATA", str(tmp_path / "module-default"))
    import app.main as main

    started = threading.Event()
    release = threading.Event()

    def blocked_process(path, options, store, ai_reader, progress):
        started.set()
        assert release.wait(5), "test did not release blocked extraction"
        return {
            "invoice": json.loads(path.read_text()),
            "text": "synthetic fixture",
            "boxes": [],
            "trace": [{"engine": "fixture", "status": "extracted"}],
            "selected_engine": "fixture",
            "completeness": 1,
        }

    monkeypatch.setattr(main, "process", blocked_process)
    app = main.create_app(tmp_path / "data")
    try:
        with TestClient(app) as client:
            client.post("/api/references/demo", headers=MUTATION).raise_for_status()
            invoice = sample("invoice.json")
            content = json.dumps(invoice).encode()
            token = preflight(client, [{"name": "invoice.json", "size": len(content)}])
            uploaded = client.post(
                "/api/invoices",
                headers=MUTATION,
                data={"options": json.dumps(OPTIONS), "preflight_token": token},
                files={"file": ("invoice.json", content, "application/json")},
            )
            assert uploaded.status_code == 200, uploaded.text
            assert started.wait(2), "mock extraction did not start"

            changed_policy = Policy(total_tolerance=Decimal("0.02")).model_dump(mode="json")
            client.post("/api/policy", headers=MUTATION, json=changed_policy).raise_for_status()
            release.set()

            deadline = time.monotonic() + 5
            while time.monotonic() < deadline:
                job = client.get(f"/api/jobs/{uploaded.json()['id']}").json()
                if job["status"] not in ("queued", "processing"):
                    break
                time.sleep(0.01)
            else:
                pytest.fail("blocked extraction did not finish")

            assert job["status"] == "error"
            assert job["validation"]["ready"] is False
            assert job["invoice"]["seller"] is None
            assert job["provenance"] == []
            assert job["error"] == (
                "Rules or references changed during extraction. "
                "Review a new processing plan and retry."
            )
    finally:
        release.set()
        app.state.pool.shutdown(wait=True)


def test_review_revision_export_gate_and_exported_invoice_immutability(api_client):
    client, _ = api_client
    client.post("/api/references/demo", headers=MUTATION).raise_for_status()
    job = upload_invoice(client, sample("invoice.json"))
    assert "REVIEW" in issue_codes(job["validation"])

    held = client.post(
        f"/api/jobs/{job['id']}/export",
        headers=MUTATION,
        json={"revision": job["revision"]},
    )
    assert held.status_code == 409

    stale = client.post(
        f"/api/jobs/{job['id']}/review",
        headers=MUTATION,
        json={"invoice": job["invoice"], "revision": job["revision"] - 1, "confirm": True},
    )
    assert stale.status_code == 409

    reviewed = review_job(client, job)
    assert reviewed["validation"]["ready"] is True
    exported = client.post(
        f"/api/jobs/{job['id']}/export",
        headers=MUTATION,
        json={"revision": reviewed["revision"]},
    )
    assert exported.status_code == 200
    content = client.get(exported.json()["url"]).content

    immutable = client.post(
        f"/api/jobs/{job['id']}/review",
        headers=MUTATION,
        json={"invoice": reviewed["invoice"], "revision": reviewed["revision"] + 1, "confirm": True},
    )
    assert immutable.status_code == 409
    client.post("/api/references/demo", headers=MUTATION).raise_for_status()
    assert client.get(exported.json()["url"]).content == content


def test_batch_export_is_transactional_when_any_invoice_is_unreviewed(api_client):
    client, _ = api_client
    client.post("/api/references/demo", headers=MUTATION).raise_for_status()
    first = review_job(client, upload_invoice(client, sample("invoice.json"), "first.json"))
    second = upload_invoice(client, sample("invoice-2.json"), "second.json")

    response = client.post(
        "/api/exports/batch",
        headers=MUTATION,
        json={
            "jobs": [
                {"id": first["id"], "revision": first["revision"]},
                {"id": second["id"], "revision": second["revision"]},
            ]
        },
    )

    assert response.status_code == 409
    assert client.get(f"/api/jobs/{first['id']}").json().get("export_id") is None
    assert client.get(f"/api/jobs/{second['id']}").json().get("export_id") is None


def test_two_invoice_api_batch_uses_transaction_numbers_across_all_sheets(api_client):
    client, _ = api_client
    client.post("/api/references/demo", headers=MUTATION).raise_for_status()
    first = review_job(client, upload_invoice(client, sample("invoice.json"), "first.json"))
    second = review_job(client, upload_invoice(client, sample("invoice-2.json"), "second.json"))
    response = client.post(
        "/api/exports/batch",
        headers=MUTATION,
        json={
            "jobs": [
                {"id": first["id"], "revision": first["revision"]},
                {"id": second["id"], "revision": second["revision"]},
            ]
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["count"] == 2

    workbook = load_workbook(io.BytesIO(client.get(response.json()["url"]).content), data_only=False)
    assert [workbook["Header"].cell(row, 1).value for row in (2, 3)] == [1, 2]
    assert [workbook["Tax_Breakdown"].cell(row, 1).value for row in (2, 3)] == [1, 2]
    assert [workbook["Details"].cell(row, 1).value for row in (2, 3, 4)] == [1, 1, 2]


def test_batch_aggregates_prior_allocations_and_rolls_back(api_client):
    client, _ = api_client
    client.post("/api/references/demo", headers=MUTATION).raise_for_status()
    invoices = []
    for number, quantity in (("ALLOC-1", "60"), ("ALLOC-2", "50")):
        invoice = sample("invoice.json")
        invoice.update(number=number, net=str(Decimal(quantity) * 60), tax=str(Decimal(quantity) * 3))
        invoice["lines"] = [{**invoice["lines"][0], "qty": quantity}]
        invoices.append(invoice)
    first = review_job(client, upload_invoice(client, invoices[0], "alloc-1.json"))
    second = review_job(client, upload_invoice(client, invoices[1], "alloc-2.json"))

    response = client.post(
        "/api/exports/batch",
        headers=MUTATION,
        json={
            "jobs": [
                {"id": first["id"], "revision": first["revision"]},
                {"id": second["id"], "revision": second["revision"]},
            ]
        },
    )

    assert response.status_code == 409
    assert "accepted received units remain" in response.json()["detail"]
    assert client.get(f"/api/jobs/{first['id']}").json().get("export_id") is None
    assert client.get(f"/api/jobs/{second['id']}").json().get("export_id") is None


def test_api_keys_are_never_echoed_and_database_contains_only_ciphertext(api_client, monkeypatch):
    client, app = api_client
    import httpx

    monkeypatch.setattr(
        httpx.Client, "get", lambda self, url, headers=None: httpx.Response(200, json={"data": []})
    )
    secret = "sk-test-super-secret-never-echo"
    response = client.post(
        "/api/connections/openai", headers=MUTATION, json={"api_key": secret}
    )
    assert response.status_code == 200
    assert secret not in response.text
    assert secret not in client.get("/api/state").text
    assert secret.encode() not in app.state.store.path.read_bytes()

    short_secret = "leaky"
    invalid = client.post(
        "/api/connections/anthropic", headers=MUTATION, json={"api_key": short_secret}
    )
    assert invalid.status_code == 422
    assert short_secret not in invalid.text


def test_rejected_api_key_is_not_saved(api_client, monkeypatch):
    client, app = api_client
    import httpx

    monkeypatch.setattr(
        httpx.Client, "get", lambda self, url, headers=None: httpx.Response(401)
    )
    secret = "sk-ant-rejected-key-never-echo"
    response = client.post(
        "/api/connections/anthropic", headers=MUTATION, json={"api_key": secret}
    )
    assert response.status_code == 400
    assert "rejected this API key" in response.json()["detail"]
    assert secret not in response.text
    assert app.state.store.secret("anthropic") is None


def test_unreachable_provider_still_saves_unverified_key(api_client, monkeypatch):
    client, app = api_client
    import httpx

    def offline(self, url, headers=None):
        raise httpx.ConnectError("offline")

    monkeypatch.setattr(httpx.Client, "get", offline)
    response = client.post(
        "/api/connections/openai", headers=MUTATION, json={"api_key": "sk-test-offline-key"}
    )
    assert response.status_code == 200
    assert response.json() == {"connected": True, "verified": False}
    assert app.state.store.secret("openai") == "sk-test-offline-key"


def test_mutating_api_requires_request_protection_header(api_client):
    client, _ = api_client
    response = client.post("/api/references/demo")
    assert response.status_code == 403
    assert response.json()["detail"] == "Missing request protection header"


def test_damaged_reference_workbook_returns_safe_client_error(api_client):
    client, _ = api_client
    response = client.post(
        "/api/references",
        headers=MUTATION,
        files={
            "file": (
                "broken.xlsx",
                b"not a zip archive",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 400
    assert response.json() == {
        "detail": "Reference workbook is damaged or is not a valid .xlsx file"
    }


def test_restart_recovers_jobs_older_than_visible_history(tmp_path):
    from app.store import Store
    from app.main import create_app
    store=Store(tmp_path)
    store.job("old-active",{"id":"old-active","status":"processing"})
    for n in range(205):store.job(str(n),{"id":str(n),"status":"review"})
    assert not any(x["id"]=="old-active" for x in store.jobs())
    app=create_app(tmp_path)
    assert app.state.store.job("old-active")["status"]=="error"
    app.state.pool.shutdown(wait=True)
