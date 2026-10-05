"""Fine-grained rules over an imported lookup catalog and through the API. Synthetic rows only."""

import io
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import fine_rules as fr
from app.fine_rules_source import LookupRulesSource
from app.models import Invoice, Line
from app.reference_lookup import ReferenceLookup
from app.store import Store
from tests.test_reference_lookup import catalog


ITEM_HASH = "d" * 64
PO_HASH = "e" * 64
# Same roles as the deployed catalog manifest.
LIVE_COLUMNS = {
    "item": {"description": ["ITEM_DESC", "ITEM_DESC_SECONDARY", "SHORT_DESC"], "sku": ["VPN"], "gtin": ["ITEM"],
             "uom": [], "pack": [], "supplier": ["SUPPLIER_NAME"], "site": ["SUPPLIER"],
             "identity": ["ITEM_PARENT"], "internal_item": ["ITEM_PARENT"]},
    "po": {"description": [], "sku": ["RMS_ITEM_ID"], "gtin": ["BARCODE"], "uom": [], "pack": [],
           "supplier": ["SUP_NAME"], "site": ["EBS_SUPPLIER_CODE"], "po": ["RMS_ORDER_NO", "EXT_ORDER_NO"],
           "internal_item": ["RMS_ITEM_ID"], "identity": ["ASN", "LOCATION"]},
}


def item_row(row, parent, barcode, vpn, desc):
    data = {"ITEM_PARENT": parent, "ITEM": barcode, "VPN": vpn, "SUPPLIER": "22001", "SUPPLIER_NAME": "ABC001RA1KWD",
            "ITEM_DESC": desc, "ITEM_DESC_SECONDARY": "", "SHORT_DESC": "", "UDA_LV_1_VALUE": "BrandA"}
    return {"kind": "item", "source_hash": ITEM_HASH, "source_sheet": "Item Master", "source_row": row,
            "keys": [parent], "data": data, "flags": []}


def po_row(row, qty, cost, barcode, parent):
    data = {"EBS_SUPPLIER_CODE": "ABC001", "LOCATION": "38091", "RMS_ORDER_NO": "13000001", "EXT_ORDER_NO": "",
            "ASN": "", "BARCODE": barcode, "RMS_ITEM_ID": parent, "QTY_RECEIVED": qty, "TOTAL COST": cost,
            "CURRENCY_CODE": "KWD", "RECEIPT_DATE": "2026-01-10", "SUP_NAME": "ABC"}
    return {"kind": "po", "source_hash": PO_HASH, "source_sheet": "POGRN", "source_row": row, "keys": ["13000001"],
            "data": data, "flags": []}


ROWS = [item_row(2, "345000001", "ULT_0012345678905", "100001", "Glow Serum Rose 30ml"),
        item_row(3, "345000002", "ULT_0098765432109", "100002", "Matte Lipstick Red 4g"),
        po_row(2, 3, 30, "0012345678905", "345000001"), po_row(3, 2, 40, "0098765432109", "345000002")]
CONFIG = {"location_master": {"38091": {"type": fr.STORE, "market": "Kuwait"}},
          "supplier_site_currency": [{"supplier_site": "22001", "market": "Kuwait", "currency": "KWD"}]}
INVOICE = Invoice(number="INV-API", supplier_name="ABC Trading LLC", date="2026-01-15", date_printed="15/01/2026",
                  currency="KWD", taxCode="VAT0", net=Decimal(70), tax=Decimal(0), lines=[
                      Line(gtin="0012345678905", sku="100001", description="Glow Serum Rose", qty=Decimal(3),
                           price=Decimal(10), net_amount=Decimal(30)),
                      Line(sku="100002", description="Matte Lipstick", qty=Decimal(2), price=Decimal(20),
                           net_amount=Decimal(40))])


def imported(tmp_path, monkeypatch, columns=LIVE_COLUMNS):
    monkeypatch.setattr("tests.test_reference_lookup.COLUMNS", columns)
    store = Store(tmp_path / "data")
    archive, manifest = catalog(tmp_path, ROWS)
    ReferenceLookup(store).import_archive(archive, manifest, batch_size=100)
    return store


def test_lookup_source_runs_rules_against_imported_catalog(tmp_path, monkeypatch):
    store = imported(tmp_path, monkeypatch)
    source = LookupRulesSource(store)
    assert [r["ITEM_PARENT"] for r in source.items_by_barcode("0012345678905")] == ["345000001"]
    assert source.items_by_barcode("0012345678905")[0]["_ref"] == "Item Master!2"
    assert len(source.pogrn_by_ebs("ABC001")) == 2 and len(source.pogrn_by_order("13000001")) == 2
    result = fr.run_invoice(INVOICE, source, fr.RulesConfig.from_dict(CONFIG))
    assert result["status"] == "Approved", result["exceptions"]
    assert [l["Item"] for l in result["lines"]] == ["345000001", "345000002"]
    assert result["header"]["Order No"] == "13000001" and result["header"]["Supplier Site"] == "22001"


def test_lookup_source_refuses_unbounded_reads(tmp_path, monkeypatch):
    source = LookupRulesSource(imported(tmp_path, monkeypatch), max_rows=1)
    with pytest.raises(ValueError, match="narrow the search"):
        source.pogrn_by_ebs("ABC001")


def test_fine_rules_api_config_run_exports_and_feedback(tmp_path, monkeypatch):
    monkeypatch.setenv("INV_STUDIO_DATA", str(tmp_path / "module-default"))
    from app.main import create_app
    imported(tmp_path, monkeypatch)
    app = create_app(tmp_path / "data")
    app.state.store.job("job-1", {"id": "job-1", "status": "review", "reviewed": False, "revision": 1,
                                  "filename": "synthetic.pdf", "invoice": INVOICE.model_dump(mode="json")})
    headers = {"x-studio-request": "1"}
    with TestClient(app) as client:
        bad = client.post("/api/fine-rules/config", json={"qty_tolerance": "-1"}, headers=headers)
        assert bad.status_code == 400
        blocked = client.post("/api/fine-rules/target.xlsx", json={"job_ids": ["job-1"]}, headers=headers)
        assert blocked.status_code == 409
        assert client.post("/api/fine-rules/config", json=CONFIG, headers=headers).status_code == 200
        run = client.post("/api/fine-rules/run", json={"job_ids": ["job-1"]}, headers=headers).json()["results"][0]
        assert run["status"] == "Approved" and run["lines"][0]["Unit Cost"] == "10"
        review = client.post("/api/fine-rules/review.xlsx", json={"job_ids": ["job-1"]}, headers=headers)
        assert "07_Match_Workbench" in load_workbook(io.BytesIO(review.content)).sheetnames
        target = client.post("/api/fine-rules/target.xlsx", json={"job_ids": ["job-1"]}, headers=headers)
        assert target.status_code == 200
        assert load_workbook(io.BytesIO(target.content))["Details"]["B2"].value == "345000001"
        proposed = client.post("/api/fine-rules/feedback", headers=headers, json={
            "invoice_line": "INV-API/1", "correction": "345000002", "evidence": "label photo"}).json()
        assert proposed["Decision"] == "Proposed"
        decided = client.post(f"/api/fine-rules/feedback/{proposed['Feedback ID']}/decision", headers=headers,
                              json={"approver": "Item steward", "decision": "Rejected"}).json()
        assert decided["Decision"] == "Rejected"
