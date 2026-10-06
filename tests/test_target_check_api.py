"""Target-sheet check wired into the app: review, confirm-time accuracy, export Checks sheet. Synthetic data only."""

import io
import json

from openpyxl import load_workbook

from tests.test_rules_wiring import H, job, production  # noqa: F401  (pytest fixture)


def confirm(client, shown, **extra):
    body = {"invoice": shown["invoice"], "revision": shown["revision"], "confirm": True, **extra}
    response = client.post("/api/jobs/job-1/review", headers=H, json=body)
    assert response.status_code == 200, response.text
    return response.json()


def test_every_target_cell_has_a_status_and_the_owner_line(production):  # noqa: F811
    app, client = production
    app.state.store.job("job-1", job())
    check = client.get("/api/jobs/job-1/target-check").json()
    assert check["summary"].startswith("Target sheet: ") and "needs checking" in check["summary"]
    assert {c["group"] for c in check["cells"]} <= {"verified", "empty_owner_rule", "empty_flagged", "needs_checking"}
    assert check["counts"]["cells"] == len(check["cells"]) and not check["holds"]
    assert all(c["status"] != "verified" or c["evidence"]["source"] for c in check["cells"])
    # The inbox gets counts only, not the cells.
    listed = next(j for j in client.get("/api/state").json()["jobs"] if j["id"] == "job-1")
    assert set(listed["rules"]["target_check"]) == {"counts", "metric", "holds", "summary", "upc"}
    assert "target_system" not in client.get("/api/jobs/job-1").json()


def test_a_cell_its_evidence_does_not_hold_holds_until_the_owner_confirms(production):  # noqa: F811
    app, client = production
    app.state.store.job("job-1", job(text="Total before tax 70.00\nVAT 0.00"))  # number, date, lines not printed
    shown = client.get("/api/jobs/job-1").json()
    issues = [i for i in shown["validation"]["issues"] if i.get("rule") == "TARGET-CHECK"]
    assert issues and all(i["level"] == "review" and i["blocking"] for i in issues)
    assert "Header.Document" in issues[0]["message"] and "INV-API" not in json.dumps(issues)
    reviewed = confirm(client, shown)
    assert reviewed["validation"]["ready"] is True  # the owner's confirm is the final say (decision 12)
    assert all(i["accepted"] for i in reviewed["validation"]["issues"] if i.get("rule") == "TARGET-CHECK")


def test_confirm_records_field_status_and_change_without_values(production):  # noqa: F811
    app, client = production
    app.state.store.job("job-1", job(text="no totals printed"))  # net and tax empty: the owner fills them
    shown = client.get("/api/jobs/job-1").json()
    reviewed = confirm(client, shown, entries={"header": {"net": "70.00", "tax": "0.50"}})
    tax = next(c for c in client.get("/api/jobs/job-1/target-check").json()["cells"] if c["column"] == "Tax Amount")
    assert (tax["status"], tax["sub"]) == ("verified", "owner_entry")  # stored entry with actor and time
    assert client.get("/api/jobs/job-1/target-check").json()["confirmed"] is True
    summary = client.get("/api/target-check/accuracy").json()
    period = summary["periods"]["all"]
    assert summary["invoices"] == 1 and period["confirms"] == 1
    assert period["fields"]["Header.Tax Amount"] == {"cells": 1, "unchanged": 0, "changed": 1, "accuracy": 0.0}
    assert period["by_status_before"]["empty_flagged"]["changed"] >= 2  # Net and Tax were real gaps
    assert period["fields"]["Header.Supplier Site"] == {"cells": 1, "unchanged": 1, "changed": 0, "accuracy": 1.0}
    assert period["overall"]["cells"] > 10 and summary["periods"]["7d"]["overall"] == period["overall"]
    text = json.dumps(summary)
    assert "0.50" not in text and "70.00" not in text and "INV-API" not in text  # no values
    supplier = reviewed["rules"]["fields"]["site"]["value"]
    assert client.get("/api/target-check/accuracy", params={"supplier": "none"}).json()["periods"]["all"]["confirms"] == 0
    assert client.get("/api/target-check/accuracy", params={"supplier": supplier}).json()["invoices"] == 1
    logged = [e for e in client.get("/api/jobs/job-1/audit").json() if e["event"] == "target_check_confirmed"]
    assert logged and "Header.Tax Amount" in logged[0]["payload"]["changed"] and "0.50" not in json.dumps(logged)


def test_export_appends_the_checks_sheet_and_the_receipt_reports_accuracy(production):  # noqa: F811
    app, client = production
    app.state.store.job("job-1", job())
    reviewed = confirm(client, client.get("/api/jobs/job-1").json())
    exported = client.post("/api/jobs/job-1/export", headers=H, json={"revision": reviewed["revision"]}).json()
    book = load_workbook(io.BytesIO(client.get(exported["url"]).content))
    assert book.sheetnames == ["Header", "Tax_Breakdown", "Details", "Checks"]
    rows = list(book["Checks"].values)
    assert rows[0][:7] == ("Transaction Number", "Sheet", "Column", "Line", "Value", "Status", "Detail")
    assert {r[1] for r in rows[1:]} >= {"Header", "Tax_Breakdown", "Details", "(check)", "(workbook)"}
    # Decision 76: Order Date is one evidence-only Checks row per invoice, never a template column.
    dated = [dict(zip(rows[0], r)) for r in rows[1:] if r[1] == "(evidence)"]
    assert [(d["Transaction Number"], d["Column"], d["Rule"]) for d in dated] == [(1, "Order Date", "POG-009")]
    assert dated[0]["Status"] in {"filled", "empty_flagged"} and "decision 76" in dated[0]["Detail"]
    assert dated[0]["Evidence Source"].endswith(": CREATED_DATE")
    assert all("Order Date" not in [c.value for c in book[name][1]] for name in ("Header", "Tax_Breakdown", "Details"))
    exported_event = [e for e in client.get("/api/jobs/job-1/audit").json() if e["event"] == "exported"][0]["payload"]
    target = exported_event["target_check"]
    assert target["summary"].startswith("Target sheet:") and target["workbook_checks"]["joins"] == "pass"
    assert set(target["metric"]["buckets"]) >= {"verified", "empty_owner_rule", "mismatch"}
