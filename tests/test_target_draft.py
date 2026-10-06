"""Export selected (draft): the target workbook for invoices nobody approved yet, read-only for the jobs
(decisions 67/70/72). Synthetic data only."""

import io
import time

import pytest
from openpyxl import load_workbook

from app.excel import rules_workbook
from app.matching import rules_view

from tests.test_rules_wiring import H, approved_result, job, production  # noqa: F401  (pytest fixture)
from tests.test_target_check_api import confirm

TEMPLATE = ("Header", "Tax_Breakdown", "Details")
KEPT = ("status", "revision", "reviewed", "export_id")


def sheets(content):
    book = load_workbook(io.BytesIO(content))
    return {name: [list(row) for row in book[name].iter_rows(values_only=True)] for name in book.sheetnames}


def draft(client, *jobs):
    return client.post("/api/exports/target-draft", headers=H, json={"jobs": [{"id": i, "revision": r} for i, r in jobs]})


def state(store, *ids):
    return {i: {k: store.job(i).get(k) for k in KEPT} for i in ids}


def files(root):
    return sorted(str(p) for p in root.rglob("*") if p.is_file())


def test_on_an_approved_invoice_the_draft_is_the_export_cell_for_cell_and_changes_nothing(production):  # noqa: F811
    app, client = production
    store = app.state.store
    store.job("job-1", job())
    reviewed = confirm(client, client.get("/api/jobs/job-1").json())
    before = state(store, "job-1")
    audit = len(client.get("/api/jobs/job-1/audit").json())
    response = draft(client, ("job-1", reviewed["revision"]))
    assert response.status_code == 200, response.text
    assert response.headers["content-disposition"].startswith('attachment; filename="DRAFT_Target_')
    assert (response.headers["x-draft-included"], response.headers["x-draft-skipped"]) == ("1", "0")
    assert state(store, "job-1") == before  # not exported, not approved, no revision bump, reviewed kept
    events = [e["event"] for e in client.get("/api/jobs/job-1/audit").json()[audit:]]
    assert "exported" not in events and "target_accuracy" not in str(events)
    got = sheets(response.content)
    exported = client.post("/api/jobs/job-1/export", headers=H, json={"revision": reviewed["revision"]}).json()
    real = sheets(client.get(exported["url"]).content)
    assert list(got) == list(real) == [*TEMPLATE, "Checks"]
    for name in TEMPLATE:
        assert got[name] == real[name]
    # Checks: the export's rows, then the draft's own disclosure row.
    assert got["Checks"][:len(real["Checks"])] == real["Checks"]
    assert [row[1] for row in got["Checks"][len(real["Checks"]):]] == ["(draft)"]


def test_a_held_invoice_gets_its_evidenced_cells_and_empty_cells_where_the_export_refuses(production):  # noqa: F811
    app, client = production
    store = app.state.store
    store.job("job-1", job(text="no totals printed"))
    shown = client.get("/api/jobs/job-1").json()
    assert shown["rules"]["fields"]["net"]["flagged"] and shown["validation"]["ready"] is False
    assert client.post("/api/jobs/job-1/export", headers=H, json={"revision": shown["revision"]}).status_code == 409
    before = state(store, "job-1")
    response = draft(client, ("job-1", shown["revision"]))
    assert response.status_code == 200, response.text
    assert state(store, "job-1") == before
    header = dict(zip(*sheets(response.content)["Header"]))
    fields = shown["rules"]["fields"]
    # A flagged total stays empty; it is never filled to look complete.
    assert header["Total Cost Ex Tax"] is None and fields["net"]["value"] is None
    assert header["Supplier Site"] == int(fields["site"]["value"])
    assert header["Document"] == fields["number"]["value"]


def test_a_mixed_selection_skips_what_has_no_current_rules_and_changes_no_job(production):  # noqa: F811
    app, client = production
    store = app.state.store
    store.job("job-1", job())
    approved = confirm(client, client.get("/api/jobs/job-1").json())
    store.job("job-2", {**job(text="no totals printed"), "id": "job-2", "filename": "held.pdf"})
    held = client.get("/api/jobs/job-2").json()
    # Still extracting: ensure_rules leaves it alone, so it has no current rules.
    store.job("job-3", {**job(), "id": "job-3", "filename": "busy.pdf", "status": "processing"})
    before = state(store, "job-1", "job-2", "job-3")
    response = draft(client, ("job-1", approved["revision"]), ("job-2", held["revision"]), ("job-3", 1))
    assert response.status_code == 200, response.text
    assert (response.headers["x-draft-included"], response.headers["x-draft-skipped"]) == ("2", "1")
    assert state(store, "job-1", "job-2", "job-3") == before
    got = sheets(response.content)
    assert [row[0] for row in got["Header"][1:]] == [1, 2]
    assert [row[0] for row in got["Tax_Breakdown"][1:]] == [1, 2]
    assert [row[0] for row in got["Details"][1:]] == [1] * len(approved["invoice"]["lines"]) + [2] * len(held["invoice"]["lines"])
    skipped = [row for row in got["Checks"] if row[1] == "(skipped)"]
    assert [(row[2], row[5]) for row in skipped] == [("busy.pdf", "Skipped")]
    assert skipped[0][7] == "Extraction has not finished"


def test_an_unknown_invoice_is_skipped_not_a_failed_draft(production):  # noqa: F811
    app, client = production
    store = app.state.store
    store.job("job-1", job())
    approved = confirm(client, client.get("/api/jobs/job-1").json())
    response = draft(client, ("job-1", approved["revision"]), ("job-gone", 1))
    assert response.status_code == 200, response.text
    assert (response.headers["x-draft-included"], response.headers["x-draft-skipped"]) == ("1", "1")
    skipped = [row for row in sheets(response.content)["Checks"] if row[1] == "(skipped)"]
    assert [(row[2], row[5], row[7]) for row in skipped] == [("job-gone", "Skipped", "Invoice not found")]
    alone = draft(client, ("job-gone", 1))
    assert alone.status_code == 409 and "Invoice not found" in alone.json()["detail"]
    assert store.job("job-gone") is None


def test_a_stale_revision_is_skipped_and_nothing_qualifying_is_a_409(production):  # noqa: F811
    app, client = production
    store = app.state.store
    store.job("job-1", job())
    shown = client.get("/api/jobs/job-1").json()
    response = draft(client, ("job-1", shown["revision"] + 1))
    assert response.status_code == 409 and "Refresh" in response.json()["detail"]
    assert store.job("job-1")["revision"] == shown["revision"]


def test_a_draft_never_learns_but_the_export_still_does(production, monkeypatch):  # noqa: F811
    app, client = production
    store, learned = app.state.store, app.state.learned
    calls = []
    monkeypatch.setattr(learned, "record_verified", lambda *a, **k: calls.append(a) or {})
    store.job("job-1", job())
    reviewed = confirm(client, client.get("/api/jobs/job-1").json())
    before = files(learned.root)
    assert draft(client, ("job-1", reviewed["revision"])).status_code == 200
    assert client.post("/api/exports/target-draft", headers=H, json={"jobs": [{"id": "job-1", "revision": reviewed["revision"]}]}).status_code == 200
    time.sleep(0.2)  # learning runs on the pool; give a wrong call time to land
    assert files(learned.root) == before and calls == []
    assert client.post("/api/jobs/job-1/export", headers=H, json={"revision": reviewed["revision"]}).status_code == 200
    for _ in range(50):
        if calls:
            break
        time.sleep(0.05)
    assert len(calls) == 1


def test_the_draft_builder_never_writes_a_flagged_or_unevidenced_value():
    view = rules_view(approved_result(status="Review"))
    with pytest.raises(ValueError, match="approved"):
        rules_workbook([view])
    fields, cells = view["fields"], view["lines"][0]["cells"]
    fields["site"] = {**fields["site"], "flagged": True}  # a value the rules flagged
    fields["po"] = {**fields["po"], "evidence": []}  # a value without evidence
    cells["Item"] = {**cells["Item"], "flagged": True}
    fields["location_type"] = {**fields["location_type"], "value": "Somewhere"}  # not a template choice
    content, evidence = rules_workbook([view], draft=True)
    got = sheets(content)
    header = dict(zip(*got["Header"]))
    assert all(fields[k]["value"] for k in ("site", "po")) and cells["Item"]["value"]
    assert header["Supplier Site"] is None and header["Order No"] is None and header["Location Type"] is None
    assert got["Details"][1][1] is None and got["Details"][1][3] is not None  # Item empty, Unit Cost still written
    # Evidence is recorded only for cells that were written.
    assert {"Header!C2", "Header!D2", "Header!F2", "Details!B2"}.isdisjoint(evidence[1]) and "Header!E2" in evidence[1]
