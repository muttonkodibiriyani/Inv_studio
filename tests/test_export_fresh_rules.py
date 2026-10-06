"""An export bumps the job revision; the rules it exported stay current for the draft and the target check. Synthetic data only."""

import io

import pytest
from openpyxl import load_workbook

from tests.test_rules_wiring import H, job, production  # noqa: F401  (pytest fixture)
from tests.test_target_check_api import confirm


def details_items(content):
    return [row[1] for row in load_workbook(io.BytesIO(content))["Details"].iter_rows(min_row=2, values_only=True)]


def draft_items(client, revision):
    response = client.post("/api/jobs/job-1/extraction-draft", headers=H,
                           json={"revision": revision, "acknowledge_unvalidated": True})
    assert response.status_code == 200, response.text
    return details_items(response.content)


def export(client, revision, route):
    if route == "single":
        response = client.post("/api/jobs/job-1/export", headers=H, json={"revision": revision})
    else:
        response = client.post("/api/exports/batch", headers=H, json={"jobs": [{"id": "job-1", "revision": revision}]})
    assert response.status_code == 200, response.text
    return details_items(client.get(response.json()["url"]).content)


@pytest.mark.parametrize("route", ["single", "batch"])
def test_after_an_export_the_draft_item_is_the_export_item_and_the_target_check_answers(production, route):  # noqa: F811
    app, client = production
    app.state.store.job("job-1", job())
    reviewed = confirm(client, client.get("/api/jobs/job-1").json())
    exported = export(client, reviewed["revision"], route)
    stored = app.state.store.job("job-1")
    assert stored["status"] == "exported" and stored["revision"] == reviewed["revision"] + 1
    # Carried forward in the export's own write, not recomputed: the same rules result, one revision on.
    assert stored["rules"]["revision"] == stored["revision"]
    assert stored["rules"]["computed_at"] == reviewed["rules"]["computed_at"]
    assert all(exported) and draft_items(client, stored["revision"]) == exported
    check = client.get("/api/jobs/job-1/target-check")
    assert check.status_code == 200, check.text
    assert check.json()["counts"] == app.state.store.job("job-1")["rules"]["target_check"]["counts"]


def test_rules_already_stale_on_an_exported_job_stay_stale(production):  # noqa: F811
    app, client = production
    app.state.store.job("job-1", job())
    reviewed = confirm(client, client.get("/api/jobs/job-1").json())
    export(client, reviewed["revision"], "single")
    # A job exported before this change: its rules sit one revision behind the export's bump.
    stored = app.state.store.job("job-1")
    stale = stored["revision"] - 1
    app.state.store.job("job-1", {**stored, "rules": {**stored["rules"], "revision": stale}})
    assert client.get("/api/jobs/job-1/target-check").status_code == 409
    assert not any(draft_items(client, stored["revision"]))  # no current rules and no typed item: Item stays empty
    assert app.state.store.job("job-1")["rules"]["revision"] == stale  # nothing is recomputed on an exported job
