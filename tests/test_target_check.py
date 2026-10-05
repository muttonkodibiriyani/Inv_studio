"""Target-sheet accuracy check: synthetic data only (no owner names, ids or amounts)."""

import copy
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

from app import target_check as tc
from app.excel import rules_workbook

ROWS = {
    ("Item Master", 10): {"ITEM_PARENT": "100000001", "SUPPLIER": "200001", "SUPPLIER_NAME": "SYN123 KWD"},
    ("Item Master", 11): {"ITEM_PARENT": "100000002", "SUPPLIER": "200001", "SUPPLIER_NAME": "SYN123 KWD"},
    ("Item Master", 12): {"ITEM_PARENT": "100000001", "SUPPLIER": "200009", "SUPPLIER_NAME": "OTHER"},
    ("PO Extract", 5): {"RMS_ORDER_NO": "300001", "LOCATION": "800001", "EBS_SUPPLIER_CODE": "syn123",
                        "CREATED_DATE": "2026-01-02", "RMS_ITEM_ID": "123456789"},
    ("PO Extract", 6): {"RMS_ORDER_NO": "300001", "LOCATION": "800001", "EBS_SUPPLIER_CODE": "syn123",
                        "CREATED_DATE": "not a date", "RMS_ITEM_ID": "x"},
}
CONFIG = {
    "location_master": {"800001": {"type": "Warehouse (W)", "market": "Synthland", "entity_currency": "KWD"}},
    "supplier_sites": [{"supplier_site": "200001", "currency": "KWD", "status": "Active"}],
    "vat_codes": [{"region": "SYN", "code": "SV5", "rate": "5", "active_from": "2020-01-01"}],
    "buyer_name": "Synthetic Buyer Co",
}
TEXT = "SYNTHETIC INVOICE\nInvoice No SYN-0001\nDate 03/02/2026\nLine 1 qty 5 unit 2.100\nTotal 10.500\nVAT 0.525\nGross 11.025\nTax SV5"


def sources(text=TEXT, config=CONFIG):
    return tc.Sources(lambda sheet, n: ROWS.get((sheet, n)), config, text,
                      site_rows=lambda site: [r for r in ROWS.values() if r.get("SUPPLIER") == site])


def ev(source, reference="", original="", kind="sheet"):
    return {"kind": kind, "source": source, "reference": reference, "original": original, "rule": "R", "confidence": ""}


def field(value, *evidence, reason=""):
    if value is None:
        return {"value": None, "evidence": [], "flagged": True, "reason": reason or "Not found", "label": "x"}
    return {"value": value, "evidence": list(evidence), "flagged": False, "reason": "", "label": "x"}


def line(n, item="100000001", cost="2.100", qty="5", code="SV5", upc=None):
    return {"line": n, "cells": {
        "Item": field(item, ev("Item Master ITEM_PARENT", "Item Master!10")) if item else field(None),
        "UPC": field(upc, ev("Invoice barcode", "page 1", upc or "", "printed")) if upc else field(None),
        "Unit Cost": field(cost, ev("Invoice line", "page 1", cost, "printed")),
        "Quantity": field(qty, ev("Invoice line", "page 1", qty, "printed")),
        "Unit Tax Code": field(code, ev("Owner VAT code table (C/PV), receiving market", "", code, "rule")),
    }}


def view(**over):
    fields = {
        "number": field("SYN-0001", ev("Invoice Document", "", "SYN-0001", "printed")),
        "site": field("200001", ev("Item Master SUPPLIER", "Item Master!10, Item Master!11")),
        "po": field("300001", ev("POGRN RMS_ORDER_NO", "PO Extract!5")),
        "location": field("800001", ev("Accepted POGRN LOCATION", "PO Extract!5")),
        "location_type": field("Warehouse (W)", ev("location master", "", "800001", "rule")),
        "date": field("2026-02-03", ev("Invoice Document Date", "", "03/02/2026", "printed")),
        "currency": field("KWD", ev("Supplier-site table", "200001|v1", "", "table")),
        "net": field("10.500", ev("Invoice printed total", "page 1", "10.500", "printed")),
        "tax": field("0.525", ev("Invoice printed total", "page 1", "0.525", "printed")),
        "taxCode": field("SV5", ev("Owner VAT code table (C/PV), receiving market", "", "SV5", "rule")),
    }
    fields.update(over.pop("fields", {}))
    return {"status": "Approved", "fields": fields, "lines": over.pop("lines", [line(1)]), "issues": [], **over}


def status(result, sheet, column, n=None):
    return next((c["status"], c["sub"]) for c in result["cells"]
                if c["sheet"] == sheet and c["column"] == column and c["line"] == n)


def test_every_cell_gets_exactly_one_status_and_all_verified():
    r = tc.check_view(view(), sources())
    assert {c["status"] for c in r["cells"]} <= set(tc.STATUSES)
    # 13 Header + 3 Tax_Breakdown + 6 Details cells.
    assert r["counts"]["cells"] == 22
    m = r["metric"]
    assert (m["cells"], m["verified"], m["empty_flagged"], m["needs_checking"], m["owner_entry"]) == (15, 14, 1, 0, 0)
    assert sum(m["buckets"].values()) == 15
    assert r["counts"]["empty_flagged"] == 5  # Ref No. 1-3, Comment, UPC (owner rule)
    assert r["summary"] == "Target sheet: 17 verified · 5 empty (flagged) · 0 needs checking"
    assert not r["holds"]
    assert all(k["status"] in (tc.PASS, tc.WARNING) for k in r["checks"]), r["checks"]


def test_check_is_read_only():
    v = view()
    before = copy.deepcopy(v)
    tc.check_view(v, sources())
    assert v == before


def test_mismatch_when_evidence_does_not_hold_the_value():
    v = view(fields={"po": field("399999", ev("POGRN RMS_ORDER_NO", "PO Extract!5"))})
    r = tc.check_view(v, sources())
    assert status(r, "Header", "Order No") == ("mismatch", "")
    assert r["holds"]
    assert [i["code"] for i in tc.issues(r)] == ["Target Mismatch"]
    assert all(i["level"] == "review" for i in tc.issues(r))


def test_missing_source_row_is_mismatch():
    v = view(fields={"po": field("300001", ev("POGRN RMS_ORDER_NO", "PO Extract!99"))})
    assert status(tc.check_view(v, sources()), "Header", "Order No") == ("mismatch", "")


def test_over_cited_data_gap_and_no_evidence_sub_buckets():
    v = view(fields={
        "site": field("200001", ev("Item Master SUPPLIER", "Item Master!10, Item Master!12")),
        "po": field("300001", ev("POGRN RMS_ORDER_NO", "PO Extract!6")),
        "location": field("800001"),
    })
    r = tc.check_view(v, sources())
    assert status(r, "Header", "Supplier Site") == ("over_cited", "")
    assert status(r, "Header", "Order No") == ("data_gap", "")
    assert status(r, "Header", "Location") == ("no_evidence", "")
    b = r["metric"]["buckets"]
    assert (b["data_gap"], b["no_evidence"], b["over_cited"], b["mismatch"], r["metric"]["needs_checking"]) == (1, 1, 1, 0, 3)
    assert [i["code"] for i in tc.issues(r)] == ["Target Unverified"]


def test_empty_cell_keeps_its_reason():
    r = tc.check_view(view(fields={"po": field(None, reason="No evidence")}), sources())
    cell = next(c for c in r["cells"] if c["column"] == "Order No")
    assert (cell["status"], cell["reason"]) == ("empty_flagged", "No evidence")


def test_printed_evidence_needs_text_and_equal_value():
    assert status(tc.check_view(view(), sources(text="")), "Header", "Document") == ("unverifiable", "")
    reason = next(c["reason"] for c in tc.check_view(view(), sources(text=""))["cells"] if c["column"] == "Document")
    assert reason == tc.SCAN_REASON  # AI read of a scan: no text layer and no box
    boxed = view(fields={"number": field("SYN-0001", ev("Invoice (ai)", "page 1 box [1, 2, 3, 4]", "SYN-0001", "printed"))})
    cell = next(c for c in tc.check_view(boxed, sources(text=""))["cells"] if c["column"] == "Document")
    assert (cell["status"], cell["reason"]) == ("unverifiable", "evidence not re-checkable")
    v = view(fields={"number": field("SYN-0002", ev("Invoice Document", "", "SYN-0001", "printed"))})
    assert status(tc.check_view(v, sources()), "Header", "Document") == ("mismatch", "")
    v = view(fields={"number": field("SYN-0009", ev("Invoice Document", "", "SYN-0009", "printed"))})
    assert status(tc.check_view(v, sources()), "Header", "Document") == ("mismatch", "")


def test_location_type_master_and_prefix_fallback():
    v = view(fields={"location_type": field("Store (S)", ev("location master", "", "800001", "rule"))})
    assert status(tc.check_view(v, sources()), "Header", "Location Type") == ("mismatch", "")
    v = view(fields={"location_type": field("Warehouse (W)", ev("prefix fallback", "", "800001", "rule"))})
    assert status(tc.check_view(v, sources()), "Header", "Location Type") == ("unverifiable", "")


def test_site_derivation_from_order_ebs_code_and_location_entity():
    derived = ev("Item Master SUPPLIER via SUPPLIER SITES, " + tc.SITE_DERIVATION, "PO Extract!5")
    r = tc.check_view(view(fields={"site": field("200001", derived)}), sources())
    assert status(r, "Header", "Supplier Site") == ("verified", "")
    config = {**CONFIG, "supplier_sites": []}
    r = tc.check_view(view(fields={"site": field("200001", derived)}), sources(config=config))
    assert status(r, "Header", "Supplier Site") == ("mismatch", "")
    shifted = ev("x " + tc.SITE_DERIVATION, "PO Extract!6")
    r = tc.check_view(view(fields={"site": field("200001", shifted)}), sources())
    assert status(r, "Header", "Supplier Site") == ("data_gap", "")


def test_reviewed_invoice_tax_code_is_checked_as_printed():
    printed = ev("Reviewed invoice tax code", "", "SV5", "rule")
    v = view(fields={"taxCode": field("SV5", printed)})
    assert status(tc.check_view(v, sources()), "Tax_Breakdown", "Tax Code") == ("verified", "")
    assert status(tc.check_view(v, sources(text="no code here")), "Tax_Breakdown", "Tax Code") == ("mismatch", "")


def test_upc_is_empty_by_owner_rule_and_filled_upc_is_a_violation():
    v = view(lines=[line(1, upc="1234567890123")])
    assert status(tc.check_view(v, sources()), "Details", "UPC", 1) == ("empty_flagged", "")
    r = tc.check_view(v, sources(), upc="barcode")
    assert status(r, "Details", "UPC", 1) == ("owner_rule_violation", "")


def test_owner_entry_needs_equal_stored_value_actor_and_time():
    entered = {"kind": "owner_entry", "source": "Entered by reviewer", "reference": "review", "original": ""}
    v = view(fields={"po": field("300002", entered)})
    entries = {"header": {"po": "300002"}, "lines": {}}
    who = {"header:po": {"actor": "reviewer", "at": "2026-02-03T10:00:00+00:00"}}
    r = tc.check_view(v, sources(), entries=entries, attribution=who)
    assert status(r, "Header", "Order No") == ("verified", "owner_entry")
    assert r["metric"]["owner_entry"] == 1 and r["metric"]["verified"] == 14
    r = tc.check_view(v, sources(), entries=entries)
    assert status(r, "Header", "Order No") == ("owner_entry_unattributed", "")
    r = tc.check_view(v, sources(), entries={"header": {"po": "300003"}}, attribution=who)
    assert status(r, "Header", "Order No") == ("owner_entry_unattributed", "")


def check(result, name):
    return next(k for k in result["checks"] if k["check"] == name)


def test_arithmetic_at_currency_decimals():
    r = tc.check_view(view(), sources())
    assert check(r, "lines_to_net")["status"] == "pass" and check(r, "lines_to_net")["decimals"] == 3
    assert check(r, "net_plus_tax_gross")["status"] == "pass"
    assert check(r, "tax_breakdown_to_header")["status"] == "pass"
    off = view(lines=[line(1, cost="2.101")])
    r = tc.check_view(off, sources())
    assert check(r, "lines_to_net")["status"] == "fail" and r["holds"]
    assert any(i["rule"] == "TARGET-LINES_TO_NET" for i in tc.issues(r))
    wrong_tax = view(fields={"tax": field("0.600", ev("Invoice printed total", "page 1", "0.600", "printed"))})
    r = tc.check_view(wrong_tax, sources(text=TEXT + " 0.600"))
    assert check(r, "tax_breakdown_to_header")["status"] == "fail"
    assert check(r, "net_plus_tax_gross")["status"] == "warning"  # 11.100 is not printed; not blocking
    assert check(tc.check_view(view(), sources(config={})), "tax_breakdown_to_header")["status"] == "skipped"


def test_two_decimal_currency_rounds_each_way():
    lines = [line(1, cost="0.333", qty="3"), line(2, cost="0.333", qty="3")]
    v = view(fields={"currency": field("SYD", ev("Supplier-site table", "200001|v1", "", "table")),
                     "net": field("2.00", ev("Invoice printed total", "page 1", "2.00", "printed"))}, lines=lines)
    assert check(tc.check_view(v, sources()), "lines_to_net")["status"] == "pass"


def test_tax_rounded_per_line_passes():
    lines = [line(1, cost="0.10", qty="1"), line(2, cost="0.10", qty="1")]
    sd = ev("Supplier-site table", "200001|v1", "", "table")
    v = view(fields={"currency": field("SYD", sd), "net": field("0.20", ev("Invoice printed total", "page 1", "0.20", "printed")),
                     "tax": field("0.02", ev("Invoice printed total", "page 1", "0.02", "printed"))}, lines=lines)
    assert check(tc.check_view(v, sources()), "tax_breakdown_to_header")["status"] == "pass"  # 0.01 + 0.01 per line
    v["fields"]["tax"] = field("0.03", ev("Invoice printed total", "page 1", "0.03", "printed"))
    assert check(tc.check_view(v, sources()), "tax_breakdown_to_header")["status"] == "fail"


def test_joins_find_orphans_and_duplicates():
    t = tc.TEMPLATE
    good = {"Header": [t["Header"], [1] + [None] * 12], "Tax_Breakdown": [t["Tax_Breakdown"], [1, "C", 1]],
            "Details": [t["Details"], [1, None, None, 1, 1, "C"]]}
    assert tc.joins(good)[0]["status"] == "pass"
    bad = copy.deepcopy(good)
    bad["Details"].append([2, None, None, 1, 1, "C"])
    bad["Header"].append([1] + [None] * 12)
    bad["Tax_Breakdown"].append([1, "C", 1])
    detail = tc.joins(bad)[0]
    assert detail["status"] == "fail"
    assert "repeat in Header" in detail["detail"] and "do not join" in detail["detail"]
    assert "more than one Tax_Breakdown" in detail["detail"]


def test_template_order_types_and_decimals():
    r = tc.check_view(view(), sources())
    assert check(r, "template")["status"] == "pass"
    sheets = tc.planned_rows(view())
    sheets["Header"][0] = list(reversed(sheets["Header"][0]))
    sheets["Header"][1][9] = "note"
    sheets["Tax_Breakdown"][1][2] = tc.Decimal("10.5001")
    sheets["Details"][1][4] = "five"
    detail = tc.template(sheets, 3)[0]["detail"]
    assert "Header columns differ" in detail and "Ref No. 1 must stay empty" in detail
    assert "more decimals than the currency" in detail and "Details Quantity is not a number" in detail


def test_checks_sheet_lists_every_cell_and_keeps_template_tabs():
    v = view()
    content, _ = rules_workbook([v], "empty")
    r = tc.check_view(v, sources())
    out, book_checks = tc.add_checks_sheet(content, [r], ["KWD"])
    assert [k["status"] for k in book_checks] == ["pass", "pass"]
    book = load_workbook(__import__("io").BytesIO(out))
    assert book.sheetnames == ["Header", "Tax_Breakdown", "Details", "Checks"]
    rows = list(book["Checks"].iter_rows(values_only=True))
    assert list(rows[0]) == tc.CHECKS_COLUMNS
    cell_rows = [x for x in rows[1:] if x[1] in ("Header", "Tax_Breakdown", "Details")]
    assert len(cell_rows) == r["counts"]["cells"]
    assert {x[5] for x in cell_rows} <= set(tc.GROUPS) and {x[6] for x in cell_rows} <= set(tc.STATUSES)
    # The written workbook agrees with the planned rows the per-invoice check used.
    written = tc.read_workbook(content)
    planned = tc.planned_rows(v)
    for name in tc.TEMPLATE:
        assert [list(map(lambda x: None if x == "" else x, row)) for row in written[name]][0] == planned[name][0]
        assert len(written[name]) == len(planned[name])


def test_private_buyer_name_never_reaches_the_check_output():
    r = tc.check_view(view(), sources())
    assert CONFIG["buyer_name"] not in repr(r) and CONFIG["buyer_name"] not in repr(tc.log_fields(r))


def test_log_fields_hold_no_values():
    v = view(fields={"po": field("399999", ev("POGRN RMS_ORDER_NO", "PO Extract!5"))})
    logged = repr(tc.log_fields(tc.check_view(v, sources())))
    assert "399999" not in logged and "SYN-0001" not in logged and "Header.Order No" in logged


def test_confirm_records_and_accuracy_summary():
    now = datetime(2026, 3, 1, tzinfo=timezone.utc)
    system = tc.check_view(view(), sources())
    entered = {"kind": "owner_entry", "source": "Entered by reviewer", "reference": "review", "original": ""}
    final = tc.check_view(view(fields={"po": field("300002", entered)}), sources())
    records = tc.confirm_records(system, final, supplier="200001", at=now - timedelta(days=10))
    assert len(records) == 15
    changed = [r for r in records if r["changed"]]
    assert [(r["field"], r["status_before"]) for r in changed] == [("Header.Order No", "verified")]
    assert all(set(r) == {"at", "supplier", "field", "line", "status_before", "changed"} for r in records)
    assert "300002" not in repr(records) and "300001" not in repr(records)
    old = tc.confirm_records(system, system, supplier="200002", at=now - timedelta(days=2))
    s = tc.accuracy_summary(records + old, now=now)
    assert s["periods"]["7d"]["overall"] == {"cells": 15, "unchanged": 15, "changed": 0, "accuracy": 1.0}
    assert s["periods"]["30d"]["overall"]["changed"] == 1 and s["periods"]["all"]["overall"]["cells"] == 30
    assert s["periods"]["all"]["fields"]["Header.Order No"] == {"cells": 2, "unchanged": 1, "changed": 1,
                                                                 "accuracy": 0.5}
    assert s["periods"]["all"]["suppliers"]["200001"]["changed"] == 1
    assert s["periods"]["all"]["by_status_before"]["verified"]["changed"] == 1
    only = tc.accuracy_summary(records + old, now=now, supplier="200002")
    assert only["periods"]["all"]["overall"]["cells"] == 15


def test_selected_order_no_is_checked_on_rms_order_no():
    sel = ("Selected by POG-001 among 2 order/location candidates under the supplier's 6-character EBS code SYN001: "
           "the only one whose quantity and value agree; not printed on the invoice")
    v = view(fields={"po": field("300001", ev(sel, "PO Extract!5"))})
    cell = next(c for c in tc.check_view(v, sources())["cells"] if c["column"] == "Order No")
    assert cell["status"] == "verified" and "not printed" in cell["reason"]
    v = view(fields={"po": field("800001", ev(sel, "PO Extract!5"))})  # held by the row, but as LOCATION
    assert status(tc.check_view(v, sources()), "Header", "Order No") == ("mismatch", "")
