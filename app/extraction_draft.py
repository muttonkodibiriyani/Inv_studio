"""An explicitly unvalidated workbook of source facts, with unknown codes blank."""
import io
from datetime import date
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from .excel import HEADERS, _id_cell, text_cell

DISCLOSURE = 'EXTRACTION_REVIEW_ONLY: Unvalidated source values. Blank fields need completion. Not approved for posting or payment; no receipts reserved.'


def rules_item(rules, n):
    """The rules' Item (Item Master ITEM_PARENT) for invoice line n, only when it has a value and evidence."""
    for row in (rules or {}).get("lines") or []:
        if row.get("line")!=n:continue
        cell=(row.get("cells") or {}).get("Item") or {}
        if cell.get("value") and cell.get("evidence") and not cell.get("flagged"):return cell
    return None


def extraction_batch_workbook(entries):
    book=Workbook();book.remove(book.active)
    book.properties.title='Invoice extraction — review copy'
    book.properties.description=DISCLOSURE
    for name,headers in HEADERS.items():
        sheet=book.create_sheet(name);sheet.append(headers);sheet.freeze_panes='A2'
        for cell in sheet[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='78521D')
            sheet.column_dimensions[cell.column_letter].width=max(18,len(cell.value)+2)
            cell.comment=Comment(DISCLOSURE,'Invoice Studio')
    for transaction,(invoice,source_name,revision,*rest) in enumerate(entries,1):
        rules=rest[0] if rest else None
        header=book['Header'];header_row=transaction+1
        try: invoice_date=date.fromisoformat(invoice.date or '')
        except ValueError: invoice_date=invoice.date
        values=[transaction,invoice.number,invoice.site,invoice.po,invoice.location,None,invoice_date,invoice.net,invoice.tax,'','','',DISCLOSURE]
        for col,value in enumerate(values,1):
            if col in (2,3,4,5,6,10,11,12,13) or isinstance(value,str):text_cell(header.cell(header_row,col),value)
            else:header.cell(header_row,col,value)
        header.cell(header_row,7).number_format='yyyy-mm-dd'
        header.cell(header_row,2).comment=Comment(f'Source: {source_name}\nReviewed snapshot revision: {revision}\n'+DISCLOSURE,'Invoice Studio')
        tax=book['Tax_Breakdown'];tax.append([transaction,None,invoice.net]);text_cell(tax.cell(transaction+1,2),invoice.taxCode)
        details=book['Details']
        for n,line in enumerate(invoice.lines,1):
            row=details.max_row+1
            # With the rules of this revision, Item is what the target workbook writes (decision 56); without them, only a typed item.
            matched=rules_item(rules,n) if rules else None
            item=(matched or {}).get("value") if rules else line.item_id
            details.append([transaction,None,None,line.price,line.qty,None])
            # A printed supplier SKU is not a proven internal target item ID.
            # A rules Item is written exactly as the target workbook writes it.
            (_id_cell if matched else text_cell)(details.cell(row,2),item)
            text_cell(details.cell(row,3),line.gtin)
            text_cell(details.cell(row,6),invoice.taxCode)
            details.cell(row,3).number_format='@'
            details.cell(row,4).number_format='0.0000';details.cell(row,5).number_format='0.0000'
            notes=[f'Supplier item code: {line.sku or "not read"}',f'Description: {line.description or "not read"}',f'Source page: {line.page or "not recorded"}']
            for field,label in (('net_amount','Printed line net'),('tax_amount','Printed line tax')):
                value=getattr(line,field,None)
                if value is not None:notes.append(f'{label}: {value}')
            if line.evidence:notes.append('Evidence: '+line.evidence)
            if matched:
                source=matched["evidence"][0]
                notes.append('Rules match, not validated: '+(source.get("source") or "rules")+(f' ({source["rule"]})' if source.get("rule") else '')+'.')
                # The rules' own review flags on this line travel with the Item they qualify.
                notes+=[f'Review flag {i.get("rule") or i.get("code")}: {i.get("message")}' for i in (rules or {}).get("issues") or [] if i.get("line")==n]
            elif item:notes.append('Internal item was supplied during review; business validation is still required.')
            else:notes.append('Internal item mapping requires confirmation.')
            if rules and line.item_id and line.item_id!=item:notes.append('Item typed in review is a review candidate only for the rules: '+line.item_id)
            details.cell(row,2).comment=Comment('\n'.join(notes),'Invoice Studio')
    output=io.BytesIO();book.save(output);return output.getvalue()


def extraction_workbook(invoice, source_name, revision, rules=None):
    """``rules`` is the job's rules view for this revision: Item is the rules' evidenced Item; a typed item only when there are no rules."""
    return extraction_batch_workbook([(invoice,source_name,revision,rules)])
