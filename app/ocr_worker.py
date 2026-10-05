"""Isolated OCR process. No API credentials are passed to this worker."""
import argparse
import gc
import json
import math
import multiprocessing
import os
import re
import tempfile
import time
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path


PADDLE_DETECTION_MODEL = "PP-OCRv5_mobile_det"
PADDLE_RECOGNITION_MODELS = {
    "en": "PP-OCRv5_mobile_rec",
    "ch": "PP-OCRv5_mobile_rec",
    "ar": "arabic_PP-OCRv5_mobile_rec",
    "fr": "latin_PP-OCRv5_mobile_rec",
    "de": "latin_PP-OCRv5_mobile_rec",
}
PADDLE_MAX_SIDE = 1600
PADDLE_RECOVERY_MAX_SIDE = 2400
PADDLE_RECOVERY_RENDER_SCALE = 3
PADDLE_RECOVERY_MAX_PAGES = 2
PADDLE_RECOVERY_BASELINE_LIMIT_SECONDS = 150
PADDLE_WORKER_BUDGET_SECONDS = 360


def paddle_model_config(language):
    try:
        recognition = PADDLE_RECOGNITION_MODELS[language]
    except KeyError:
        raise ValueError("Unsupported OCR language") from None
    return {
        "text_detection_model_name": PADDLE_DETECTION_MODEL,
        "text_recognition_model_name": recognition,
        "text_det_limit_side_len": PADDLE_MAX_SIDE,
        "text_det_limit_type": "max",
    }


def spatial_text(boxes):
    """Rebuild reading rows while leaving raw OCR boxes unchanged."""
    pages = {}
    for index, entry in enumerate(boxes):
        text = str(entry.get("text", "")).strip()
        if not text:
            continue
        page = entry.get("page", 1)
        box = entry.get("box")
        positioned = None
        if isinstance(box, (list, tuple)) and len(box) == 4 and all(
            isinstance(value, (int, float)) and math.isfinite(value) for value in box
        ):
            left, top, right, bottom = (float(value) for value in box)
            if right >= left and bottom >= top:
                positioned = {"text": text, "left": left, "top": top, "bottom": bottom,
                              "center": (top + bottom) / 2, "height": max(1.0, bottom - top)}
        pages.setdefault(page, {"positioned": [], "loose": []})[
            "positioned" if positioned else "loose"
        ].append(positioned or (index, text))

    output = []
    for page in sorted(pages, key=lambda value: (not isinstance(value, int), str(value))):
        rows = []
        for item in sorted(pages[page]["positioned"], key=lambda value: (value["center"], value["left"])):
            best = None
            for row in rows:
                center = sum(row["centers"]) / len(row["centers"])
                height = sorted(row["heights"])[len(row["heights"]) // 2]
                distance = abs(item["center"] - center)
                overlap = min(item["bottom"], row["bottom"]) - max(item["top"], row["top"])
                aligned = distance <= max(3.0, 0.55 * max(item["height"], height))
                if aligned or overlap >= 0.25 * min(item["height"], height):
                    if best is None or distance < best[0]:
                        best = (distance, row)
            if best is None:
                rows.append({"items": [item], "centers": [item["center"]],
                             "heights": [item["height"]], "top": item["top"],
                             "bottom": item["bottom"]})
            else:
                row = best[1]
                row["items"].append(item)
                row["centers"].append(item["center"])
                row["heights"].append(item["height"])
                row["top"] = min(row["top"], item["top"])
                row["bottom"] = max(row["bottom"], item["bottom"])
        rows.sort(key=lambda row: sum(row["centers"]) / len(row["centers"]))
        output.extend("  ".join(item["text"] for item in sorted(row["items"], key=lambda x: x["left"]))
                      for row in rows)
        output.extend(text for _, text in sorted(pages[page]["loose"]))
    return "\n".join(output)


def paddle_reader(language, max_side=None):
    # Cached official model names skip mutable defaults and network discovery.
    os.environ.setdefault("PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK", "True")
    from paddleocr import PaddleOCR
    config=paddle_model_config(language)
    if max_side is not None:
        config["text_det_limit_side_len"]=max_side
    return PaddleOCR(**config,use_doc_orientation_classify=False,
                     use_doc_unwarping=False,use_textline_orientation=False,
                     text_recognition_batch_size=16,device="cpu",enable_mkldnn=False,cpu_threads=2)


def templates_from(paths):
    from invoice2data.extract.loader import read_templates
    return [t for path in paths if Path(path).is_dir() for t in read_templates(str(path))]


def _learned_meta(result, templates):
    # A learned template (app.learned) carries its provenance and row conventions in a 'learned' block.
    name=result.get("template_name")
    for template in templates:
        if template.get("template_name")==name:
            return template.get("learned") if isinstance(template.get("learned"),dict) else None
    return None


def template_extract(text, templates, with_meta=False):
    # Reuse invoice2data's parser after a local OCR/layout engine has read text.
    from invoice2data import extract_data
    class TextReader:
        @staticmethod
        def to_text(_): return text
    result=extract_data("ocr-text", templates=templates, input_module=TextReader, ai_fallback=False)
    if not result: return (None,None) if with_meta else None
    meta=_learned_meta(result,templates)
    fields={"invoice_number":"number","amount":"net"}
    allowed={"number","supplier_name","buyer_name","seller","site","buyer","po","location","date_printed","date","currency","origin","market","taxCode","net","tax","lines"}
    out={fields.get(k,k):v for k,v in result.items() if fields.get(k,k) in allowed}
    if meta is not None:
        # A learned regex that matches more than once yields a list: one value when they agree, else nothing.
        for key in [k for k,v in out.items() if k!="lines" and isinstance(v,list)]:
            if len({str(v) for v in out[key]})==1:out[key]=out[key][0]
            else:del out[key]
        for field,candidates in (meta.get("fallbacks") or {}).items():
            # The learner ranked further anchors for this field; the first that finds one value wins.
            if out.get(field) not in (None,""):continue
            for candidate in candidates:
                try:found={m for m in re.findall(candidate["regex"],text)}
                except re.error:continue
                if len(found)!=1:continue
                value=found.pop()
                if field=="date":
                    try:value=datetime.strptime(value,candidate.get("date_format") or "%Y-%m-%d").strftime("%Y-%m-%d")
                    except ValueError:continue
                elif field in ("net","tax"):
                    try:value=float(value.replace(",",""))
                    except ValueError:continue
                out[field]=value
                break
    if isinstance(out.get("date"),(date,datetime)):out["date"]=out["date"].strftime("%Y-%m-%d")
    from .models import Line
    lines=[]
    for row in out.get("lines",[]):
        row=dict(row)
        if "description" not in row and "name" in row:row["description"]=row["name"]
        if meta is not None:
            from .learned import finish_row
            row=finish_row(row,meta)
        lines.append({k:v for k,v in row.items() if k in Line.model_fields})
    out["lines"]=lines
    if meta is not None:
        for field in ("net","tax"):
            if isinstance(out.get(field),float):
                value=Decimal(str(out[field]))
                places=(meta.get("decimals") or {}).get(field)  # keep the decimals the supplier prints
                if places is not None:value=value.quantize(Decimal(1).scaleb(-int(places)))
                out[field]=format(value,"f")
    return (out,meta) if with_meta else out


LEARNED_HEADER_FIELDS=("number","date","currency","net","tax","po")


def _same_arithmetic(baseline,candidate):
    left=(baseline or {}).get("lines") or [];right=(candidate or {}).get("lines") or []
    if len(left)!=len(right):return False
    for a,b in zip(left,right):
        for field in ("qty","price","net_amount"):
            if _visible_decimal(a.get(field))!=_visible_decimal(b.get(field)):return False
    return True


def _document_reconciles(invoice):
    lines=(invoice or {}).get("lines") or []
    if not lines:return False
    net=_visible_decimal((invoice or {}).get("net"))
    amounts=[_visible_decimal(line.get("net_amount")) for line in lines]
    if net is not None:
        return None not in amounts and sum(amounts)==net
    checks=_line_reconciliation(invoice)
    return all(check is True for check in checks)


def apply_learned(text,learned_templates,baseline,method):
    """Learned supplier templates stay behind the built-in readers.

    They read alone only when the built-in reading has no lines; when both agree on every
    row's qty, price and amount they may overlay only the fields the template was proven to
    fix on its verified source; otherwise they replace the built-in reading only when the
    learned one reconciles with the document and the built-in one does not.
    """
    if not text or not learned_templates:return baseline,method,None
    try:candidate,meta=template_extract(text,learned_templates,with_meta=True)
    except Exception as error:return baseline,method,{"mode":"error","error":type(error).__name__}
    if candidate is None or meta is None:return baseline,method,None
    note={"template":meta.get("supplier_key"),"source_id":meta.get("source_id")}
    if not (baseline or {}).get("lines"):
        return candidate,"learned_template",dict(note,mode="read")
    if _same_arithmetic(baseline,candidate):
        merged=json.loads(json.dumps(baseline));changed=[]
        for field in meta.get("fixes") or []:
            if field in LEARNED_HEADER_FIELDS:
                if candidate.get(field) not in (None,"") and str(candidate[field])!=str(merged.get(field)):
                    merged[field]=candidate[field];changed.append(field)
            elif field!="rows":
                for row,learned_row in zip(merged["lines"],candidate["lines"]):
                    value=learned_row.get(field)
                    if value not in (None,"") and str(value)!=str(row.get(field)):
                        row[field]=value
                        if field not in changed:changed.append(field)
        if changed:return merged,method+"+learned",dict(note,mode="overlay",fields=changed)
        return baseline,method,None
    if _document_reconciles(candidate) and not _document_reconciles(baseline):
        return candidate,"learned_template",dict(note,mode="replace")
    return baseline,method,None


def structured_extract(text, boxes, templates, tables=None):
    parsed = template_extract(text, templates) if text else None
    if parsed is not None:
        return parsed, "template"
    if boxes:
        # A recognised columnar 'TAX INVOICE' layout is read by its own geometry for every engine.
        from .columnar_tax_invoice import extract as columnar_extract
        columnar = columnar_extract(text, boxes)
        if columnar is not None:
            from .models import Invoice
            return Invoice.model_validate(columnar[0]).model_dump(mode="json"), "columnar_tax_invoice"
    if tables or boxes:
        from .docling_extract import extract_invoice_from_tables
        candidate=extract_invoice_from_tables(text,tables or [],boxes=boxes)
        if candidate is not None:
            # Complement table rows only with explicit header facts from this
            # same reader's text/geometry. Never borrow a different engine's rows.
            from .layout_extract import extract_invoice
            header=extract_invoice(text,boxes) or {}
            fields=("number","po","date_printed","date","currency","net","tax","supplier_name","buyer_name")
            compatible=all(candidate.get(key) in (None,"") or header.get(key) in (None,"")
                           or str(candidate[key])==str(header[key]) for key in ("number","po"))
            if compatible:
                for key in fields:
                    if candidate.get(key) in (None,"") and header.get(key) not in (None,""):
                        candidate[key]=header[key]
            from .models import Invoice
            return Invoice.model_validate(candidate).model_dump(mode="json"), "table"
    if text:
        from .layout_extract import extract_invoice
        candidate = extract_invoice(text, boxes)
        if candidate is not None:
            from .models import Invoice
            return Invoice.model_validate(candidate).model_dump(mode="json"), "layout"
    return None, "text_only"


def digital(path, with_tables=False):
    import pdfplumber
    suffix=path.suffix.lower()
    if suffix in (".txt",".csv",".json"):return path.read_text(encoding="utf-8-sig"),[]
    if suffix==".docx":
        import zipfile
        from defusedxml import ElementTree
        with zipfile.ZipFile(path) as z:
            root=ElementTree.fromstring(z.read("word/document.xml"))
        ns={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        return "\n".join("".join(t.text or "" for t in p.findall(".//w:t",ns)) for p in root.findall(".//w:p",ns)),[]
    if suffix==".xlsx":
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True);chunks=[]
        for sheet in book:
            if sheet.max_row>10000 or sheet.max_column>60:raise ValueError("Invoice spreadsheet is too large")
            chunks.append(sheet.title)
            chunks.extend(" | ".join(str(x) if x is not None else "" for x in row) for row in sheet.iter_rows(values_only=True))
        book.close();return "\n".join(chunks),[]
    if suffix!=".pdf":return "",[]
    pages=[]; boxes=[];tables=[]
    with pdfplumber.open(path) as pdf:
        for n,page in enumerate(pdf.pages,1):
            pages.append(page.extract_text() or "")
            if with_tables:
                tables.extend({"page":n,"rows":rows,"source":"pdfplumber","size":[page.width,page.height]} for rows in page.extract_tables() if rows)
            for word in page.extract_words()[:4000]:
                boxes.append({"text":word["text"],"page":n,"box":[word["x0"],word["top"],word["x1"],word["bottom"]],"size":[page.width,page.height]})
    return ("\n".join(pages),boxes,tables) if with_tables else ("\n".join(pages),boxes)


def paddle(path,language,*,render_scale=2,max_side=None):
    import pypdfium2 as pdfium
    import numpy as np
    from PIL import Image, ImageOps, ImageSequence
    reader=paddle_reader(language,max_side=max_side)
    boxes=[];line_boxes=[]
    with tempfile.TemporaryDirectory() as temp:
        pages=[]
        if path.suffix.lower()==".pdf":
            doc=pdfium.PdfDocument(path)
            for n,page in enumerate(doc):
                image=page.render(scale=render_scale).to_pil().convert("RGB")
                target=Path(temp)/f"{n}.png";image.save(target);pages.append(target)
            doc.close()
        else:
            with Image.open(path) as original:
                for n,frame in enumerate(ImageSequence.Iterator(original)):
                    if n>=20:raise ValueError("Too many image frames")
                    target=Path(temp)/f"{n}.png"
                    ImageOps.exif_transpose(frame).convert("RGB").save(target);pages.append(target)
        for n,file in enumerate(pages,1):
            with Image.open(file) as original:
                image=ImageOps.exif_transpose(original).convert("RGB")
                image_limit=3600 if render_scale>2 else 2600
                image.thumbnail((image_limit,image_limit)); size=list(image.size)
                for result in reader.predict(np.array(image),return_word_box=True):
                    res=result.json
                    if isinstance(res,str):res=json.loads(res)
                    res=res.get("res",res)
                    texts=res.get("rec_texts",[]);scores=res.get("rec_scores",[]);polys=res.get("rec_polys",[])
                    for idx,value in enumerate(texts):
                        poly=polys[idx] if idx<len(polys) else []
                        box=[min(p[0] for p in poly),min(p[1] for p in poly),max(p[0] for p in poly),max(p[1] for p in poly)] if len(poly) else None
                        line_box={"text":value,"page":n,"confidence":float(scores[idx]) if idx<len(scores) else None,"box":box,"size":size}
                        line_boxes.append(line_box)
                        word_texts=res.get("text_word",[])
                        word_regions=res.get("text_word_boxes",[])
                        if idx<len(word_texts) and idx<len(word_regions) and len(word_texts[idx])==len(word_regions[idx]) and word_texts[idx]:
                            for word,region in zip(word_texts[idx],word_regions[idx]):
                                if len(region)!=4:continue
                                boxes.append({"text":word,"page":n,"confidence":float(scores[idx]) if idx<len(scores) else None,
                                    "box":[float(v) for v in region],"size":size,"geometry":"word","parent_text":value})
                            continue
                        boxes.append(line_box)
    # Keep real text-line strings for templates; word geometry serves tables
    # without adding artificial spaces inside header labels or numeric fields.
    text=spatial_text(line_boxes)
    # The optional recovery pass runs in a fresh child. Release the baseline
    # reader before spawning it so both model graphs are not resident together.
    del reader
    gc.collect()
    return text,boxes


def _visible_decimal(value):
    try:
        parsed=Decimal(str(value).strip().replace(",",""))
    except (InvalidOperation,ValueError):
        return None
    return parsed if parsed.is_finite() else None


def _same_visible_value(field,left,right):
    if left in (None,"") or right in (None,""):
        return left in (None,"") and right in (None,"")
    if field in {"net","tax"}:
        left_decimal=_visible_decimal(left);right_decimal=_visible_decimal(right)
        return left_decimal is not None and left_decimal==right_decimal
    if field in {"number","po"}:
        normal=lambda value:"".join(char for char in str(value).casefold() if char.isalnum())
    else:
        normal=lambda value:" ".join(str(value).casefold().split())
    return normal(left)==normal(right)


def _missing_quantity_price(candidate):
    return sum(
        value in (None,"")
        for line in (candidate or {}).get("lines",[])
        for value in (line.get("qty"),line.get("price"))
    )


def _line_reconciliation(candidate):
    result=[]
    for line in (candidate or {}).get("lines",[]):
        qty=_visible_decimal(line.get("qty"));price=_visible_decimal(line.get("price"))
        printed=_visible_decimal(line.get("net_amount"))
        result.append(None if None in (qty,price,printed) else qty*price==printed)
    return result


def paddle_recovery_criteria(baseline,alternate):
    baseline_lines=(baseline or {}).get("lines",[])
    alternate_lines=(alternate or {}).get("lines",[])
    preserved_headers=all(
        baseline.get(field) in (None,"")
        or _same_visible_value(field,baseline.get(field),alternate.get(field))
        for field in ("number","date","currency","po")
    )
    totals_equal=all(
        _same_visible_value(field,baseline.get(field),alternate.get(field))
        for field in ("net","tax")
    )
    baseline_reconciliation=_line_reconciliation(baseline)
    alternate_reconciliation=_line_reconciliation(alternate)
    reconciliation_no_worse=(
        alternate_reconciliation.count(True)>=baseline_reconciliation.count(True)
        and alternate_reconciliation.count(False)<=baseline_reconciliation.count(False)
        and all(
            status is not True
            or (index<len(alternate_reconciliation) and alternate_reconciliation[index] is True)
            for index,status in enumerate(baseline_reconciliation)
        )
    )
    criteria={
        "row_count_not_lower":len(alternate_lines)>=len(baseline_lines),
        "preserved_headers":preserved_headers,
        "totals_equal":totals_equal,
        "missing_qty_price_strictly_lower":(
            _missing_quantity_price(alternate)<_missing_quantity_price(baseline)
        ),
        "line_net_reconciliation_no_worse":reconciliation_no_worse,
    }
    return criteria,all(criteria.values())


def _pdf_page_count(path):
    if path.suffix.lower()!=".pdf":return None
    import pypdfium2 as pdfium
    document=pdfium.PdfDocument(path)
    try:return len(document)
    finally:document.close()


def _paddle_recovery_child(path,language,output):
    """Run the optional expensive pass outside the baseline worker process."""
    started=time.monotonic()
    try:
        text,boxes=paddle(
            Path(path),language,render_scale=PADDLE_RECOVERY_RENDER_SCALE,
            max_side=PADDLE_RECOVERY_MAX_SIDE,
        )
        payload={"status":"completed","seconds":time.monotonic()-started,
                 "text":text,"boxes":boxes}
    except Exception as error:
        payload={"status":"failed","seconds":time.monotonic()-started,
                 "error":type(error).__name__}
    target=Path(output)
    target.write_text(json.dumps(payload),encoding="utf-8")
    target.chmod(0o600)


def _run_paddle_recovery_pass(path,language,timeout):
    """Return one optional OCR pass while enforcing a hard child deadline."""
    with tempfile.TemporaryDirectory() as temp:
        output=Path(temp)/"recovery.json"
        context=multiprocessing.get_context("spawn")
        process=context.Process(
            target=_paddle_recovery_child,
            args=(str(path),language,str(output)),
            daemon=True,
        )
        started=time.monotonic();process.start();process.join(timeout)
        elapsed=time.monotonic()-started
        if process.is_alive():
            process.terminate();process.join(5)
            if process.is_alive():process.kill();process.join(2)
            return {"status":"timed_out","seconds":elapsed}
        if process.exitcode!=0 or not output.exists():
            return {"status":"failed","seconds":elapsed}
        try:return json.loads(output.read_text(encoding="utf-8"))
        except (OSError,json.JSONDecodeError):
            return {"status":"failed","seconds":elapsed}


def maybe_recover_paddle(
    path,language,templates,baseline_text,baseline_boxes,baseline_candidate,
    baseline_method,baseline_seconds,worker_elapsed,
    worker_budget=PADDLE_WORKER_BUDGET_SECONDS,
):
    """Try one higher-resolution pass and retain it only on strict improvement."""
    missing=_missing_quantity_price(baseline_candidate)
    lines=(baseline_candidate or {}).get("lines",[])
    page_count=_pdf_page_count(path)
    estimated_alternate=max(15.0,baseline_seconds*1.35)
    remaining=max(0.0,min(PADDLE_WORKER_BUDGET_SECONDS,worker_budget)-worker_elapsed)
    eligibility={
        "has_lines":bool(lines),
        "missing_qty_or_price":missing>0,
        "pdf_at_most_two_pages":page_count is not None and page_count<=PADDLE_RECOVERY_MAX_PAGES,
        "baseline_within_time_limit":baseline_seconds<=PADDLE_RECOVERY_BASELINE_LIMIT_SECONDS,
        "estimated_time_available":remaining>=estimated_alternate+20.0,
    }
    metadata={
        "attempted":False,
        "selected_pass":"baseline",
        "baseline_seconds":round(baseline_seconds,3),
        "alternate_seconds":None,
        "eligibility":eligibility,
        "criteria":None,
    }
    if not all(eligibility.values()):
        return baseline_text,baseline_boxes,baseline_candidate,baseline_method,metadata

    metadata["attempted"]=True
    alternate=_run_paddle_recovery_pass(path,language,max(1.0,remaining-20.0))
    metadata["alternate_status"]=alternate.get("status","failed")
    metadata["alternate_seconds"]=round(float(alternate.get("seconds",0)),3)
    if alternate.get("status")!="completed":
        return baseline_text,baseline_boxes,baseline_candidate,baseline_method,metadata
    alternate_text=alternate["text"];alternate_boxes=alternate["boxes"]
    try:
        alternate_candidate,alternate_method=structured_extract(
            alternate_text,alternate_boxes,templates,[]
        )
    except Exception as parse_error:
        metadata["criteria"]={"alternate_parse_succeeded":False}
        metadata["alternate_parse_error"]=type(parse_error).__name__
        return baseline_text,baseline_boxes,baseline_candidate,baseline_method,metadata
    if alternate_candidate is None:
        metadata["criteria"]={"alternate_parse_succeeded":False}
        return baseline_text,baseline_boxes,baseline_candidate,baseline_method,metadata
    criteria,accepted=paddle_recovery_criteria(baseline_candidate,alternate_candidate)
    metadata["criteria"]={"alternate_parse_succeeded":True,**criteria}
    if not accepted:
        return baseline_text,baseline_boxes,baseline_candidate,baseline_method,metadata
    metadata["selected_pass"]="higher_resolution"
    return (
        alternate_text,alternate_boxes,alternate_candidate,
        f"{alternate_method}_higher_resolution",metadata,
    )


def docling(path,language="en"):
    from types import SimpleNamespace

    from docling.document_converter import DocumentConverter, PdfFormatOption, ImageFormatOption
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import (
        OcrMode,
        PdfPipelineOptions,
        RapidOcrOptions,
        TableFormerMode,
    )
    from .docling_extract import document_payload
    # Fully scanned documents need a full-page render for small invoice table
    # text. Mixed PDFs still need OCR, but only where native PDF cells do not
    # already cover a layout region. This avoids both the old all-or-nothing
    # mixed-PDF decision and duplicate OCR on native pages.
    native_pages=[]
    if path.suffix.lower()==".pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            native_pages=[len((page.extract_text() or "").strip())>=20 for page in pdf.pages]
    has_native_text=any(native_pages)
    needs_ocr=path.suffix.lower()!=".pdf" or not native_pages or not all(native_pages)
    ocr_mode=OcrMode.PDF_AWARE_LAYOUT_REGIONS if has_native_text else OcrMode.FULL_PAGE
    options=PdfPipelineOptions(do_ocr=needs_ocr,do_table_structure=True,generate_parsed_pages=True)
    options.table_structure_options.mode=TableFormerMode.ACCURATE
    options.table_structure_options.do_cell_matching=True
    # Paddle gives RapidOCR access to its current PP-OCRv6 English models in
    # Docling 2.133. The former torch backend was limited to PP-OCRv4 here.
    options.ocr_options=RapidOcrOptions(
        backend="paddle",lang=["ch" if language=="ch" else "iso:"+language],
        mode=ocr_mode,scale=3.0,model_size="small",
    )
    converter=DocumentConverter(format_options={InputFormat.PDF:PdfFormatOption(pipeline_options=options),
        InputFormat.IMAGE:ImageFormatOption(pipeline_options=options)})
    result=converter.convert(path)
    doc=result.document
    # RapidOCR stores its measured detections in textline_cells. Docling's
    # assembled document can merge many of those cells into a small number of
    # text blocks, while document_payload intentionally reads word_cells for
    # fine geometry. Present a non-mutating page view containing both native
    # words and OCR cells so table extraction retains the reader's real boxes.
    payload_pages=[]
    for page in result.pages:
        parsed=getattr(page,"parsed_page",None)
        if parsed is None:
            payload_pages.append(page)
            continue
        words=list(getattr(parsed,"word_cells",[]) or [])
        seen={id(cell) for cell in words}
        words.extend(cell for cell in (getattr(parsed,"textline_cells",[]) or [])
                     if getattr(cell,"from_ocr",False) and id(cell) not in seen)
        payload_pages.append(SimpleNamespace(
            parsed_page=SimpleNamespace(word_cells=words),
            size=getattr(page,"size",None),
        ))
    text,boxes,tables=document_payload(doc,payload_pages)
    for box in boxes:
        if box.get("from_ocr"):
            box["geometry"]="ocr_cell"
    return text,boxes,tables


def main():
    p=argparse.ArgumentParser();p.add_argument("--engine",choices=["invoice2data","paddleocr","docling"],required=True)
    p.add_argument("--file",type=Path,required=True);p.add_argument("--output",type=Path,required=True)
    p.add_argument("--templates",type=Path,action="append",default=[]);p.add_argument("--language",default="en")
    p.add_argument("--budget-seconds",type=int,default=PADDLE_WORKER_BUDGET_SECONDS)
    p.add_argument("--learned-templates",type=Path,action="append",default=[],
                   help="private directory of learned supplier templates (app.learned); missing dirs are ignored")
    p.add_argument("--learned-only",action="store_true",
                   help="diagnostic: read with the learned templates alone, skipping the built-in readers")
    p.add_argument("--text-only",action="store_true",
                   help="return the text layer and word boxes only (evidence for values read elsewhere); no parsing")
    args=p.parse_args()
    if not args.learned_templates and os.getenv("INV_STUDIO_LEARNED_TEMPLATES"):
        args.learned_templates=[Path(item) for item in os.getenv("INV_STUDIO_LEARNED_TEMPLATES").split(os.pathsep) if item]
    try:
        worker_started=time.monotonic();tables=[];parser_error=None;recovery=None;learned=None
        if args.engine=="invoice2data":
            if args.file.suffix.lower()==".pdf":text,boxes,tables=digital(args.file,with_tables=True)
            else:text,boxes=digital(args.file)
        elif args.engine=="paddleocr":
            baseline_started=time.monotonic()
            text,boxes=paddle(args.file,args.language)
            baseline_seconds=time.monotonic()-baseline_started
        else:text,boxes,tables=docling(args.file,args.language)
        if args.text_only:
            args.output.write_text(json.dumps({"text":text[:150000],"boxes":boxes[:10000],"invoice":None,
                "extraction_method":"text_only","tables":tables,"parser_error":None},default=str))
            return
        templates=templates_from(args.templates)
        learned_templates=templates_from(args.learned_templates)
        if args.file.suffix.lower()==".json":
            from .models import Invoice
            parsed=Invoice.model_validate_json(text).model_dump(mode="json")
            extraction_method="template"
        elif args.learned_only:
            try:
                parsed,meta=template_extract(text,learned_templates,with_meta=True) if text else (None,None)
                extraction_method="learned_template" if parsed is not None else "text_only"
                learned=None if meta is None else {"mode":"only","template":meta.get("supplier_key"),
                                                   "source_id":meta.get("source_id")}
            except Exception as parse_error:
                parsed=None;extraction_method="text_only";parser_error=type(parse_error).__name__
        else:
            try:
                parsed,extraction_method=structured_extract(text,boxes,templates,tables)
            except Exception as parse_error:
                parsed=None;extraction_method="text_only";parser_error=type(parse_error).__name__
        if args.engine=="paddleocr" and parser_error is None and not args.learned_only:
            text,boxes,parsed,extraction_method,recovery=maybe_recover_paddle(
                args.file,args.language,templates,text,boxes,parsed,extraction_method,
                baseline_seconds,time.monotonic()-worker_started,args.budget_seconds,
            )
        if learned_templates and parser_error is None and not args.learned_only:
            parsed,extraction_method,learned=apply_learned(text,learned_templates,parsed,extraction_method)
        payload={"text":text[:150000],"boxes":boxes[:10000],
            "invoice":parsed,"extraction_method":extraction_method,"tables":tables,
            "parser_error":parser_error}
        if recovery is not None:payload["recovery"]=recovery
        if learned is not None:payload["learned"]=learned
        args.output.write_text(json.dumps(payload,default=str))
    except Exception as e:
        # No document or credential content in a user-facing error.
        args.output.write_text(json.dumps({"error":type(e).__name__,"hint":"Engine could not read this document. Check its installation, model cache and input format."}))
        raise SystemExit(1)


if __name__=="__main__":main()
