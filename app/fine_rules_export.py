"""Review and target workbooks for the fine-grained rules engine.

The review workbook always exports and carries the rulebook's working sheets
with their exact headers. The target workbook keeps the three template tabs
unchanged and is refused unless every invoice is Approved (R-016).
"""

import io
import re
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.excel import HEADERS, text_cell
from app.fine_rules import FEEDBACK_COLUMNS, TARGET_HEADER, TARGET_LINE


RAW_INVOICE = ["Invoice File", "Invoice Number", "Line No.", "Supplier Name", "Description Raw", "Candidate VPN",
               "Candidate Barcode", "Quantity", "Unit Cost", "Line Amount", "Currency", "Source Row"]
WORKBENCH = ["Invoice Number", "Line No.", "Supplier Code", "Entity Suffix", "Market", "Location Code",
             "Location Type", "VPN", "Barcode", "ITEM_PARENT", "ITEM", "Brand", "Currency", "Match Method", "Rule ID",
             "Confidence", "Validation Status", "Exception ID", "Invoice Value Before Tax", "POGRN Quantity",
             "POGRN Value Before Tax", "PO Source", "POGRN RMS Order No"]
POGRN_VALIDATION = ["Invoice Number", "Supplier Name Raw", "Item Master Supplier Name", "Supplier Value",
                    "EBS Supplier Code 6D", "Invoice Quantity", "Invoice Value Before Tax", "POGRN Quantity",
                    "POGRN Value Before Tax", "Qty Match", "Pre-Tax Value Match", "POGRN RMS Order No",
                    "Derived PO Number", "PO Source", "POGRN Row Reference", "Validation Status", "Exception Reason",
                    "POGRN Location ID", "Location Type", "Market", "Aggregated QTY_RECEIVED", "Quantity Variance",
                    "Quantity Result", "Aggregated TOTAL COST", "Value Variance", "Value Result", "Receipt Dates",
                    "Supplier Check", "Location Check", "Market Check", "Candidate Pass Count"]
EXCEPTIONS = ["Exception ID", "Invoice Number", "Line No.", "Exception Type", "Description", "Candidates / Evidence",
              "Proposed Resolution", "Owner", "Status", "Resolution", "Approved By", "Resolved Date", "Rule ID",
              "Blocking"]
TAX = ["Invoice Number", "Line No.", "Tax Code", "Tax Rate", "Taxable Amount", "Tax Amount", "Country Code",
       "Source Row"]
LINEAGE = ["Invoice Number", "Target", "Line No.", "Original Value", "Transformed Value", "Rule ID", "Source",
           "Reference Row", "Confidence"]

NUMERIC = (int, float, Decimal)


def _sheet(workbook, name, columns, rows):
    sheet = workbook.create_sheet(name)
    sheet.append(columns)
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173D3B")
    for row in rows:
        r = sheet.max_row + 1
        for c, column in enumerate(columns, 1):
            value = row.get(column)
            if isinstance(value, NUMERIC) and not isinstance(value, bool):
                sheet.cell(r, c, value)
            else:
                # Identifiers stay text (leading zeros) and nothing is written as a formula.
                text_cell(sheet.cell(r, c), "" if value is None else value)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = max(14, min(48, len(str(column[0].value)) + 2))
    return sheet


def review_workbook(results, feedback=()):
    workbook = Workbook()
    workbook.remove(workbook.active)
    flat = lambda key: [row for result in results for row in result.get(key, [])]  # noqa: E731
    _sheet(workbook, "06_Raw_Invoice", RAW_INVOICE, flat("raw_invoice"))
    _sheet(workbook, "07_Match_Workbench", WORKBENCH, flat("workbench"))
    _sheet(workbook, "06A_POGRN_Validation", POGRN_VALIDATION, flat("pogrn_validation"))
    _sheet(workbook, "08_Exceptions", EXCEPTIONS,
           [{**e, "Blocking": "Yes" if e.get("blocking") else "No"} for e in flat("exceptions")])
    _sheet(workbook, "09_Output_Header", TARGET_HEADER, [r["header"] for r in results])
    _sheet(workbook, "10_Output_Lines", TARGET_LINE, flat("lines"))
    _sheet(workbook, "11_Output_Tax", TAX, flat("tax"))
    _sheet(workbook, "12_Feedback_Log", FEEDBACK_COLUMNS, list(feedback))
    _sheet(workbook, "Lineage", LINEAGE, [
        {"Invoice Number": r["header"].get("Document", ""), "Target": x["target"], "Line No.": x["line"] or "",
         "Original Value": x["original"], "Transformed Value": x["value"], "Rule ID": x["rule"], "Source": x["source"],
         "Reference Row": x["reference"], "Confidence": x["confidence"]}
        for r in results for x in r.get("lineage", [])])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def target_workbook(results):
    """R-016: the three template tabs, only when every invoice and line passed."""
    if not results or any(r["status"] != "Approved" for r in results):
        raise ValueError("Only invoices with status Approved can be exported to the target template")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, columns in HEADERS.items():
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    for result in results:
        h, t = result["header"], result["transaction"]
        sheet = workbook["Header"]
        row = sheet.max_row + 1
        values = [t, h["Document"], h["Supplier Site"], h["Order No"], h["Location"], h["Location Type"],
                  date.fromisoformat(h["Document Date"]), h["Net Amount"], h["Tax Amount"], "", "", "", ""]
        for c, value in enumerate(values, 1):
            if c in (2, 3, 4, 5, 6, 10, 11, 12, 13):
                text_cell(sheet.cell(row, c), value)
            else:
                sheet.cell(row, c, value)
        sheet.cell(row, 7).number_format = "m/d/yyyy"
        tax = workbook["Tax_Breakdown"]
        tax.append([t, None, h["Net Amount"]])
        text_cell(tax.cell(tax.max_row, 2), result["lines"][0]["Unit Tax Code"] if result["lines"] else "")
        details = workbook["Details"]
        for line in result["lines"]:
            details.append([t, None, None, line["Unit Cost"], line["Quantity"], None])
            n = details.max_row
            for c, value in ((2, line["Item"]), (3, line["UPC"]), (6, line["Unit Tax Code"])):
                text_cell(details.cell(n, c), value)
            details.cell(n, 3).number_format = "@"
            for c in (4, 5):
                details.cell(n, c).number_format = "0.0000"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def safe_filename(value):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value)[:80] or "invoices"
