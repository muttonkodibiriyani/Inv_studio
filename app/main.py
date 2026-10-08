import hashlib
import io
import json
import os
import re
import shutil
import threading
import secrets
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from contextvars import copy_context
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlencode, urlparse
import yaml
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import Field, SecretStr, ValidationError
from starlette.middleware.trustedhost import TrustedHostMiddleware
from starlette.concurrency import run_in_threadpool
from .authentication import CloudIdentity, actor
from .engines import capabilities, needs_scan_evidence, process, verify_scan
from .evidence import annotate as annotate_evidence, verify_with_ocr
from .learned import LearnedStore
from .excel import UPC_MODES, workbook, batch_workbook, rules_workbook
from . import target_check as tc
from .extraction_draft import extraction_workbook, extraction_batch_workbook
from .drafts import ManualDraftRequest, build_draft_workbook, draft_filename, draft_public_metadata
from .deletion import DeletionError, delete_invoices, remove_uploads
from .reference_lookup import ReferenceLookup
from .fine_rules import RulesConfig, decide_feedback, feedback_entry, ocr_read, run_batch
from .fine_rules_export import review_workbook, target_workbook
from .fine_rules_source import LookupRulesSource
from .product_candidates import ProductCandidates
from .matching import accepted_ids, enrich, key, owner_entries, rules_key, rules_validation, rules_view, validate
from .models import Invoice, Line, Policy, ProcessingOptions, StrictModel
from .oauth import ChatGPTAuth
from .providers import Providers
from .references import import_references, reference_workbook
from .store import Store

ROOT=Path(__file__).resolve().parent.parent
MIME_XLSX="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_UPLOAD=12_000_000


class Review(StrictModel):
    invoice: Invoice
    revision: int
    confirm: bool = False
    # Reviewer-entered target values for cells the fine rules left empty: {"header": {...}, "lines": {"1": {...}}}.
    entries: dict | None = None
    # The owner's pick among the rules' supplier_site_candidates (a supplier code); "" clears it.
    supplier_code: str | None = None


class Revision(StrictModel):
    revision: int


class TargetExportConfig(StrictModel):
    upc: str = Field(default="empty",pattern="^(barcode|empty)$")


class FineRulesRequest(StrictModel):
    job_ids: list[str] = Field(min_length=1,max_length=200)


class FeedbackRequest(StrictModel):
    invoice_line: str = Field(max_length=200)
    original: str = Field(default="",max_length=500)
    correction: str = Field(max_length=500)
    reason: str = Field(default="",max_length=1000)
    evidence: str = Field(max_length=2000)
    rule_candidate: str = Field(default="",max_length=200)


class FeedbackDecision(StrictModel):
    approver: str = Field(max_length=200)
    decision: str = Field(pattern="^(Approved|Rejected)$")
    version: str = Field(default="",max_length=80)


class ExtractionDraftRequest(Revision):
    acknowledge_unvalidated: bool = Field(default=False,strict=True)


class BatchItem(Revision):
    id: str


class Batch(StrictModel):
    jobs: list[BatchItem] = Field(min_length=1,max_length=100)


class ExtractionBatch(Batch):
    acknowledge_unvalidated: bool = Field(default=False,strict=True)


class DeleteItem(StrictModel):
    id: str = Field(min_length=1,max_length=200)
    revision: int = Field(ge=1,strict=True)


class DeleteJobsRequest(StrictModel):
    jobs: list[DeleteItem] = Field(min_length=1,max_length=100)
    confirm_permanent: bool = Field(default=False,strict=True)


class Retry(StrictModel):
    options: ProcessingOptions
    preflight_token: str


class Connection(StrictModel):
    api_key: SecretStr = Field(min_length=10,max_length=500)


class SubscriptionToken(StrictModel):
    setup_token: SecretStr = Field(min_length=20,max_length=4096)


class ChatGPTBundle(StrictModel):
    bundle: dict


class Settings(StrictModel):
    provider: str = "openai"
    model: str = Field(default="",max_length=150,pattern=r"^[A-Za-z0-9._:/-]*$")
    ai_fallback: bool = True


class Account(StrictModel):
    id: str


class SignIn(StrictModel):
    account_id: str | None = None


class FilePlan(StrictModel):
    name: str = Field(min_length=1,max_length=300)
    size: int = Field(ge=1,le=MAX_UPLOAD)


class Preflight(StrictModel):
    options: ProcessingOptions
    files: list[FilePlan] = Field(min_length=1,max_length=100)


class PreflightToken(StrictModel):
    token: str


class DemoRequest(StrictModel):
    preflight_token: str


def text_key(value):
    """A line text cell keyed for "is this the same line" checks. The review client sends text back from an
    <input type=text> (CR/LF dropped), trimmed and with '' as null; a client that shows a newline as a space sends
    'a b' instead (D134). Removing all whitespace makes both compare equal to the stored 'a\\nb' (D130)."""
    if not isinstance(value,str):return value
    return re.sub(r"\s+","",value) or None


BACKGROUND_WAIT=45  # seconds one request waits for a bulk download; under the 60 s Hosting proxy
BACKGROUND_KEEP=1800  # seconds an uncollected bulk download is kept


def create_app(data_dir=None):
    app=FastAPI(title="Inv Studio",version="0.1.0")
    # TEST-ONLY: the synthetic demo validation references (matching.enrich/validate).
    # Production values come only from the fine-rules result with evidence.
    demo_references=os.getenv("INV_STUDIO_DEMO_REFERENCES")=="1"
    if demo_references and os.getenv("INV_STUDIO_CLOUD")=="1":raise RuntimeError("INV_STUDIO_DEMO_REFERENCES is test-only and cannot run in the cloud")
    identity=CloudIdentity()
    store_type=Store
    if identity.cloud:
        from .cloud_store import PostgresStore
        store_type=PostgresStore
    store=store_type(Path(data_dir or os.getenv("INV_STUDIO_DATA",ROOT/".data")))
    if identity.cloud:store.sync_templates()
    learned=None
    if os.getenv("INV_STUDIO_LEARN","1")!="0":
        learned=LearnedStore(store.root,persist=getattr(store,"persist_blob",None) if identity.cloud else None,
                             remove=getattr(store,"remove_blob",None) if identity.cloud else None)
    auth=ChatGPTAuth(store);providers=Providers(store,auth,learned)
    lookup=ReferenceLookup(store)
    products=ProductCandidates(lookup)
    pool=ThreadPoolExecutor(max_workers=2,thread_name_prefix="invoice")
    slots=threading.BoundedSemaphore(20)
    plans={};plans_lock=threading.RLock()
    app.state.store=store;app.state.providers=providers;app.state.auth=auth;app.state.pool=pool;app.state.learned=learned
    allowed_origins={x.strip() for x in os.getenv("INV_STUDIO_ALLOWED_ORIGINS","").split(",") if x.strip()}
    hosts=["127.0.0.1","localhost","testserver"]
    if identity.cloud:hosts += ["*.run.app"]+[urlparse(x).netloc for x in allowed_origins]
    app.state.identity=identity
    app.add_middleware(TrustedHostMiddleware,allowed_hosts=hosts)
    # Recovery must inspect all persisted jobs, not the 200-entry UI history.
    with store.connection(True) as connection:
        for row in connection.execute("SELECT id,payload FROM jobs"):
            job=json.loads(row["payload"])
            if job["status"] in ("processing","queued"):
                job.update(status="error",error="The server restarted during processing. Retry this invoice.")
                store.job(job["id"],job,connection)

    @app.middleware("http")
    async def local_guard(request,call_next):
        claims=None
        if identity.cloud and request.url.path.startswith("/api/") and request.url.path!="/api/public-config":
            try:claims=await run_in_threadpool(identity.verify,request.headers.get("authorization"))
            except PermissionError as exc:return JSONResponse({"detail":str(exc)},403)
            except ValueError as exc:return JSONResponse({"detail":str(exc)},401)
            request.state.identity=claims
        if request.method not in ("GET","HEAD","OPTIONS"):
            if request.headers.get("x-studio-request")!="1":return JSONResponse({"detail":"Missing request protection header"},403)
            origin=request.headers.get("origin")
            if origin and urlparse(origin).netloc!=request.headers.get("host") and origin not in allowed_origins:
                return JSONResponse({"detail":"Cross-origin requests are not accepted"},403)
            try:size=int(request.headers.get("content-length","0"))
            except ValueError:return JSONResponse({"detail":"Invalid content length"},400)
            if size>MAX_UPLOAD+200000:return JSONResponse({"detail":"Maximum upload size is 12 MB"},413)
        actor_token=actor.set(claims["uid"] if claims else "local-operator")
        try:response=await call_next(request)
        finally:actor.reset(actor_token)
        response.headers["X-Content-Type-Options"]="nosniff"
        response.headers["Referrer-Policy"]="no-referrer"
        response.headers["Cache-Control"]="no-store"
        response.headers["Content-Security-Policy"]="default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data: blob:; frame-src 'self' blob:; connect-src 'self'; object-src 'none'; base-uri 'none'; frame-ancestors 'none'"
        if identity.cloud:
            auth_domain=os.getenv("INV_STUDIO_FIREBASE_AUTH_DOMAIN","")
            response.headers["Content-Security-Policy"]=("default-src 'self'; script-src 'self' https://apis.google.com; style-src 'self'; img-src 'self' data: blob: https://*.googleusercontent.com; frame-src 'self' blob: https://"+auth_domain+"; connect-src 'self' https://identitytoolkit.googleapis.com https://securetoken.googleapis.com https://www.googleapis.com; object-src 'none'; base-uri 'none'; frame-ancestors 'none'")
        return response

    @app.exception_handler(ValidationError)
    @app.exception_handler(RequestValidationError)
    async def invalid_request(_,exc):
        # Pydantic's default error includes the submitted input. Never echo a key.
        return JSONResponse({"detail":"Invalid input: "+"; ".join(".".join(map(str,e["loc"]))+": "+e["msg"] for e in exc.errors())},422)

    @app.exception_handler(ValueError)
    async def bad_value(_,exc):return JSONResponse({"detail":str(exc)[:400]},400)

    def references(c=None):return store.get("references",{},c) if demo_references else {}
    def policy(c=None):return Policy.model_validate(store.get("policy",{},c))
    def rule_signature(opts):
        return hashlib.sha256(json.dumps({"options":opts.model_dump(),"references":references().get("version"),"policy":policy().model_dump(mode="json")},sort_keys=True).encode()).hexdigest()
    def require_plan(token,opts,filename,size):
        with plans_lock:
            plan=plans.get(token)
            if not plan or plan["expires"]<time.time() or not plan["confirmed"]:
                raise HTTPException(409,"Review and confirm the processing rules before running this invoice")
            if plan["signature"]!=rule_signature(opts):raise HTTPException(409,"Rules or references changed. Review the processing plan again.")
            match=next((x for x in plan["files"] if x["name"]==filename and x["size"]==size),None)
            if not match:raise HTTPException(409,"This file is not in the confirmed processing plan")
            return plan,match
    def job_or_404(jid,c=None):
        j=store.job(jid,c=c)
        if not j:raise HTTPException(404,"Invoice not found")
        return j
    def assert_editable(j):
        if j.get("export_id"):raise HTTPException(409,"This invoice was exported. Download its existing export.")
        if j["status"] in ("queued","processing"):raise HTTPException(409,"Wait for extraction to finish")
    def evaluate(j,c=None,validation_context=None):
        inv=Invoice.model_validate(j["invoice"])
        if validation_context is None:
            refs=references(c);rules=policy(c)
            if c:led=store.ledger(c)
            else:
                with store.connection() as conn:led=store.ledger(conn)
        else:refs,rules,led=validation_context
        if j.get("export_id"):led=[x for x in led if x.get("job_id")!=j["id"]]
        if demo_references:j["validation"]=validate(inv,refs,rules,led,j.get("reviewed",False))
        else:j["validation"]=rules_validation(fresh_rules(j),led,j.get("reviewed",False))
        if not j.get("export_id") and j["status"] not in ("queued","processing","error"):
            j["status"]="ready" if j["validation"]["ready"] else "review"
        return j

    def rules_signature(c):
        # The stored result is stale once the mapping tables or the imported catalog change.
        sources=[r[0] for r in c.execute("SELECT payload FROM lookup_sources ORDER BY id")]
        # The UPC export mode decides what the target check scores in the Details UPC column.
        return hashlib.sha256(json.dumps([store.get("fine_rules_config",{},c),sources,target_upc(c)],sort_keys=True,default=str).encode()).hexdigest()
    def compute_rules(invoice,j,entries=None):
        """ONE fine-rules run for this invoice; review fields, banner and download all read its view."""
        with store.connection() as c:config=store.get("fine_rules_config",{},c);signature=rules_signature(c)
        entry={"invoice":invoice,"filename":j["filename"],"text":j.get("text",""),"boxes":j.get("boxes",[]),"job_id":j["id"]}
        if j.get("owner_supplier_code"):entry["owner_supplier_code"]=j["owner_supplier_code"]
        # Decision 44: lines read by local OCR or the AI from a scan may get the I/1, O/0 VPN lookup.
        entry["ocr_lines"]=ocr_read(j)
        result=run_batch([entry],LookupRulesSource(store),RulesConfig.from_dict(config))[0]
        entries=j.get("owner_entries") if entries is None else entries
        view=rules_view(plain(result),entry["text"],entry["boxes"],j.get("evidence"),entries)
        view.update(signature=signature,evidence_hash=evidence_hash(j),computed_at=datetime.now(timezone.utc).isoformat())
        entries,attribution=picked_site_entry(view,j,entries,j.get("owner_entry_attribution"))
        attach_target_check(view,config,entry["text"],entries,attribution)
        return view
    def picked_site_entry(view,j,entries,attribution):
        """A Supplier Site the rules took from the owner's supplier-code pick is scored as that owner's entry,
        attributed to the stored pick's actor and time; without them it stays an unattributed owner entry."""
        site=(view.get("fields") or {}).get("site") or {}
        first=(site.get("evidence") or [{}])[0]
        if first.get("kind")!="owner_entry" or first.get("rule")!="OWNER-PICK" or not j.get("owner_supplier_code"):
            return entries,attribution
        entries=owner_entries(entries) if entries else {"header":{},"lines":{}}
        entries={**entries,"header":{**entries["header"],"site":site.get("value")}}
        who=j.get("owner_supplier_pick") or {}
        return entries,{**(attribution or {}),"header:site":who}
    def attach_target_check(view,config,text,entries,attribution):
        """Read-only target-sheet check (app/target_check): deterministic, no AI or cloud call; the owner config
        (buyer_name included) is compared in memory only. Its failures are review-level issues (decision 12)."""
        sources=tc.Sources(tc.store_rows(store),config,text,tc.store_site_rows(store))
        with store.connection() as c:upc=target_upc(c)
        result=tc.check_view(view,sources,upc=upc,entries=entries,attribution=attribution)
        view["target_check"]=result
        view["issues"]=[*view.get("issues",[]),*tc.issues(result)]
    def evidence_hash(j):
        # Extraction evidence (boxes from a deferred OCR pass) can arrive without a revision change.
        return hashlib.sha256(json.dumps(j.get("evidence") or {},sort_keys=True,default=str).encode()).hexdigest()
    def fresh_rules(j):
        r=j.get("rules")
        return r if r and r.get("revision")==j["revision"] else None
    def mark_exported(j,**fields):
        """The export bumps the revision; rules that were current for the exported revision stay current, so the draft
        and the target check keep reading what was exported. Nothing is recomputed on an exported job; stale stays stale."""
        current=fresh_rules(j) is not None
        j.update(status="exported",revision=j["revision"]+1,**fields)
        if current:j["rules"]={**j["rules"],"revision":j["revision"]}
    def ensure_rules(jid):
        """Re-run the rules when the invoice revision, mapping tables or catalog changed."""
        if demo_references:return
        j=job_or_404(jid)
        if j["status"] in ("queued","processing","error") or j.get("export_id"):return
        with store.connection() as c:signature=rules_signature(c)
        if fresh_rules(j) and j["rules"].get("signature")==signature and j["rules"].get("evidence_hash")==evidence_hash(j):return
        view=compute_rules(Invoice.model_validate(j["invoice"]),j)
        with store.connection(True) as c:
            current=job_or_404(jid,c)
            if current["revision"]!=j["revision"] or evidence_hash(current)!=view["evidence_hash"]:return
            view["revision"]=current["revision"]
            # New review-level issues from changed tables were not seen by the reviewer: ask for a fresh confirm.
            reset=bool(current.get("reviewed")) and accepted_ids(current.get("rules"))!=accepted_ids(view)
            if reset:current["reviewed"]=False
            current["rules"]=view
            if not current.get("target_system") and not current.get("owner_entries"):current["target_system"]=view["target_check"]
            evaluate(current,c);store.job(jid,current,c)
            if reset:store.audit("review_reset",{"job_id":jid,"revision":current["revision"],"reason":"fine rules issues changed"},c)
            store.audit("fine_rules_applied",{"job_id":jid,"revision":current["revision"],"status":view["status"],
                        "item_lines":view["item_lines"],"approved":False},c)
    def scan_evidence(jid,path,opts,selected_engine):
        """After an AI-only scan read is saved: a local OCR pass that adds boxes and review flags to the
        job's evidence. The AI values, revision and status never change; the pass is recorded in the trace."""
        started=time.monotonic();boxes=None
        try:
            boxes,_=verify_scan(path,store.root,opts.language)
            entry={"engine":"paddleocr","status":"evidence","method":"ocr_text","seconds":round(time.monotonic()-started,2),
                   "reason":"Local OCR text layer read for evidence only; the AI values are unchanged"}
        except Exception as exc:
            entry={"engine":"paddleocr","status":"evidence_failed","seconds":round(time.monotonic()-started,2),
                   "reason":("Local OCR evidence pass failed: "+str(exc))[:240]}
        with plans_lock,store.connection(True) as c:
            job=job_or_404(jid,c)
            if job.get("status")!="review" or job.get("selected_engine")!=selected_engine:return
            if boxes is not None:
                job["evidence"]=verify_with_ocr(job.get("invoice") or {},job.get("evidence"),boxes)
            job["trace"]=list(job.get("trace") or [])+[entry]
            store.job(jid,job,c)

    def run(jid,opts):
        try:
            with plans_lock:
                job=job_or_404(jid)
                if job.get("confirmed_signature") != rule_signature(opts):
                    raise ValueError("Confirmed rules changed")
                job.update(status="processing",error=None);store.job(jid,job)
            def progress(engine,message,details=None):
                j=job_or_404(jid)
                previous=j.get("progress") or {}
                j["progress"]={"engine":engine,"message":message,
                    "started_at":previous.get("started_at") if previous.get("engine")==engine else datetime.now(timezone.utc).isoformat()}
                if details:
                    j["trace"]=details.get("trace",j.get("trace",[]))
                    j["progress"]["characters"]=details.get("characters",0)
                store.job(jid,j)
            if identity.cloud:store.ensure_blob(Path(job["path"]))
            result=process(Path(job["path"]),opts,store,providers.extract,progress)
            lines=Invoice.model_validate(result["invoice"]).model_dump(mode="json")["lines"]
            view=None
            if not demo_references:
                kept=reread_entries(job,lines)
                view=compute_rules(Invoice.model_validate(result["invoice"]),{**job,**result,**(kept[0] if kept else {})})
            with plans_lock,store.connection(True) as c:
                job=job_or_404(jid,c)
                unchanged=job.get("confirmed_signature")==rule_signature(opts)
                raw=Invoice.model_validate(result["invoice"])
                header_fields=sum(getattr(raw,name) not in (None,"") for name in Invoice.model_fields if name!="lines")
                extracted_lines=sum(any(value not in (None,"") for value in line.model_dump().values()) for line in raw.lines)
                outcome="fields_extracted" if header_fields or extracted_lines else "text_read" if result.get("text","").strip() else "extraction_failed"
                if demo_references:inv,provenance=enrich(raw,references(c)) if unchanged else (raw,[])
                else:inv,provenance=raw,[]
                kept=reread_entries(job,lines)
                split_note=(job.get("split") or {}).get("note")  # from the upload, so a re-read never stacks notes
                job.update(result);job.update(evidence=annotate_evidence(result.get("evidence"),result["invoice"],filename=job.get("filename")))
                if split_note:job["extraction_note"]=" ".join(x for x in (split_note,job.get("extraction_note")) if x)
                if kept:
                    job.update(kept[0])
                    note=f"{kept[1]} line correction{' was' if kept[1]==1 else 's were'} cleared because the invoice was read again."
                    job["extraction_note"]=" ".join(x for x in (note,job.get("extraction_note")) if x)
                job.update(invoice=inv.model_dump(mode="json"),provenance=provenance,
                            extraction_status=outcome,
                            status="review" if unchanged else "error",reviewed=False,progress=None,revision=job["revision"]+1,
                            error=None if unchanged else "Rules or references changed during extraction. Review a new processing plan and retry.")
                if view is not None:
                    job["rules"]={**view,"revision":job["revision"]}
                    if not job.get("owner_entries"):job["target_system"]=view["target_check"]
                evaluate(job,c);store.job(jid,job,c)
                store.audit(outcome,{"job_id":jid,"engine":job["selected_engine"],"trace":job["trace"],
                            "header_fields":header_fields,"line_items":extracted_lines,
                            "text_characters":len(result.get("text","")),"approved":False,
                            "confirmed_rules_unchanged":unchanged,"owner_line_entries_cleared":kept[1] if kept else 0},c)
            if unchanged and needs_scan_evidence(result,opts) and os.getenv("INV_STUDIO_VERIFY_SCANS","1")!="0":
                # The result is saved and reviewable; a failing evidence pass must not turn it into an error.
                try:scan_evidence(jid,Path(job["path"]),opts,result["selected_engine"])
                except Exception as exc:
                    store.audit("evidence_pass_failed",{"job_id":jid,"error_type":type(exc).__name__,"approved":False})
        except Exception as exc:
            job=job_or_404(jid);job.update(status="error",error=f"Processing failed ({type(exc).__name__}). Retry with another engine or review the document.",progress=None)
            store.job(jid,job)
            store.audit("extraction_failed",{"job_id":jid,"error_type":type(exc).__name__,"approved":False})
        finally:slots.release()

    def submit(content,filename,opts,token):
        if not slots.acquire(blocking=False):raise HTTPException(429,"Processing queue is full. Wait for an invoice to finish.")
        try:
            suffix=Path(filename).suffix.lower()
            if suffix not in (".pdf",".png",".jpg",".jpeg",".webp",".bmp",".tif",".tiff",".docx",".xlsx",".csv",".txt",".json"):
                raise ValueError("Supported: PDF, PNG, JPEG, WebP, BMP, TIFF, DOCX, XLSX, CSV, TXT and invoice JSON")
            if not content or len(content)>MAX_UPLOAD:raise ValueError("Invoice must be between 1 byte and 12 MB")
            if suffix==".pdf":
                import pypdfium2 as pdfium
                if not content.startswith(b"%PDF-"):raise ValueError("File is not a PDF")
                try:
                    pdf=pdfium.PdfDocument(content);count=len(pdf);pdf.close()
                except Exception:raise ValueError("PDF is damaged or encrypted") from None
                if not 1<=count<=20:raise ValueError("Use invoices with 1 to 20 pages")
            elif suffix in (".png",".jpg",".jpeg",".webp",".bmp",".tif",".tiff"):
                from PIL import Image
                try:
                    with Image.open(io.BytesIO(content)) as image:
                        if image.format not in ("PNG","JPEG","WEBP","BMP","TIFF") or image.width*image.height>25_000_000 or getattr(image,"n_frames",1)>20:raise ValueError()
                        image.verify()
                except Exception:raise ValueError("Image is invalid or exceeds 25 megapixels / 20 frames") from None
            elif suffix in (".docx",".xlsx"):
                import zipfile
                try:
                    with zipfile.ZipFile(io.BytesIO(content)) as archive:
                        if sum(x.file_size for x in archive.infolist())>30_000_000:raise ValueError()
                        marker="word/document.xml" if suffix==".docx" else "xl/workbook.xml"
                        if marker not in archive.namelist():raise ValueError()
                except Exception:raise ValueError("Office document is invalid or expands beyond 30 MB") from None
            else:
                try:content.decode("utf-8-sig")
                except UnicodeDecodeError:raise ValueError("Text/CSV/JSON invoices must use UTF-8 encoding") from None
            parts=[(content,Path(filename).name[:200],{})];refusal=None
            if suffix==".pdf":
                from .multi_invoice import analyse_pdf, part_filename, split_note, split_pdf
                try:outcome=analyse_pdf(content)
                except Exception:outcome=None
                if outcome is not None and outcome.multiple:
                    if outcome.refusal:refusal=outcome.refusal
                    else:
                        count=len(outcome.segments);name=Path(filename).name[:200]
                        parts=[(part,part_filename(name,n,count,segment.number or ""),
                                {"split":{"source_filename":name,"source_sha256":hashlib.sha256(content).hexdigest(),"part":n,"of":count,
                                          "invoice_number":segment.number,"pages":segment.pages,
                                          "note":split_note(n,count,outcome.numbers,segment.pages)},
                                 "extraction_note":split_note(n,count,outcome.numbers,segment.pages)})
                               for n,(part,segment) in enumerate(zip(split_pdf(content,outcome.segments),outcome.segments),1)]
            extra=len(parts)-1
            for n in range(extra):
                if not slots.acquire(blocking=False):
                    for _ in range(n):slots.release()
                    raise HTTPException(429,"Processing queue is full. Wait for an invoice to finish.")
            jobs=[]
            try:
                for part,part_name,fields in parts:
                    jid=uuid.uuid4().hex;path=store.root/"uploads"/(jid+suffix);path.write_bytes(part);os.chmod(path,0o600)
                    jobs.append({"id":jid,"filename":part_name,"size":len(part),"path":str(path),"sha256":hashlib.sha256(part).hexdigest(),
                         "created_at":datetime.now(timezone.utc).isoformat(),"status":"queued","options":opts.model_dump(),
                         "invoice":Invoice().model_dump(mode="json"),"revision":1,"reviewed":False,"trace":[],"provenance":[],"completeness":0,"selected_engine":"pending","validation":{"ready":False,"issues":[],"matches":[]},**fields})
                with plans_lock:
                    try:
                        plan,entry=require_plan(token,opts,filename,len(content))
                        for job in jobs:
                            if identity.cloud:store.persist_blob(Path(job["path"]))
                            job["confirmed_signature"]=plan["signature"]
                            if refusal:
                                job.update(status="error",error=refusal,extraction_status="multi_invoice_refused")
                                store.job(job["id"],job);store.audit("multi_invoice_refused",{"job_id":job["id"],"filename":job["filename"],"sha256":job["sha256"]})
                                continue
                            store.job(job["id"],job);store.audit("uploaded",{"job_id":job["id"],"filename":job["filename"],"sha256":job["sha256"],
                                                              **({"split":{k:v for k,v in job["split"].items() if k not in ("invoice_number","note")}} if job.get("split") else {})})
                        for job in jobs:
                            if not refusal:pool.submit(copy_context().run,run,job["id"],opts)
                        plan["files"].remove(entry)
                    except Exception:
                        for job in jobs:
                            if store.job(job["id"]):
                                job.update(status="error",error="Could not queue this invoice. Review the plan and retry.");store.job(job["id"],job)
                            else:Path(job["path"]).unlink(missing_ok=True)
                        raise
            except Exception:
                for _ in range(extra):slots.release()
                raise
            if refusal:slots.release()
            return public(jobs[0])
        except Exception:slots.release();raise

    def public(j,full=True):
        # The confirm-time snapshot and records are served by the target-check endpoints, not with every job.
        j={k:v for k,v in j.items() if k not in ("path","target_system","target_accuracy")}
        if not full:
            j.pop("text",None);j.pop("boxes",None);j.pop("evidence",None)
            check=(j.get("rules") or {}).get("target_check")
            if check:j["rules"]={**j["rules"],"target_check":{k:check[k] for k in ("counts","metric","holds","summary","upc")}}
        return j

    @app.get("/api/public-config")
    def public_config():return identity.config()

    @app.get("/api/session")
    def user_session(request:Request):return getattr(request.state,"identity",{"email":"local operator","uid":"local-operator"})

    @app.get("/api/state")
    def state():
        vertex_configured=bool(os.getenv("VERTEX_PROJECT_ID"))
        with store.connection() as c:
            accounts=store.get("chatgpt_accounts",[],c)
            active_account=store.get("chatgpt_active",c=c)
            settings=store.get("settings",Settings().model_dump(),c)
            if vertex_configured and not settings.get("model"):
                settings={**settings,"provider":"vertex","model":os.getenv("VERTEX_MODEL","gemini-3.7-flash")}
            connections={"openai":bool(store.secret("openai",c=c)),"anthropic":bool(store.secret("anthropic",c=c)),
                "chatgpt":any(a["id"]==active_account and a["connected"] for a in accounts),
                "claude_local":providers.claude_subscription.connected(c=c) and bool(shutil.which("claude")),"vertex":vertex_configured}
            jobs=store.jobs(c)
            refs=references(c);rules=policy(c);ledger=store.ledger(c)
            exports=[{"id":r["id"],"job_id":r["job_id"],**{k:v for k,v in json.loads(r["payload"]).items() if k in ("number","created_at")}} for r in c.execute("SELECT id,job_id,payload FROM exports ORDER BY rowid DESC")]
        return {"engines":capabilities(),"connections":connections,
                "accounts":accounts,"active_account":active_account,"settings":settings,
                "policy":rules.model_dump(mode="json"),"references":{"version":refs.get("version"),"imported_at":refs.get("imported_at"),"counts":{k:len(refs.get(k,[])) for k in ("sites","routes","items","orders","receipts","taxRules")}} if refs else None,
                "jobs":[public(evaluate(j,validation_context=(refs,rules,ledger)),False) for j in jobs],"exports":exports,
                "legacy_references":demo_references,
                "samples":{"demo":{"name":"SYNTHETIC-demo-invoice.pdf","size":(ROOT/"samples/invoice.pdf").stat().st_size}}}

    @app.post("/api/preflight")
    def preflight(body:Preflight):
        warnings=[];blocking=[];opts=body.options
        if opts.engine=="invoice2data":warnings.append("invoice2data reads native text and supplier templates. Scanned PDFs and images use PaddleOCR locally as the input reader; the trace names both components.")
        if opts.engine in ("auto","invoice2data","paddleocr"):
            warnings.append("Short scanned PDFs with unread quantities or prices may receive one higher-resolution local OCR pass, within a six-minute reader limit. It is retained only when consistency checks improve; all values still require review. Local OCR files wait their turn to keep the workspace responsive.")
        if demo_references and not references():warnings.append("No reference files loaded. Extraction can run, but Excel export stays on hold until references are imported and checked.")
        if not demo_references:
            with store.connection() as c:
                missing=[n for n,v in (("owner catalog",c.execute("SELECT 1 FROM lookup_sources LIMIT 1").fetchone()),("mapping tables",store.get("fine_rules_config",{},c))) if not v]
            if missing:warnings.append(f"No {' or '.join(missing)} imported. Extraction can run, but the rules cannot check this batch and Excel export stays on hold until they are imported.")
        if opts.ai_fallback or opts.engine=="ai":
            if not opts.model:warnings.append("No AI model selected. If local reading fails, this batch will wait for an AI connection or manual review.")
            warnings.append("AI fallback sends the invoice document or its extracted text to the selected provider. Provider usage limits and charges may apply.")
        warnings.append("Each file must contain one invoice. For a combined PDF containing several invoices, split it into separate files first.")
        warnings.append("Missing identity, price, quantity, receipt or tax evidence cannot be guessed. Exceptions require review.")
        token=secrets.token_urlsafe(32)
        with plans_lock:
            expired=[k for k,v in plans.items() if v["expires"]<time.time()]
            for k in expired:plans.pop(k)
            if len(plans)>200:raise HTTPException(429,"Too many pending plans; try again later")
            plans[token]={"expires":time.time()+600,"signature":rule_signature(opts),"confirmed":False,"files":[x.model_dump() for x in body.files]}
            summary={**opts.model_dump(),"reference_version":references().get("version"),"policy":policy().model_dump(mode="json"),"file_count":len(body.files)}
        return {"token":token,"summary":summary,"warnings":warnings,"blocking":blocking}

    @app.post("/api/preflight/confirm")
    def confirm_plan(body:PreflightToken):
        with plans_lock:
            plan=plans.get(body.token)
            if not plan or plan["expires"]<time.time():raise HTTPException(409,"Processing plan expired. Review it again.")
            plan["confirmed"]=True
        store.audit("processing_plan_confirmed",{"signature":plan["signature"],"files":plan["files"]})
        return {"confirmed":True}

    @app.post("/api/invoices")
    async def upload(file:UploadFile=File(...),options:str=Form("{}"),preflight_token:str=Form("")):
        content=await file.read(MAX_UPLOAD+1);name=file.filename or "invoice.pdf";opts=ProcessingOptions.model_validate_json(options)
        return submit(content,name,opts,preflight_token)

    @app.post("/api/jobs/delete")
    def delete_jobs(body:DeleteJobsRequest):
        if not body.confirm_permanent:
            raise HTTPException(400,"Confirm permanent invoice deletion")
        if len({item.id for item in body.jobs})!=len(body.jobs):
            raise ValueError("Select each invoice only once")
        try:
            with plans_lock:
                with store.connection(True) as connection:
                    outcome=delete_invoices(store,body.jobs,connection)
                # After the commit: a failed delete leaves every invoice and its source file whole.
                outcome["source_files_not_removed"]=remove_uploads(store,outcome.pop("upload_paths"))
        except DeletionError as error:
            raise HTTPException(error.status_code,error.detail) from None
        if learned is not None:
            try:learned.forget_sources([item.id for item in body.jobs])
            except Exception as error:store.audit("learned_failed",{"stage":"forget","error_type":type(error).__name__})
        return outcome

    @app.get("/api/jobs/{jid}")
    def get_job(jid:str):
        ensure_rules(jid)
        return public(evaluate(job_or_404(jid)))

    @app.get("/api/jobs/{jid}/document")
    def document(jid:str):
        j=job_or_404(jid)
        try:
            if identity.cloud:store.ensure_blob(Path(j["path"]))
        except Exception as error:
            if getattr(error,"code",None)!=404 and 404 not in getattr(error,"args",()):raise
        if not Path(j["path"]).is_file():raise HTTPException(404,"The source file of this invoice is no longer stored.")
        return FileResponse(j["path"],filename=j["filename"],content_disposition_type="inline")

    @app.post("/api/jobs/{jid}/extraction-draft")
    def download_extraction(jid:str,body:ExtractionDraftRequest):
        if not body.acknowledge_unvalidated:raise HTTPException(400,"Acknowledge that this is an unvalidated review copy")
        # The same current rules the target export reads, so the draft's Item is the export's Item (decision 56).
        ensure_rules(jid)
        j=job_or_404(jid)
        if j["status"] in ("queued","processing"):raise HTTPException(409,"Wait for extraction to finish")
        if j["revision"]!=body.revision:raise HTTPException(409,"Invoice changed. Refresh before downloading.")
        invoice=Invoice.model_validate(j["invoice"])
        if not invoice.number and not invoice.lines:raise HTTPException(409,"No extracted invoice fields are available")
        content=extraction_workbook(invoice,j["filename"],j["revision"],fresh_rules(j))
        store.audit("extraction_draft_downloaded",{"job_id":jid,"revision":j["revision"],"approved":False,"line_items":len(invoice.lines)})
        return Response(content,media_type=MIME_XLSX,headers={"Content-Disposition":f'attachment; filename="EXTRACTION_REVIEW_ONLY_{jid[:8]}.xlsx"'})

    def supplier_pick(j,body):
        """The owner's supplier code for the rules: only one of the candidates the rules last offered on this job.
        A changed supplier name drops an earlier pick; the rules re-check the pick against their own candidates."""
        if body.supplier_code is None:
            same=body.invoice.supplier_name==(j.get("invoice") or {}).get("supplier_name")
            return j.get("owner_supplier_code") if same else None
        code=body.supplier_code.strip()
        if not code:return None
        offered={str(c.get("supplier_code")) for c in (j.get("rules") or {}).get("supplier_site_candidates") or []}
        if code not in offered and code!=j.get("owner_supplier_code"):raise ValueError("Supplier code is not one of the candidates the rules offered")
        return code
    def carry_unchecked(invoice,stored):
        """D102(3): the review client does not send barcode_unchecked. The stored value stays on the line at the same
        index while that line still has no gtin and the same description and code (compared through text_key); a
        reviewer-typed gtin clears it, and a value the client sends is never kept: the field is the reader's alone (W1).
        Inserting or deleting a line above shifts the index, so the key no longer matches and the value is dropped, never
        moved onto another line."""
        old=(stored or {}).get("lines") or [];lines=[]
        for i,l in enumerate(invoice.lines):
            was=old[i] if i<len(old) and isinstance(old[i],dict) else {}
            same=not l.gtin and not was.get("gtin") and \
                (text_key(l.description),text_key(l.sku))==(text_key(was.get("description")),text_key(was.get("sku")))
            lines.append(l.model_copy(update={"barcode_unchecked":was.get("barcode_unchecked") if same else None}))
        return invoice.model_copy(update={"lines":lines})
    def aligned_evidence(invoice,j):
        """F5 (D146(2), D153): the reader's line reviews are kept by line index, so they stay only while every saved
        line is the stored line at the same index (same count, same description and sku through text_key). Any insert,
        delete or replacement drops them, header evidence stays."""
        evidence=j.get("evidence")
        old=(j.get("invoice") or {}).get("lines") or []
        def key(line):return text_key(line.get("description")),text_key(line.get("sku"))
        if not isinstance(evidence,dict) or not evidence.get("lines"):return evidence
        if len(old)==len(invoice.lines) and all(isinstance(was,dict) and key(was)==key(l.model_dump()) for was,l in zip(old,invoice.lines)):
            return evidence
        return {**evidence,"lines":[]}
    @app.post("/api/jobs/{jid}/review")
    def review(jid:str,body:Review):
        view=entries=pick=None;cleared=0
        if not demo_references:
            j=job_or_404(jid);assert_editable(j)
            if j["revision"]!=body.revision:raise HTTPException(409,"Invoice changed. Refresh before saving.")
            # Before the lines comparison below, so a client that drops the field does not clear the line entries.
            body.invoice=carry_unchecked(body.invoice,j["invoice"])
            if body.entries is not None:entries=owner_entries(body.entries)
            else:
                # Line entries are keyed by line number; they do not survive a change to the lines.
                entries=owner_entries(j.get("owner_entries"))
                if not same_lines(body.invoice.lines,j):
                    cleared=sum(len(cells or {}) for cells in (entries.get("lines") or {}).values());entries["lines"]={}
            attribution=entry_attribution(j.get("owner_entries"),entries,j.get("owner_entry_attribution"))
            pick=supplier_pick(j,body)
            picked_by=j.get("owner_supplier_pick") if pick and pick==j.get("owner_supplier_code") else {"actor":actor.get(),"at":datetime.now(timezone.utc).isoformat()} if pick else None
            view=compute_rules(body.invoice,{**j,"evidence":aligned_evidence(body.invoice,j),"owner_entry_attribution":attribution,"owner_supplier_code":pick,"owner_supplier_pick":picked_by},entries)
        with store.connection(True) as c:
            j=job_or_404(jid,c);assert_editable(j)
            if j["revision"]!=body.revision:raise HTTPException(409,"Invoice changed. Refresh before saving.")
            body.invoice=carry_unchecked(body.invoice,j["invoice"])
            inv,provenance=enrich(body.invoice,references(c)) if demo_references else (body.invoice,[])
            before=j["invoice"]
            j.update(evidence=aligned_evidence(body.invoice,j),invoice=inv.model_dump(mode="json"),reviewed=body.confirm,status="review",revision=j["revision"]+1)
            if view is not None:
                j["rules"]={**view,"revision":j["revision"]};j["owner_entries"]=entries;j["owner_entry_attribution"]=attribution
                if pick:j.update(owner_supplier_code=pick,owner_supplier_pick=picked_by)
                else:j.pop("owner_supplier_code",None);j.pop("owner_supplier_pick",None)
                if body.confirm:
                    # Confirm-time accuracy: per field, its status before and changed yes/no; never a value.
                    supplier=(view["fields"].get("site") or {}).get("value") or ""
                    j["target_accuracy"]=[{**r,"job_id":jid} for r in tc.confirm_records(j.get("target_system"),view["target_check"],supplier)]
                    store.audit("target_check_confirmed",{"job_id":jid,"revision":j["revision"],**tc.log_fields(view["target_check"]),
                                "changed":sorted({r["field"] for r in j["target_accuracy"] if r["changed"]})},c)
            # Keep original derivation evidence, add new derivations and record edits separately.
            j["provenance"]+=provenance
            evaluate(j,c);store.job(jid,j,c)
            store.audit("reviewed" if body.confirm else "edited",{"job_id":jid,"before":before,"after":j["invoice"],"revision":j["revision"],"reference_version":references(c).get("version"),
                        **({"owner_entries":entries,"accepted":j["validation"].get("accepted",[]),"owner_supplier_pick":bool(pick),
                           "owner_line_entries_cleared":cleared} if view is not None else {})},c)
        return public(j)

    @app.post("/api/jobs/{jid}/retry")
    def retry(jid:str,body:Retry):
        j=job_or_404(jid);assert_editable(j)
        if not slots.acquire(blocking=False):raise HTTPException(429,"Processing queue is full")
        try:
            with plans_lock:
                plan,entry=require_plan(body.preflight_token,body.options,j["filename"],j.get("size",Path(j["path"]).stat().st_size))
                with store.connection(True) as c:
                    j=job_or_404(jid,c);assert_editable(j)
                    j.update(status="queued",reviewed=False,revision=j["revision"]+1,options=body.options.model_dump(),confirmed_signature=plan["signature"])
                    store.job(jid,j,c);store.audit("retry",{"job_id":jid,"options":body.options.model_dump()},c)
                try:pool.submit(copy_context().run,run,jid,body.options)
                except Exception:
                    j.update(status="error",error="Could not queue this invoice. Review the plan and retry.");store.job(jid,j)
                    raise HTTPException(503,"Processing queue is unavailable") from None
                plan["files"].remove(entry)
            return public(j)
        except Exception:slots.release();raise

    @app.post("/api/jobs/{jid}/export")
    def export(jid:str,body:Revision):
        ensure_rules(jid)
        with store.connection(True) as c:
            j=job_or_404(jid,c)
            if j.get("export_id"):return {"id":j["export_id"],"url":"/api/exports/"+j["export_id"]}
            assert_editable(j)
            if j["revision"]!=body.revision:raise HTTPException(409,"Invoice changed. Refresh before exporting.")
            evaluate(j,c)
            if not j["validation"]["ready"]:raise HTTPException(409,"Invoice is held. Resolve every validation issue and confirm evidence review.")
            inv=Invoice.model_validate(j["invoice"]);eid=uuid.uuid4().hex
            if demo_references:
                content=workbook(inv,j["validation"])
                receipt={"id":eid,"job_id":jid,"invoice_key":key(inv),"number":inv.number,"created_at":datetime.now(timezone.utc).isoformat(),
                         "reference_version":references(c).get("version"),"policy":policy(c).model_dump(mode="json"),
                         "allocations":[m["allocation"] for m in j["validation"]["matches"]]}
            else:
                accepted=j["validation"].get("accepted",[])
                content,evidence=rules_workbook([{**j["rules"],"owner_accepted":accepted}],target_upc(c))
                content,target=with_checks_sheet(content,[j["rules"]])
                receipt={"id":eid,"job_id":jid,"invoice_key":rules_key(j["rules"]),"number":inv.number,"created_at":datetime.now(timezone.utc).isoformat(),
                         "source":"fine_rules","rules_signature":j["rules"]["signature"],"rules_revision":j["rules"]["revision"],
                         "rules_status":j["rules"]["status"],"owner_accepted":accepted,"owner_entries":j.get("owner_entries") or {},
                         "allocations":[],"evidence":evidence.get(1,{}),"target_check":target[1]}
            c.execute("INSERT INTO exports VALUES (?,?,?,?,?)",(eid,key(inv),jid,json.dumps(receipt),content))
            mark_exported(j,export_id=eid);store.job(jid,j,c);store.audit("exported",audit_receipt(receipt),c)
        learn_later(j)
        return {"id":eid,"url":"/api/exports/"+eid}

    def stored_lines(j):
        """The job's lines in the current model's shape: a job saved before a field was added (part_code, FT3) compares
        equal to the same lines read or saved now, so an unchanged save or re-read keeps its line entries (decision 65)."""
        return Invoice.model_validate(j["invoice"]).model_dump(mode="json")["lines"] if j.get("invoice") else None
    def same_lines(lines,j):
        """F1 (D132): a save keeps the line entries only when every line is the stored one, compared as Line values:
        numbers by value (a stored 1.00 equals a sent 1) and text through text_key, so the client's trimmed,
        newline-free text matches the stored cell. A job with no stored invoice counts as changed."""
        if not j.get("invoice"):return False
        old=Invoice.model_validate(j["invoice"]).lines
        def key(line):return tuple(text_key(getattr(line,f)) for f in Line.model_fields)
        return len(old)==len(lines) and all(key(a)==key(b) for a,b in zip(lines,old))
    def reread_entries(j,lines):
        """Line entries are keyed by line number, so a re-read that changes the lines drops them; header
        entries stay. Returns the job fields to keep and the count of cleared line cells, or None."""
        entries=j.get("owner_entries") or {}
        cleared=sum(len(cells or {}) for cells in (entries.get("lines") or {}).values())
        if not cleared or lines==stored_lines(j):return None
        attribution={k:v for k,v in (j.get("owner_entry_attribution") or {}).items() if not k.startswith("line:")}
        return {"owner_entries":{**entries,"lines":{}},"owner_entry_attribution":attribution},cleared
    def entry_attribution(before,after,previous):
        """Actor and time per owner entry; an unchanged entry keeps its first attribution."""
        before,previous,out=before or {},previous or {},{}
        now=datetime.now(timezone.utc).isoformat()
        for name,value in (after.get("header") or {}).items():
            k=f"header:{name}"
            out[k]=previous[k] if k in previous and (before.get("header") or {}).get(name)==value else {"actor":actor.get(),"at":now}
        for line,cells in (after.get("lines") or {}).items():
            for name,value in cells.items():
                k=f"line:{line}:{name}"
                old=((before.get("lines") or {}).get(str(line)) or {}).get(name)
                out[k]=previous[k] if k in previous and old==value else {"actor":actor.get(),"at":now}
        return out
    def with_checks_sheet(content,views):
        """Append the 'Checks' sheet (cells by status with evidence, row and workbook checks) to a written
        target workbook; returns it with each transaction's target-sheet accuracy for its receipt (counts only)."""
        results=[]
        for n,view in enumerate(views,1):
            check=view.get("target_check") or {}
            cells=[{**x,"value":str(n)} if x["column"]=="Transaction Number" else x for x in check.get("cells",[])]
            results.append({**check,"transaction":n,"cells":cells,"checks":check.get("checks",[])})
        currencies=[((v.get("fields") or {}).get("currency") or {}).get("value") for v in views]
        content,workbook_checks=tc.add_checks_sheet(content,results,currencies)
        content=order_date_rows(content,views)
        receipts={r["transaction"]:{"summary":r.get("summary",""),"counts":r.get("counts",{}),"metric":r.get("metric",{}),
                  "checks":{k["check"]:k["status"] for k in r["checks"]},
                  "workbook_checks":{k["check"]:k["status"] for k in workbook_checks}} for r in results}
        return content,receipts
    def order_date_rows(content,views):
        """Decisions 76/83: per transaction, a Checks row with its POGRN Order Date and one with its Receipt Date(s)
        check; evidence only, the template sheets are untouched."""
        rows=[]
        for n,v in enumerate(views,1):
            for column,d in (("Order Date",v.get("order_date")),("Receipt Date",v.get("receipt_date"))):
                if not d:continue
                status="flagged" if d.get("flag") else "filled" if d.get("value") is not None else "empty_flagged"
                reason="; ".join(x for x in (d.get("flag"),d.get("reason")) if x)
                detail="Evidence only; not a template column (decision 76)" if column=="Order Date" else \
                       "Date check only; flags, never blocks (decision 83)"
                rows.append({"Transaction Number":n,"Sheet":"(evidence)","Column":column,"Value":d.get("value") or "",
                             "Status":status,"Detail":detail,"Reason":reason,"Evidence Kind":d.get("evidence_kind") or "",
                             "Evidence Source":d.get("source") or "","Evidence Reference":d.get("reference") or "","Rule":d.get("rule") or ""})
        if not rows:return content
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(content));sheet=book["Checks"]
        for row in rows:
            sheet.append([None]*len(tc.CHECKS_COLUMNS));r=sheet.max_row
            for name,value in row.items():
                if value=="":continue
                cell=sheet.cell(r,tc.CHECKS_COLUMNS.index(name)+1)
                if isinstance(value,int):cell.value=value
                else:cell.value=str(value);cell.data_type="s"  # text only, never a formula
        buffer=io.BytesIO();book.save(buffer);return buffer.getvalue()
    def target_upc(c=None):return store.get("target_export",{},c).get("upc","empty")
    def audit_receipt(receipt):
        # Cell evidence holds owner rows; the audit keeps its size, the receipt keeps the rows.
        return {**{k:v for k,v in receipt.items() if k!="evidence"},**({"evidence_cells":len(receipt["evidence"])} if "evidence" in receipt else {})}

    @app.get("/api/exports/{eid}/evidence")
    def export_evidence(eid:str):
        with store.connection() as c:r=c.execute("SELECT payload FROM exports WHERE id=?",(eid,)).fetchone()
        if not r:raise HTTPException(404,"Export not found")
        receipt=json.loads(r[0])
        return {"id":eid,"job_id":receipt.get("job_id"),"transaction_number":receipt.get("transaction_number",1),"cells":receipt.get("evidence",{})}

    @app.get("/api/target-export/config")
    def target_export_config():return {"upc":target_upc(),"modes":list(UPC_MODES)}

    @app.post("/api/target-export/config")
    def save_target_export_config(body:TargetExportConfig):
        with store.connection(True) as c:
            store.set("target_export",body.model_dump(),c);store.audit("target_export_config_changed",body.model_dump(),c)
        return body.model_dump()

    @app.get("/api/exports/{eid}")
    def download(eid:str):
        with store.connection() as c:r=c.execute("SELECT workbook FROM exports WHERE id=?",(eid,)).fetchone()
        if not r:raise HTTPException(404,"Export not found")
        return Response(bytes(r[0]),media_type=MIME_XLSX,headers={"Content-Disposition":f'attachment; filename="Merch_Inv_{eid[:8]}.xlsx"'})

    @app.post("/api/exports/extraction-batch")
    def download_extraction_batch(body:ExtractionBatch):
        if not body.acknowledge_unvalidated:raise HTTPException(400,"Acknowledge that this is an unvalidated review copy")
        if len({x.id for x in body.jobs})!=len(body.jobs):raise ValueError("Select each invoice only once")
        for request in body.jobs:ensure_rules(request.id)
        entries=[];snapshots=[]
        with store.connection() as c:
            for request in body.jobs:
                j=job_or_404(request.id,c)
                if j["status"] in ("queued","processing"):raise HTTPException(409,"Wait for selected invoices to finish processing")
                if j["revision"]!=request.revision:raise HTTPException(409,"A selected invoice changed. Refresh before downloading.")
                invoice=Invoice.model_validate(j["invoice"])
                if not invoice.number and not invoice.lines:raise HTTPException(409,"A selected invoice has no extracted fields. Complete it before downloading.")
                entries.append((invoice,j["filename"],j["revision"],fresh_rules(j)))
                snapshots.append({"job_id":j["id"],"revision":j["revision"],"transaction_number":len(entries),"line_items":len(invoice.lines)})
        content=extraction_batch_workbook(entries)
        store.audit("extraction_batch_downloaded",{"invoices":snapshots,"approved":False,"count":len(entries)})
        return Response(content,media_type=MIME_XLSX,headers={"Content-Disposition":'attachment; filename="EXTRACTION_REVIEW_ONLY_BATCH.xlsx"'})

    @app.post("/api/exports/target-draft")
    def download_target_draft(body:Batch):
        """The target workbook for selected invoices nobody approved yet (decisions 67/70/72): the export's own builder
        in draft mode, so an empty cell is one the export would refuse. Read-only for the jobs: no export, no revision
        bump, no learning, no accuracy record. An invoice without current rules is skipped with its reason."""
        if demo_references:raise HTTPException(409,"The draft target workbook needs the fine rules")
        if len({x.id for x in body.jobs})!=len(body.jobs):raise ValueError("Select each invoice only once")
        views=[];included=[];skipped=[]
        for request in body.jobs:
            # Before ensure_rules, which 404s an unknown id: a deleted invoice is skipped, not a failed draft.
            if not store.job(request.id):skipped.append(({"id":request.id,"filename":""},"Invoice not found"));continue
            # The same refresh as opening the invoice; it returns early on exported or unfinished invoices.
            ensure_rules(request.id)
            j=store.job(request.id)
            if j["status"] in ("queued","processing","error"):reason="Extraction has not finished"
            elif j["revision"]!=request.revision:reason="Invoice changed. Refresh before downloading."
            elif not fresh_rules(j):reason="No current fine-rules result for this revision"
            else:reason=None
            if reason:skipped.append((j,reason));continue
            views.append(j["rules"]);included.append(j)
        if not views:raise HTTPException(409,"None of the selected invoices has a current fine-rules result. "+skipped[0][1])
        content,_=rules_workbook(views,target_upc(),draft=True)
        content,_=with_checks_sheet(content,views)
        content=draft_disclosure(content,skipped)
        store.audit("target_draft_downloaded",{"approved":False,"count":len(included),
                    "invoices":[{"job_id":j["id"],"revision":j["revision"],"transaction_number":n} for n,j in enumerate(included,1)],
                    "skipped":[{"job_id":j["id"],"reason":r} for j,r in skipped]})
        stamp=datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        return Response(content,media_type=MIME_XLSX,headers={"Content-Disposition":f'attachment; filename="DRAFT_Target_{stamp}.xlsx"',
                        "Cache-Control":"no-store","X-Draft-Included":str(len(included)),"X-Draft-Skipped":str(len(skipped))})
    def draft_disclosure(content,skipped):
        """The Checks sheet says the workbook is a draft and lists each skipped invoice with its reason."""
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(content));sheet=book["Checks"]
        rows=[("(draft)","Draft","Not exported, approved or confirmed. An empty cell has no evidenced value yet.")]
        rows+=[("(skipped)",j.get("filename") or j["id"],reason) for j,reason in skipped]
        for kind,column,reason in rows:
            sheet.append([None]*len(tc.CHECKS_COLUMNS));n=sheet.max_row
            for name,value in (("Sheet",kind),("Column",column),("Status","Draft" if kind=="(draft)" else "Skipped"),("Reason",reason)):
                cell=sheet.cell(n,tc.CHECKS_COLUMNS.index(name)+1);cell.value=str(value);cell.data_type="s"
        book.properties.title="DRAFT target workbook";buffer=io.BytesIO();book.save(buffer);return buffer.getvalue()

    @app.post("/api/exports/batch")
    def export_batch(body:Batch):
        if len({x.id for x in body.jobs})!=len(body.jobs):raise ValueError("Select each invoice only once")
        for request in body.jobs:ensure_rules(request.id)
        with store.connection(True) as c:
            selected=[];entries=[];ledger=store.ledger(c);batch_id=uuid.uuid4().hex
            for index,request in enumerate(body.jobs,1):
                j=job_or_404(request.id,c);assert_editable(j)
                if j["revision"]!=request.revision:raise HTTPException(409,f"Invoice {j['filename']} changed. Refresh before exporting.")
                inv=Invoice.model_validate(j["invoice"])
                if demo_references:result=validate(inv,references(c),policy(c),ledger,j.get("reviewed",False))
                else:result=rules_validation(fresh_rules(j),ledger,j.get("reviewed",False))
                if not result["ready"]:raise HTTPException(409,f"Invoice {j['filename']} is held: {result['issues'][0]['message']}")
                receipt={"id":uuid.uuid4().hex,"job_id":j["id"],"batch_id":batch_id,"transaction_number":index,
                         "invoice_key":key(inv) if demo_references else rules_key(j["rules"]),"number":inv.number,"created_at":datetime.now(timezone.utc).isoformat(),
                         "reference_version":references(c).get("version"),"policy":policy(c).model_dump(mode="json"),
                         "allocations":[m["allocation"] for m in result["matches"]] if demo_references else []}
                if not demo_references:
                    receipt.update(rules_status=j["rules"]["status"],owner_accepted=result["accepted"],owner_entries=j.get("owner_entries") or {})
                ledger.append(receipt);entries.append((inv,result) if demo_references else {**j["rules"],"owner_accepted":result["accepted"]})
                selected.append((j,receipt))
            if demo_references:content=batch_workbook(entries)
            else:
                content,evidence=rules_workbook(entries,target_upc(c))
                content,target=with_checks_sheet(content,entries)
                for _,receipt in selected:
                    row=receipt["transaction_number"];receipt["source"]="fine_rules"
                    receipt["evidence"]=evidence.get(row,{});receipt["target_check"]=target[row]
            c.execute("INSERT INTO batches VALUES (?,?,?)",(batch_id,json.dumps([r for _,r in selected]),content))
            for j,receipt in selected:
                c.execute("INSERT INTO exports VALUES (?,?,?,?,?)",(receipt["id"],receipt["invoice_key"],j["id"],json.dumps(receipt),content))
                mark_exported(j,export_id=receipt["id"],batch_id=batch_id,transaction_number=receipt["transaction_number"])
                store.job(j["id"],j,c);store.audit("exported",audit_receipt(receipt),c)
        for j,_ in selected:learn_later(j)
        return {"id":batch_id,"url":"/api/batches/"+batch_id,"count":len(selected)}

    @app.get("/api/batches/{bid}")
    def download_batch(bid:str):
        with store.connection() as c:r=c.execute("SELECT workbook FROM batches WHERE id=?",(bid,)).fetchone()
        if not r:raise HTTPException(404,"Batch not found")
        return Response(bytes(r[0]),media_type=MIME_XLSX,headers={"Content-Disposition":f'attachment; filename="Merch_Inv_Batch_{bid[:8]}.xlsx"'})

    @app.get("/api/jobs/{jid}/target-check")
    def target_check(jid:str):
        """Read-only: the invoice's target-sheet cells by status with evidence, its row checks and the owner line."""
        ensure_rules(jid)
        j=job_or_404(jid);check=(fresh_rules(j) or {}).get("target_check")
        if not check:raise HTTPException(409,"The target check runs with the fine rules; it is not available for this invoice yet")
        return {"job_id":jid,"revision":j["revision"],**check,"issues":tc.issues(check),"confirmed":bool(j.get("target_accuracy"))}

    @app.get("/api/target-check/accuracy")
    def target_accuracy(supplier:str|None=None):
        """Read-only confirm-time accuracy against owner truth (MEASURE reads this): per field and overall,
        by status before confirm and per supplier code, for 7 days, 30 days and all time. No values."""
        records=[];invoices=0
        with store.connection() as c:
            for r in c.execute("SELECT payload FROM jobs"):
                found=json.loads(r["payload"]).get("target_accuracy") or []
                invoices+=bool(found);records+=found
        return {**tc.accuracy_summary(records,supplier=supplier),"invoices":invoices}

    @app.get("/api/jobs/{jid}/audit")
    def audit(jid:str):
        job_or_404(jid)
        with store.connection() as c:
            return [{"event":r["event"],"at":r["at"],"payload":json.loads(r["payload"])} for r in c.execute("SELECT * FROM audit ORDER BY id") if json.loads(r["payload"]).get("job_id")==jid]

    def save_refs(refs):
        if not demo_references:raise HTTPException(404,"Validation references are test-only. Production uses the owner catalog and mapping tables.")
        with plans_lock,store.connection(True) as c:
            # Prior exports remain in the ledger. Baseline invoiced must exclude this app's exports.
            store.set("references",refs,c)
            for r in c.execute("SELECT id,payload FROM jobs").fetchall():
                j=json.loads(r["payload"])
                if not j.get("export_id"):
                    j.update(reviewed=False,revision=j["revision"]+1)
                    store.job(r["id"],j,c)
            store.audit("reference_import",{"version":refs["version"]},c)
        return {"version":refs["version"],"detail":"References imported. Unexported invoices require a new review."}

    @app.post("/api/references")
    async def reference_upload(file:UploadFile=File(...)):
        content=await file.read(MAX_UPLOAD+1)
        if len(content)>MAX_UPLOAD:raise HTTPException(413,"Reference file is too large")
        return save_refs(import_references(content,file.filename or ""))

    @app.post("/api/references/demo")
    def reference_demo():
        if not demo_references:raise HTTPException(404,"Demo validation references are test-only")
        return save_refs(import_references((ROOT/"samples/reference.json").read_bytes(),"reference.json"))

    @app.post("/api/demo")
    def demo(body:DemoRequest):
        content=(ROOT/"samples/invoice.pdf").read_bytes();name="SYNTHETIC-demo-invoice.pdf";opts=ProcessingOptions()
        return submit(content,name,opts,body.preflight_token)

    @app.get("/api/samples/reference.xlsx")
    def sample_reference():return Response(reference_workbook(json.loads((ROOT/"samples/reference.json").read_text())),media_type=MIME_XLSX,headers={"Content-Disposition":'attachment; filename="Reference_Template.xlsx"'})

    @app.get("/api/samples/invoice.pdf")
    def sample_invoice():return FileResponse(ROOT/"samples/invoice.pdf",filename="Synthetic_Invoice.pdf")

    @app.get("/api/samples/template.yml")
    def sample_template():return FileResponse(ROOT/"app/templates/demo.yml",filename="supplier-template.yml")

    @app.post("/api/templates")
    async def upload_template(file:UploadFile=File(...)):
        content=await file.read(65537)
        if len(content)>65536:raise ValueError("Template limit is 64 KB")
        data=yaml.safe_load(content)
        if not isinstance(data,dict) or not isinstance(data.get("keywords"),list) or not isinstance(data.get("fields"),dict):raise ValueError("Template needs keywords and fields")
        if not data["keywords"] or any(not isinstance(x,str) for x in data["keywords"]):raise ValueError("Template keywords must be non-empty strings")
        # invoice2data templates are declarative, and parsing runs in a timed subprocess.
        allowed=set(Invoice.model_fields)|{"invoice_number","amount"}
        if set(data["fields"])-allowed:raise ValueError("Template includes unsupported invoice fields")
        name=hashlib.sha256(content).hexdigest()[:20]+".yml"
        (store.root/"templates"/name).write_text(yaml.safe_dump(data))
        if identity.cloud:store.persist_blob(store.root/"templates"/name)
        store.audit("template_import",{"name":name})
        return {"name":name,"detail":"Supplier template registered. Test it on representative invoices."}

    @app.post("/api/policy")
    def save_policy(body:Policy):
        if any(not re.fullmatch(r"[A-Z]{3}",k) or not 0<=v<=4 for k,v in body.currency_decimals.items()):raise ValueError("Currency decimal mapping must use ISO-style codes and 0–4 places")
        with plans_lock,store.connection(True) as c:
            store.set("policy",body.model_dump(mode="json"),c)
            for row in c.execute("SELECT id,payload FROM jobs").fetchall():
                j=json.loads(row["payload"])
                if not j.get("export_id"):j.update(reviewed=False,revision=j["revision"]+1);store.job(row["id"],j,c)
            store.audit("policy_changed",body.model_dump(mode="json"),c)
        return body

    @app.post("/api/connections/{provider}")
    def connect(provider:str,body:Connection):
        if provider not in ("openai","anthropic"):raise ValueError("Use API keys only for OpenAI or Anthropic API connections")
        key=body.api_key.get_secret_value().strip()
        verified=providers.verify_key(provider,key)
        store.secret(provider,key);return {"connected":True,"verified":verified}

    @app.post("/api/subscriptions/claude/import")
    def import_claude_subscription(body:SubscriptionToken):
        result=providers.claude_subscription.connect(body.setup_token.get_secret_value())
        store.audit("subscription_connected",{"provider":"claude_local"})
        return result

    @app.delete("/api/subscriptions/claude")
    def disconnect_claude_subscription():
        providers.claude_subscription.disconnect()
        store.audit("subscription_disconnected",{"provider":"claude_local"})
        return {"connected":False}

    @app.post("/api/chatgpt/export")
    def export_chatgpt_bundle(body:Account):
        if identity.cloud:raise HTTPException(403,"Export a subscription registration from your local Invoice Studio installation")
        bundle=auth.export_bundle(body.id)
        return Response(json.dumps(bundle),media_type="application/json",headers={"Cache-Control":"no-store","Pragma":"no-cache","Content-Disposition":'attachment; filename="invoice-studio-chatgpt-credentials.json"'})

    @app.post("/api/chatgpt/import")
    def import_chatgpt_bundle(body:ChatGPTBundle):
        if not identity.cloud:raise HTTPException(403,"Credential transfer is for your authenticated self-hosted workspace")
        account=auth.import_bundle(body.bundle)
        store.audit("subscription_connected",{"provider":"chatgpt"})
        return account

    @app.get("/api/reference-lookup/summary")
    def lookup_summary():return lookup.summary()

    @app.get("/api/reference-lookup/search")
    def lookup_search(kind:str,q:str,cursor:str="",limit:int=25,site:str="",supplier:str="",po:str=""):
        return lookup.search(kind,q,cursor,limit,site=site,supplier=supplier,po=po)

    @app.get("/api/reference-lookup/products")
    def product_candidates(q:str,limit:int=25,site:str="",supplier:str="",po:str="",invoice_po:str="",price:str="",qty:str="",currency:str="",uom:str=""):
        return products.search(q,limit=limit,site=site,supplier=supplier,po=po,invoice_po=invoice_po,price=price,qty=qty,currency=currency,uom=uom)

    def fine_rules_results(job_ids):
        config=RulesConfig.from_dict(store.get("fine_rules_config",{}))
        entries=[]
        for jid in dict.fromkeys(job_ids):
            j=job_or_404(jid)
            if j["status"] in ("queued","processing"):raise HTTPException(409,"Wait for extraction to finish")
            entries.append({"invoice":Invoice.model_validate(j["invoice"]),"filename":j["filename"],
                            "text":j.get("text",""),"boxes":j.get("boxes",[]),"job_id":jid})
        return run_batch(entries,LookupRulesSource(store),config)

    def plain(value):
        # Money and quantities stay exact strings in JSON.
        if isinstance(value,dict):return {k:plain(v) for k,v in value.items() if not str(k).startswith("_")}
        if isinstance(value,(list,tuple)):return [plain(v) for v in value]
        if isinstance(value,set):return sorted(plain(v) for v in value)
        if isinstance(value,(int,float,str,bool)) or value is None:return value
        return str(value)

    @app.get("/api/fine-rules/config")
    def fine_rules_config():return store.get("fine_rules_config",{})

    @app.post("/api/fine-rules/config")
    async def save_fine_rules_config(request:Request):
        body=await request.json()
        if not isinstance(body,dict):raise ValueError("Fine-rules configuration must be an object")
        RulesConfig.from_dict(body)
        with store.connection(True) as c:
            store.set("fine_rules_config",body,c)
            # The configuration is private business data: the audit keeps field names and sizes, never values.
            store.audit("fine_rules_config_changed",{k:(len(v) if isinstance(v,(list,dict,str)) else type(v).__name__) for k,v in body.items()},c)
        return body

    @app.post("/api/fine-rules/run")
    def fine_rules_run(body:FineRulesRequest):
        results=fine_rules_results(body.job_ids)
        store.audit("fine_rules_run",{"job_ids":body.job_ids,"statuses":[r["status"] for r in results],"approved":False})
        return {"results":plain(results)}

    @app.post("/api/fine-rules/review.xlsx")
    def fine_rules_review(body:FineRulesRequest):
        content=review_workbook(fine_rules_results(body.job_ids),store.get("fine_rules_feedback",[]))
        store.audit("fine_rules_review_downloaded",{"job_ids":body.job_ids,"approved":False})
        return Response(content,media_type=MIME_XLSX,headers={"Content-Disposition":'attachment; filename="ULTA_Rules_Review.xlsx"',"Cache-Control":"no-store"})

    @app.post("/api/fine-rules/target.xlsx")
    def fine_rules_target(body:FineRulesRequest):
        if demo_references:
            results=fine_rules_results(body.job_ids)
            try:content=target_workbook(results)
            except ValueError as exc:raise HTTPException(409,str(exc))
        else:
            # The same stored result the review screen shows; never a second, different run.
            views=[]
            for jid in dict.fromkeys(body.job_ids):
                ensure_rules(jid);j=job_or_404(jid)
                if not fresh_rules(j):raise HTTPException(409,f"Invoice {j['filename']} has no current fine-rules result")
                views.append(j["rules"])
            try:content,_=rules_workbook(views,target_upc())
            except ValueError as exc:raise HTTPException(409,str(exc))
        store.audit("fine_rules_target_downloaded",{"job_ids":body.job_ids,"sha256":hashlib.sha256(content).hexdigest()})
        return Response(content,media_type=MIME_XLSX,headers={"Content-Disposition":'attachment; filename="ULTA_Target.xlsx"',"Cache-Control":"no-store"})

    # A bulk download over dozens of invoices outlasts the 60 s Firebase Hosting proxy (owner, 2026-10-08: 74-163 s), so
    # the browser got no file although the server finished. The work runs in a thread; the start and each poll wait at
    # most BACKGROUND_WAIT seconds for it, so a request is always open while it runs (Cloud Run CPU is request-billed).
    background_tasks={};background_lock=threading.Lock()
    background_kinds={"target-draft":(Batch,download_target_draft),"extraction-batch":(ExtractionBatch,download_extraction_batch),
                      "export-batch":(Batch,export_batch),"fine-rules-target":(FineRulesRequest,fine_rules_target),
                      "fine-rules-review":(FineRulesRequest,fine_rules_review)}
    def background_work(task,handler,body):
        try:
            result=handler(body)
            task["response"]=result if isinstance(result,Response) else JSONResponse(result)
        except HTTPException as exc:task["response"]=JSONResponse({"detail":exc.detail},exc.status_code)
        except (ValidationError,ValueError) as exc:task["response"]=JSONResponse({"detail":str(exc)[:400]},400)
        except Exception:task["response"]=JSONResponse({"detail":"The download failed on the server. Try again."},500)
        finally:task["done"].set()
    def background_answer(tid,task):
        if not task["done"].wait(BACKGROUND_WAIT):return JSONResponse({"id":tid,"status":"running"},202)
        with background_lock:background_tasks.pop(tid,None)
        return task["response"]

    @app.post("/api/background/{kind}")
    def background_start(kind:str,body:dict):
        if kind not in background_kinds:raise HTTPException(404,"Unknown download")
        model,handler=background_kinds[kind];request=model.model_validate(body)
        tid=uuid.uuid4().hex;task={"actor":actor.get(),"started":time.monotonic(),"done":threading.Event()}
        with background_lock:
            for old in [k for k,t in background_tasks.items() if time.monotonic()-t["started"]>BACKGROUND_KEEP]:del background_tasks[old]
            background_tasks[tid]=task
        threading.Thread(target=copy_context().run,args=(background_work,task,handler,request),daemon=True).start()
        return background_answer(tid,task)

    @app.get("/api/background/tasks/{tid}")
    def background_poll(tid:str):
        with background_lock:task=background_tasks.get(tid)
        if not task or task["actor"]!=actor.get():raise HTTPException(404,"This download is no longer running. Start it again.")
        return background_answer(tid,task)

    @app.get("/api/fine-rules/feedback")
    def fine_rules_feedback():return store.get("fine_rules_feedback",[])

    @app.post("/api/fine-rules/feedback")
    def add_fine_rules_feedback(body:FeedbackRequest):
        with store.connection(True) as c:
            log=feedback_entry(store.get("fine_rules_feedback",[],c),body.invoice_line,body.original,body.correction,
                               body.reason,body.evidence,body.rule_candidate)
            store.set("fine_rules_feedback",log,c);store.audit("fine_rules_feedback_proposed",log[-1],c)
        return log[-1]

    @app.post("/api/fine-rules/feedback/{fid}/decision")
    def decide_fine_rules_feedback(fid:str,body:FeedbackDecision):
        with store.connection(True) as c:
            log=decide_feedback(store.get("fine_rules_feedback",[],c),fid,body.approver,body.decision,body.version)
            store.set("fine_rules_feedback",log,c)
            entry=next(x for x in log if x["Feedback ID"]==fid);store.audit("fine_rules_feedback_decided",entry,c)
            if learned is not None:
                try:
                    if entry["Decision"]=="Approved":
                        record=learned.promote_feedback(entry,store.jobs(c))
                        store.audit("learned_correction",{"feedback_id":fid,"supplier_key":record["supplier_key"]},c)
                    else:learned.retract_feedback(fid)
                except Exception as error:store.audit("learned_failed",{"stage":"feedback","error_type":type(error).__name__},c)
        return entry

    def learn_from_job(j):
        # A verified (exported) invoice teaches its supplier's template and the AI examples.
        try:
            from .ocr_worker import structured_extract, templates_from
            try:baseline,_=structured_extract(j.get("text") or "",j.get("boxes") or [],
                                              templates_from([ROOT/"templates",store.root/"templates"]),j.get("tables") or [])
            except Exception:baseline=None
            outcome=learned.record_verified(j["invoice"],j.get("text") or "",j["id"],"export",baseline)
            store.audit("learned_recorded",{"job_id":j["id"],"supplier_key":outcome.get("supplier_key"),
                        "template":outcome.get("template"),"examples":outcome.get("examples"),
                        "fixes":outcome.get("fixes",[]),"reason":outcome.get("reason")})
        except Exception as error:
            store.audit("learned_failed",{"job_id":j.get("id"),"stage":"record","error_type":type(error).__name__})

    def learn_later(j):
        if learned is not None:pool.submit(copy_context().run,learn_from_job,dict(j))

    @app.get("/api/learned")
    def learned_summary():
        if learned is None:return {"enabled":False,"suppliers":[],"corrections":0,"unassigned_corrections":0}
        return {"enabled":True,**learned.summary()}

    @app.post("/api/learned/jobs/{jid}")
    def learn_job(jid:str):
        if learned is None:raise HTTPException(409,"Learning is switched off")
        j=job_or_404(jid)
        if j["status"]!="exported":raise HTTPException(409,"Only exported invoices count as verified")
        learn_from_job(j)
        return learned.summary()

    @app.delete("/api/learned/{key}")
    def forget_learned(key:str):
        if learned is None:raise HTTPException(409,"Learning is switched off")
        existed=learned.forget(key)
        store.audit("learned_forgotten",{"supplier_key":key,"existed":existed})
        return {"forgotten":existed}

    @app.post("/api/jobs/{jid}/draft")
    def manual_draft(jid:str,body:ManualDraftRequest):
        job_or_404(jid)
        content=build_draft_workbook(body)
        metadata=draft_public_metadata(body,job_id=jid)
        metadata["workbook_sha256"]=hashlib.sha256(content).hexdigest()
        store.audit("manual_draft_downloaded",metadata)
        return Response(content,media_type=MIME_XLSX,headers={"Content-Disposition":f'attachment; filename="{draft_filename(body)}"',"Cache-Control":"no-store"})

    @app.delete("/api/connections/{provider}")
    def disconnect(provider:str):
        if provider not in ("openai","anthropic"):raise ValueError("Unknown API connection")
        store.secret(provider,delete=True);return {"connected":False}

    @app.get("/api/connections/{provider}/models")
    def model_list(provider:str):
        if provider not in ("openai","anthropic","chatgpt","claude_local","vertex"):raise ValueError("Unknown provider")
        try:return {"models":providers.models(provider)}
        except ValueError:raise
        except Exception:raise ValueError("Could not retrieve models. Check the connection and try again.") from None

    @app.post("/api/settings")
    def save_settings(body:Settings):
        if body.provider not in ("openai","anthropic","chatgpt","claude_local","vertex"):raise ValueError("Unknown provider")
        store.set("settings",body.model_dump());return body

    @app.post("/api/chatgpt/start")
    def sign_in(body:SignIn,request:Request):
        if identity.cloud:raise ValueError("Use ChatGPT subscription sign-in from the local application")
        port=request.url.port or 80
        return {"url":auth.start(port,body.account_id)}

    @app.get("/auth/callback")
    def callback(request:Request):
        if identity.cloud:raise HTTPException(404,"Local subscription callback is not enabled in cloud mode")
        try:auth.finish(dict(request.query_params));return RedirectResponse("/?connected=chatgpt")
        except Exception:return RedirectResponse("/?"+urlencode({"connection_error":"ChatGPT sign-in could not be verified. Start again from AI connections."}))

    @app.post("/api/chatgpt/select")
    def select_account(body:Account):
        if not any(a["id"]==body.id and a["connected"] for a in auth.accounts()):raise ValueError("Connect this account before selecting it")
        store.set("chatgpt_active",body.id);return {"id":body.id}

    @app.delete("/api/chatgpt/{aid}")
    def remove_account(aid:str):return {"detail":auth.disconnect(aid)}

    @app.get("/health")
    def health():return {"status":"ok","version":"0.1.0"}

    @app.get("/")
    def home():return FileResponse(ROOT/"app/static/index.html")

    app.mount("/static",StaticFiles(directory=ROOT/"app/static"),name="static")
    return app


app=create_app()
